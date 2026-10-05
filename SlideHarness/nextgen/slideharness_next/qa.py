from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
import math
import re
import shutil
import subprocess

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE

from .models import LayoutPlan


PLACEHOLDER_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("lorem_ipsum", re.compile(r"\b(?:lorem|ipsum)\b", re.IGNORECASE)),
    ("todo_marker", re.compile(r"\bTODO\b", re.IGNORECASE)),
    ("xxx_marker", re.compile(r"(?<![A-Za-z0-9])X{3,}(?![A-Za-z0-9])", re.IGNORECASE)),
    ("insert_marker", re.compile(r"\[(?:insert|replace)[^\]]*\]", re.IGNORECASE)),
    ("powerpoint_placeholder", re.compile(r"\bclick to add (?:title|text|subtitle)\b", re.IGNORECASE)),
    ("template_instruction_en", re.compile(r"\bthis.{0,80}\b(?:page|slide).{0,80}\blayout\b", re.IGNORECASE | re.DOTALL)),
    ("template_instruction_ja", re.compile(r"(?:タイトルを入力してください|ここには.{0,80}(?:入力|記入)してください|この文章はダミーです)", re.DOTALL)),
)


@dataclass(frozen=True)
class Finding:
    category: str
    code: str
    severity: str
    slide_id: str | None
    message: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "category": self.category,
            "code": self.code,
            "severity": self.severity,
            "slide_id": self.slide_id,
            "message": self.message,
        }


def _shape_bounds(shape: Any, prs: Presentation) -> tuple[float, float, float, float]:
    return (
        shape.left / prs.slide_width,
        shape.top / prs.slide_height,
        shape.width / prs.slide_width,
        shape.height / prs.slide_height,
    )


def _content_shapes(slide: Any) -> list[Any]:
    return [shape for shape in slide.shapes if str(shape.name).startswith("SHN_CONTENT:")]


def _occupancy(shapes: list[Any], prs: Presentation, columns: int = 160, rows: int = 90) -> tuple[float, list[list[int]]]:
    grid = [[0 for _ in range(columns)] for _ in range(rows)]
    for shape in shapes:
        x, y, w, h = _shape_bounds(shape, prs)
        x0 = max(0, min(columns, int(x * columns)))
        y0 = max(0, min(rows, int(y * rows)))
        x1 = max(x0, min(columns, int((x + w) * columns + 0.999)))
        y1 = max(y0, min(rows, int((y + h) * rows + 0.999)))
        for row in range(y0, y1):
            for column in range(x0, x1):
                grid[row][column] = 1
    occupied = sum(sum(row) for row in grid)
    return occupied / (columns * rows), grid


def _largest_empty_rectangle(grid: list[list[int]]) -> float:
    if not grid or not grid[0]:
        return 1.0
    rows, columns = len(grid), len(grid[0])
    heights = [0] * columns
    largest = 0
    for row in range(rows):
        for column in range(columns):
            heights[column] = heights[column] + 1 if grid[row][column] == 0 else 0
        stack: list[int] = []
        for column in range(columns + 1):
            current = heights[column] if column < columns else 0
            while stack and current < heights[stack[-1]]:
                height = heights[stack.pop()]
                left = stack[-1] + 1 if stack else 0
                largest = max(largest, height * (column - left))
            stack.append(column)
    return largest / (rows * columns)


def _font_sizes(shape: Any) -> list[float]:
    sizes: list[float] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size:
                    sizes.append(run.font.size.pt)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size:
                            sizes.append(run.font.size.pt)
    return sizes


def _shape_text(shape: Any) -> str:
    chunks: list[str] = []
    if getattr(shape, "has_text_frame", False):
        chunks.append(shape.text_frame.text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _weighted_text_width(text: str, font_pt: float) -> float:
    units = 0.0
    for character in text:
        if character.isspace():
            units += 0.30
        elif ord(character) >= 0x2E80:
            units += 1.0
        elif character.isalnum():
            units += 0.56
        else:
            units += 0.45
    return units * font_pt


def _likely_text_overflow(shape: Any) -> bool:
    """Conservatively detect fixed-size text boxes that cannot contain their text."""

    if not getattr(shape, "has_text_frame", False) or not shape.text_frame.text.strip():
        return False
    frame = shape.text_frame
    if frame.auto_size != MSO_AUTO_SIZE.NONE:
        return False
    sizes = _font_sizes(shape)
    if not sizes:
        return False
    font_pt = max(sizes)
    emu_per_point = 12700
    available_width = (shape.width - frame.margin_left - frame.margin_right) / emu_per_point
    available_height = (shape.height - frame.margin_top - frame.margin_bottom) / emu_per_point
    if available_width <= 0 or available_height <= 0:
        return True
    line_count = 0
    for paragraph in frame.paragraphs:
        explicit_lines = paragraph.text.splitlines() or [""]
        for line in explicit_lines:
            width = _weighted_text_width(line, font_pt)
            if frame.word_wrap is False:
                if width > available_width * 1.05:
                    return True
                line_count += 1
            else:
                line_count += max(1, math.ceil(width / available_width))
    required_height = line_count * font_pt * 1.18
    # Source templates often use intentionally tight text frames whose glyphs
    # extend beyond the nominal box without visual clipping. Block only an
    # unmistakable capacity mismatch; the render layer handles subtler cases.
    return required_height > available_height * 1.90


def _placeholder_matches(slide: Any) -> list[tuple[str, str]]:
    matches: list[tuple[str, str]] = []
    for shape in slide.shapes:
        text = _shape_text(shape)
        for code, pattern in PLACEHOLDER_PATTERNS:
            if pattern.search(text):
                matches.append((str(shape.name), code))
    return matches


def inspect_pptx(
    pptx_path: str | Path,
    plans: list[LayoutPlan],
    composition_mode: str = "reconstruct",
    allow_shadows: bool = False,
    information_density: str = "medium",
    decorativeness: str = "medium",
    structured_visual_mode: str = "native",
    slide_policies: dict[str, Any] | None = None,
) -> dict[str, Any]:
    if information_density not in {"low", "medium", "high"}:
        raise ValueError(f"unsupported information density: {information_density}")
    if decorativeness not in {"low", "medium", "high"}:
        raise ValueError(f"unsupported decorativeness: {decorativeness}")
    if structured_visual_mode not in {"auto", "native", "generated-image"}:
        raise ValueError(f"unsupported structured visual mode: {structured_visual_mode}")
    slide_policies = slide_policies or {}
    prs = Presentation(str(pptx_path))
    findings: list[Finding] = []
    slides_report: list[dict[str, Any]] = []
    if len(prs.slides) != len(plans):
        findings.append(Finding("Technical", "slide_count_mismatch", "blocking", None, f"expected {len(plans)} slides, found {len(prs.slides)}"))
    for index, slide in enumerate(prs.slides):
        plan = plans[index] if index < len(plans) else None
        slide_id = plan.slide_id if plan else f"S{index + 1:03d}"
        effective_visual_mode = str(slide_policies.get(slide_id, {}).get("structured_visual_mode") or structured_visual_mode)
        shapes = _content_shapes(slide)
        occupancy, grid = _occupancy(shapes, prs)
        largest_void = _largest_empty_rectangle(grid)
        sparse_recipe = bool(plan and plan.recipe in {"cover", "section", "statement", "closing_cta"})
        if composition_mode == "source-slides" and sparse_recipe:
            minimum, maximum, void_limit = (0.06, 0.90, 0.70)
        elif composition_mode == "source-slides":
            minimum, maximum, void_limit = (0.22, 0.95, 0.45)
        elif plan and plan.recipe == "image_text":
            minimum, maximum, void_limit = (0.60, 0.95, 0.35)
        elif sparse_recipe:
            minimum, maximum, void_limit = (0.30, 0.70, 0.35)
        else:
            minimum, maximum, void_limit = (0.55, 0.90, 0.35)
        if occupancy < minimum:
            findings.append(Finding("Design", "low_meaningful_occupancy", "blocking", slide_id, f"meaningful occupancy {occupancy:.1%} is below {minimum:.0%}"))
        if occupancy > maximum:
            findings.append(Finding("Design", "overcrowded", "blocking", slide_id, f"meaningful occupancy {occupancy:.1%} exceeds {maximum:.0%}"))
        if largest_void > void_limit:
            findings.append(Finding("Design", "unexplained_large_void", "blocking", slide_id, f"largest empty rectangle is {largest_void:.1%} of slide"))
        if plan and plan.recipe == "title_body":
            body_characters = sum(len(str(item)) for element in plan.elements if element.kind == "body" for item in element.data.get("items", []))
            if plan.expected_item_count < 2 and body_characters < 140:
                findings.append(Finding("Design", "recipe_underfill", "blocking", slide_id, "title_body has too little content; use statement or a smaller recipe"))
        if plan and plan.recipe == "card_grid_2x2" and plan.expected_item_count not in {3, 4}:
            findings.append(Finding("Design", "recipe_item_count_mismatch", "blocking", slide_id, "card_grid_2x2 requires 3 or 4 supporting units"))
        if plan and plan.recipe == "three_column" and plan.expected_item_count != 3:
            findings.append(Finding("Design", "recipe_item_count_mismatch", "blocking", slide_id, "three_column requires exactly 3 supporting units"))
        if plan and plan.recipe == "two_column_equal" and plan.expected_item_count != 2:
            findings.append(Finding("Design", "recipe_item_count_mismatch", "blocking", slide_id, "two_column_equal requires exactly 2 supporting units"))
        out_of_bounds = []
        unresolved = []
        small_fonts = []
        likely_overflow = []
        empty_containers = []
        placeholder_matches = _placeholder_matches(slide)
        drop_shadow_count = sum(len(slide._element.xpath(f".//a:{effect_name}")) for effect_name in ("outerShdw", "prstShdw"))
        for shape in shapes:
            x, y, w, h = _shape_bounds(shape, prs)
            if x < -0.001 or y < -0.001 or x + w > 1.001 or y + h > 1.001:
                out_of_bounds.append(shape.name)
            if "UNRESOLVED_ASSET" in str(shape.name):
                unresolved.append(shape.name)
            for size in _font_sizes(shape):
                if size < 10:
                    small_fonts.append((shape.name, size))
            if _likely_text_overflow(shape):
                likely_overflow.append(shape.name)
            semantic_name = str(shape.name)
            if (
                getattr(shape, "has_text_frame", False)
                and shape.text_frame.text.strip() == ""
                and not any(term in semantic_name for term in ("-marker", "-accent", "-arrow", "fallback"))
                and any(term in semantic_name for term in ("-card-", "-column-", "-metric-", "-callout", "-step-"))
            ):
                empty_containers.append(shape.name)
        if out_of_bounds:
            findings.append(Finding("Technical", "out_of_bounds", "blocking", slide_id, f"{len(out_of_bounds)} content shapes exceed slide bounds"))
        if unresolved:
            findings.append(Finding("Content", "unresolved_asset", "blocking", slide_id, f"{len(unresolved)} generated-image request(s) are unresolved"))
        if small_fonts:
            findings.append(Finding("Design", "tiny_text", "blocking", slide_id, f"{len(small_fonts)} text run(s) are below 10 pt"))
        if likely_overflow:
            findings.append(Finding("Design", "likely_text_overflow", "blocking", slide_id, f"{len(likely_overflow)} fixed-size text container(s) are too small for their text"))
        if empty_containers:
            findings.append(Finding("Design", "empty_container", "blocking", slide_id, f"{len(empty_containers)} empty semantic container(s)"))
        if placeholder_matches:
            details = ", ".join(f"{name} ({code})" for name, code in placeholder_matches[:5])
            findings.append(Finding("Content", "leftover_placeholder", "blocking", slide_id, f"{len(placeholder_matches)} placeholder marker(s) remain: {details}"))
        if drop_shadow_count and not allow_shadows:
            findings.append(Finding("Design", "unrequested_drop_shadow", "blocking", slide_id, f"{drop_shadow_count} drop shadow effect(s) remain without an explicit instruction"))
        editable = {
            "charts": sum(1 for shape in slide.shapes if getattr(shape, "has_chart", False)),
            "tables": sum(1 for shape in slide.shapes if getattr(shape, "has_table", False)),
            "process_shapes": sum(1 for shape in slide.shapes if "-step-" in str(shape.name)),
        }
        picture_count = sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        if plan and effective_visual_mode == "native" and plan.recipe == "generated_visual":
            findings.append(Finding("Technical", "structured_visual_not_native", "blocking", slide_id, "native mode requires editable PowerPoint chart or flow objects"))
        if plan and effective_visual_mode == "generated-image" and plan.recipe in {"native_bar_chart", "native_chart", "process"}:
            findings.append(Finding("Technical", "structured_visual_not_generated", "blocking", slide_id, "generated-image mode requires the generated_visual recipe"))
        if plan and plan.recipe == "generated_visual" and picture_count == 0:
            findings.append(Finding("Technical", "generated_visual_missing_image", "blocking", slide_id, "generated visual recipe produced no image"))
        if plan and plan.recipe in {"native_bar_chart", "native_chart"} and editable["charts"] == 0:
            findings.append(Finding("Technical", "chart_not_native", "blocking", slide_id, "native chart recipe produced no editable chart"))
        if plan and plan.recipe == "table" and editable["tables"] == 0:
            findings.append(Finding("Technical", "table_not_native", "blocking", slide_id, "table recipe produced no editable table"))
        if plan and plan.recipe == "process" and editable["process_shapes"] < plan.expected_item_count:
            findings.append(Finding("Technical", "flow_not_editable", "blocking", slide_id, "process recipe produced too few editable step shapes"))
        slides_report.append(
            {
                "slide_id": slide_id,
                "recipe": plan.recipe if plan else None,
                "meaningful_occupancy": round(occupancy, 4),
                "largest_empty_rectangle": round(largest_void, 4),
                "occupancy_policy": composition_mode,
                "information_density": information_density,
                "decorativeness": decorativeness,
                "structured_visual_mode": effective_visual_mode,
                "content_shape_count": len(shapes),
                "drop_shadow_count": drop_shadow_count,
                "placeholder_match_count": len(placeholder_matches),
                "likely_text_overflow_count": len(likely_overflow),
                "editable": editable,
            }
        )
    blocking = [finding for finding in findings if finding.severity == "blocking"]
    return {
        "schema_version": "0.11-mvp.4",
        "status": "PASS" if not blocking else "FAIL",
        "summary": {"slide_count": len(prs.slides), "blocking_count": len(blocking), "finding_count": len(findings)},
        "slides": slides_report,
        "findings": [finding.to_dict() for finding in findings],
    }


def _montage(image_paths: list[Path], target: Path) -> Path:
    thumbnails: list[Image.Image] = []
    thumb_width = 720
    for path in image_paths:
        image = Image.open(path).convert("RGB")
        height = round(image.height * thumb_width / image.width)
        thumbnails.append(image.resize((thumb_width, height)))
    if not thumbnails:
        raise ValueError("no rendered slide images")
    gap = 24
    columns = min(2, len(thumbnails))
    rows = (len(thumbnails) + columns - 1) // columns
    cell_height = max(image.height for image in thumbnails)
    canvas = Image.new("RGB", (columns * thumb_width + (columns + 1) * gap, rows * cell_height + (rows + 1) * gap), "#D9DDE3")
    for index, image in enumerate(thumbnails):
        x = gap + (index % columns) * (thumb_width + gap)
        y = gap + (index // columns) * (cell_height + gap)
        canvas.paste(image, (x, y))
    target.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(target)
    for image in thumbnails:
        image.close()
    return target


def render_pptx(pptx_path: str | Path, render_dir: str | Path) -> dict[str, Any]:
    source = Path(pptx_path).resolve()
    target_dir = Path(render_dir).resolve()
    target_dir.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftocairo = shutil.which("pdftocairo")
    if not soffice or not pdftocairo:
        return {"status": "SKIP", "reason": "soffice or pdftocairo is unavailable", "images": []}
    for stale in [*target_dir.glob("slide-*.png"), target_dir / f"{source.stem}.pdf", target_dir / "montage.png"]:
        if stale.is_file():
            stale.unlink()
    with TemporaryDirectory(prefix="slideharness-next-lo-") as profile:
        profile_uri = Path(profile).resolve().as_uri()
        conversion = subprocess.run(
            [soffice, f"-env:UserInstallation={profile_uri}", "--headless", "--convert-to", "pdf", "--outdir", str(target_dir), str(source)],
            capture_output=True,
            text=True,
            timeout=90,
        )
    pdf_path = target_dir / f"{source.stem}.pdf"
    if conversion.returncode != 0 or not pdf_path.is_file():
        return {"status": "FAIL", "reason": conversion.stderr or conversion.stdout or "PDF conversion failed", "images": []}
    prefix = target_dir / "slide"
    raster = subprocess.run([pdftocairo, "-png", "-r", "110", str(pdf_path), str(prefix)], capture_output=True, text=True, timeout=90)
    image_paths = sorted(target_dir.glob("slide-*.png"))
    findings: list[dict[str, Any]] = []
    for index, path in enumerate(image_paths, start=1):
        with Image.open(path).convert("RGB") as image:
            extrema = ImageStat.Stat(image).extrema
            dynamic_range = max(high - low for low, high in extrema)
            if dynamic_range < 12:
                findings.append({"category": "Technical", "code": "blank_render", "severity": "blocking", "slide_id": f"S{index:03d}", "message": f"rendered slide dynamic range is only {dynamic_range}"})
    if raster.returncode != 0 or not image_paths:
        findings.append({"category": "Technical", "code": "rasterization_failed", "severity": "blocking", "slide_id": None, "message": raster.stderr or "no rendered PNG files"})
    montage_path = _montage(image_paths, target_dir / "montage.png") if image_paths else None
    return {
        "status": "PASS" if not findings else "FAIL",
        "pdf": str(pdf_path),
        "images": [str(path) for path in image_paths],
        "montage": str(montage_path) if montage_path else None,
        "findings": findings,
    }
