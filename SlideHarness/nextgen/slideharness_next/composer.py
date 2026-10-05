from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any
import re

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from .models import ChromeElement, DeckSpec, LayoutPlan, PlannedElement, TemplateSystem
from .charts import choose_chart_type
from .provenance import apply_source_notes
from .source_adapter import SourceSlideEvidence, clone_slide, populate_source_slide


def _rgb(value: str) -> RGBColor:
    cleaned = (value or "000000").lstrip("#")[:6]
    return RGBColor.from_string(cleaned)


def _mix(left: str, right: str, amount: float) -> str:
    a = tuple(int(left[i : i + 2], 16) for i in (0, 2, 4))
    b = tuple(int(right[i : i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(x * (1 - amount) + y * amount):02X}" for x, y in zip(a, b))


def _luminance(value: str) -> float:
    channels = [int(value[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    linear = [channel / 12.92 if channel <= 0.03928 else ((channel + 0.055) / 1.055) ** 2.4 for channel in channels]
    return linear[0] * 0.2126 + linear[1] * 0.7152 + linear[2] * 0.0722


def _contrast_text(background: str) -> str:
    return "FFFFFF" if _luminance(background) < 0.44 else "111827"


def _coords(prs: Presentation, element: PlannedElement) -> tuple[int, int, int, int]:
    return (
        round(prs.slide_width * element.x),
        round(prs.slide_height * element.y),
        round(prs.slide_width * element.w),
        round(prs.slide_height * element.h),
    )


def _name(shape: Any, element: PlannedElement, suffix: str = "") -> None:
    shape.name = f"SHN_CONTENT:{element.element_id}{suffix}"


def _font_scale(prs: Presentation) -> float:
    standard_width = Inches(13.333333)
    standard_height = Inches(7.5)
    return max(0.85, min(2.25, min(prs.slide_width / standard_width, prs.slide_height / standard_height)))


def _set_font(run: Any, font_name: str, size: float, color: str, bold: bool = False, scale: float = 1.0) -> None:
    run.font.name = font_name
    run.font.size = Pt(size * scale)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _configure_text_frame(text_frame: Any, margin: float = 0.08, vertical: MSO_ANCHOR = MSO_ANCHOR.TOP) -> None:
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)
    text_frame.vertical_anchor = vertical


def _add_text(slide: Any, prs: Presentation, element: PlannedElement, text: str, font_name: str, size: float, color: str, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT, vertical: MSO_ANCHOR = MSO_ANCHOR.TOP) -> Any:
    shape = slide.shapes.add_textbox(*_coords(prs, element))
    _name(shape, element)
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0.02, vertical=vertical)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    _set_font(run, font_name, size, color, bold, _font_scale(prs))
    return shape


def _shape_box(slide: Any, prs: Presentation, element: PlannedElement, fill: str, line: str | None = None, radius: bool = True) -> Any:
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, *_coords(prs, element))
    _name(shape, element)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line or fill)
    shape.line.width = Pt(1)
    return shape


def _add_rich_box(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, *, fill: str | None = None, heading: str = "", body: str = "", eyebrow: str = "") -> Any:
    fill_color = fill or _mix(tokens.background, tokens.primary, 0.06)
    shape = _shape_box(slide, prs, element, fill_color, _mix(tokens.background, tokens.primary, 0.18))
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0.17)
    if eyebrow:
        p = frame.paragraphs[0]
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = eyebrow
        _set_font(run, tokens.body_font, 11, tokens.primary, True, _font_scale(prs))
    else:
        frame.paragraphs[0]._element.getparent().remove(frame.paragraphs[0]._element)
    p = frame.add_paragraph()
    p.space_after = Pt(7)
    run = p.add_run()
    run.text = heading
    _set_font(run, tokens.heading_font, 18, tokens.text, True, _font_scale(prs))
    if body:
        p = frame.add_paragraph()
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = body
        _set_font(run, tokens.body_font, 12.5, tokens.muted_text, scale=_font_scale(prs))
    return shape


def _render_title(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    _add_text(slide, prs, element, str(element.data.get("text", "")), tokens.heading_font, 26, tokens.text, True)
    underline = PlannedElement(f"{element.element_id}-accent", "decoration", element.x, element.y + element.h - 0.007, 0.075, 0.006, "decoration", {})
    shape = _shape_box(slide, prs, underline, tokens.primary, tokens.primary, radius=False)
    _name(shape, underline)


def _render_message(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    _add_text(slide, prs, element, str(element.data.get("text", "")), tokens.body_font, 11.5, tokens.muted_text)


def _render_cover_panel(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    panel = _shape_box(slide, prs, element, tokens.primary, tokens.primary)
    panel.line.width = Pt(0)
    for index, (x, y, w, h, color) in enumerate(
        [
            (element.x + 0.035, element.y + 0.06, element.w * 0.55, element.h * 0.22, tokens.secondary),
            (element.x + element.w * 0.34, element.y + element.h * 0.39, element.w * 0.48, element.h * 0.18, tokens.accent),
            (element.x + 0.055, element.y + element.h * 0.70, element.w * 0.62, element.h * 0.12, _mix(tokens.primary, tokens.background, 0.3)),
        ]
    ):
        decoration = PlannedElement(f"{element.element_id}-layer-{index + 1}", "decoration", x, y, w, h, "decoration", {})
        _shape_box(slide, prs, decoration, color, color)


def _render_section_panel(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    panel = _shape_box(slide, prs, element, _mix(tokens.background, tokens.primary, 0.08), _mix(tokens.background, tokens.primary, 0.22))
    accent = PlannedElement(f"{element.element_id}-accent", "decoration", element.x, element.y, 0.018, element.h, "decoration", {})
    _shape_box(slide, prs, accent, tokens.primary, tokens.primary, radius=False)


def _render_statement_panel(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    panel = _shape_box(slide, prs, element, _mix(tokens.background, tokens.primary, 0.055), _mix(tokens.background, tokens.primary, 0.24))
    panel.line.width = Pt(1.2 * _font_scale(prs))
    quote = PlannedElement(f"{element.element_id}-quote", "decoration", element.x + 0.035, element.y + 0.045, 0.055, 0.095, "decoration", {})
    _add_text(slide, prs, quote, "“", tokens.heading_font, 34, tokens.primary, True)


def _render_special_text(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    text = str(element.data.get("text", ""))
    role = element.role
    if role == "cover_title":
        _add_text(slide, prs, element, text, tokens.heading_font, 30, tokens.text, True, vertical=MSO_ANCHOR.MIDDLE)
    elif role == "cover_subtitle":
        _add_text(slide, prs, element, text, tokens.body_font, 14, tokens.muted_text, vertical=MSO_ANCHOR.MIDDLE)
    elif role == "section_title":
        _add_text(slide, prs, element, text, tokens.heading_font, 28, tokens.text, True, vertical=MSO_ANCHOR.MIDDLE)
    elif role == "section_message":
        _add_text(slide, prs, element, text, tokens.body_font, 13, tokens.muted_text, vertical=MSO_ANCHOR.MIDDLE)
    elif role == "statement":
        _add_text(slide, prs, element, text, tokens.heading_font, 22, tokens.text, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    elif role == "statement_label":
        _add_text(slide, prs, element, text, tokens.body_font, 11.5, tokens.primary, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    elif role == "closing_title":
        _add_text(slide, prs, element, text, tokens.heading_font, 27, tokens.text, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    elif role == "closing_body":
        _add_text(slide, prs, element, text, tokens.body_font, 13.5, tokens.muted_text, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def _render_body(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    shape = _shape_box(slide, prs, element, _mix(tokens.background, tokens.primary, 0.045), _mix(tokens.background, tokens.primary, 0.16))
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0.25)
    items = list(element.data.get("items", []))
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(item)
        paragraph.level = 0
        paragraph.space_after = Pt(12)
        paragraph.font.name = tokens.body_font
        paragraph.font.size = Pt(17 * _font_scale(prs))
        paragraph.font.color.rgb = _rgb(tokens.text)
        paragraph.font.bold = index == 0


def _render_card(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, index: int) -> None:
    accent = [tokens.primary, tokens.secondary, tokens.accent, _mix(tokens.primary, tokens.secondary, 0.5)][index % 4]
    shape = _add_rich_box(
        slide,
        prs,
        element,
        tokens,
        heading=str(element.data.get("heading", "")),
        body=str(element.data.get("body", "")),
        eyebrow="",
    )
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.line.color.rgb = _rgb(_mix(tokens.background, accent, 0.35))
    marker = PlannedElement(f"{element.element_id}-marker", "decoration", element.x, element.y, 0.009, element.h, "decoration", {})
    _shape_box(slide, prs, marker, accent, accent, radius=False)
    number = PlannedElement(f"{element.element_id}-number", "decoration", element.x + element.w - 0.095, element.y + 0.025, 0.07, 0.065, "decoration", {})
    _add_text(slide, prs, number, str(element.data.get("eyebrow", f"0{index + 1}")), tokens.heading_font, 15, accent, True, PP_ALIGN.RIGHT)


def _render_comparison(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, index: int) -> None:
    fill = _mix(tokens.background, tokens.primary if index == 0 else tokens.secondary, 0.08)
    shape = _shape_box(slide, prs, element, fill, _mix(fill, tokens.text, 0.14))
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0.22)
    p = frame.paragraphs[0]
    p.space_after = Pt(12)
    run = p.add_run()
    run.text = str(element.data.get("heading", ""))
    _set_font(run, tokens.heading_font, 19, tokens.text, True, _font_scale(prs))
    for item in element.data.get("items", []):
        p = frame.add_paragraph()
        p.text = f"• {item}"
        p.space_after = Pt(7)
        p.font.name = tokens.body_font
        p.font.size = Pt(13 * _font_scale(prs))
        p.font.color.rgb = _rgb(tokens.text)


def _render_column_box(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, index: int) -> None:
    palette = [tokens.primary, tokens.secondary, tokens.accent]
    accent = palette[index % len(palette)]
    shape = _add_rich_box(
        slide,
        prs,
        element,
        tokens,
        fill=_mix(tokens.background, accent, 0.065),
        heading=str(element.data.get("heading", "")),
        body=str(element.data.get("body", "")),
    )
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.line.color.rgb = _rgb(_mix(tokens.background, accent, 0.34))
    marker = PlannedElement(f"{element.element_id}-marker", "decoration", element.x, element.y, element.w, 0.012, "decoration", {})
    _shape_box(slide, prs, marker, accent, accent, radius=False)


def _render_metric(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, index: int) -> None:
    colors = [tokens.primary, tokens.secondary, tokens.accent, _mix(tokens.primary, tokens.secondary, 0.5)]
    color = colors[index % len(colors)]
    shape = _shape_box(slide, prs, element, color, color)
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0.15, vertical=MSO_ANCHOR.MIDDLE)
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(element.data.get("value", "—"))
    _set_font(run, tokens.heading_font, 30, _contrast_text(color), True, _font_scale(prs))
    p = frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(element.data.get("label", ""))
    _set_font(run, tokens.body_font, 12, _contrast_text(color), True, _font_scale(prs))


def _render_chart(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    categories = [str(item) for item in element.data.get("categories", [])]
    series = list(element.data.get("series", []))
    chart_type = choose_chart_type(dict(element.data))
    chart_types = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "stacked-bar": XL_CHART_TYPE.BAR_STACKED,
        "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "scatter": XL_CHART_TYPE.XY_SCATTER,
    }
    if chart_type == "scatter":
        data = XyChartData()
        for item in series:
            series_data = data.add_series(str(item.get("name", "Series")))
            x_values = list(item.get("x_values") or [])
            y_values = list(item.get("y_values") or item.get("values") or [])
            for x_value, y_value in zip(x_values, y_values):
                series_data.add_data_point(float(x_value), float(y_value))
    else:
        data = CategoryChartData()
        data.categories = categories
        for item in series:
            data.add_series(str(item.get("name", "Series")), [float(value) for value in item.get("values", [])])
    chart_shape = slide.shapes.add_chart(chart_types[chart_type], *_coords(prs, element), data)
    chart_shape.name = f"SHN_CONTENT:{element.element_id}"
    description = f"Editable {chart_type} chart. " + "; ".join(
        f"{item.get('name', 'Series')}: {', '.join(str(value) for value in (item.get('y_values') or item.get('values') or []))}"
        for item in series
    )
    properties = chart_shape._element.xpath(".//p:cNvPr")
    if properties:
        properties[0].set("descr", description[:2000])
    chart = chart_shape.chart
    if chart_type == "scatter":
        # python-pptx serializes XY_SCATTER with a line-marker scatter style.
        # Hide the connecting line so the rendered result remains a point plot.
        for rendered_series in chart.series:
            rendered_series.format.line.fill.background()
    chart.has_legend = len(series) > 1 or chart_type in {"pie", "doughnut"}
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.has_title = False
    for plot in chart.plots:
        # python-pptx does not expose data-label XML for scatter plots.
        # Keep the chart native/editable and leave labels to PowerPoint defaults.
        if chart_type == "scatter":
            continue
        plot.has_data_labels = True
        plot.data_labels.show_value = True
        plot.data_labels.number_format = "0%" if chart_type in {"pie", "doughnut"} else "#,##0.##"
        if chart_type in {"pie", "doughnut"}:
            plot.data_labels.show_category_name = True
            plot.data_labels.show_percentage = True
            plot.data_labels.position = XL_LABEL_POSITION.BEST_FIT
        elif chart_type in {"line", "area", "scatter"}:
            plot.data_labels.position = XL_LABEL_POSITION.ABOVE
        else:
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.name = tokens.body_font
        plot.data_labels.font.size = Pt(10.5 * _font_scale(prs))
    try:
        chart.value_axis.has_major_gridlines = chart_type not in {"pie", "doughnut"}
        chart.value_axis.tick_labels.font.name = tokens.body_font
        chart.value_axis.tick_labels.font.size = Pt(10 * _font_scale(prs))
        chart.value_axis.tick_labels.number_format = "#,##0.##"
    except (ValueError, AttributeError):
        pass
    try:
        chart.category_axis.tick_labels.font.name = tokens.body_font
        chart.category_axis.tick_labels.font.size = Pt(10 * _font_scale(prs))
    except (ValueError, AttributeError):
        pass
    palette = [tokens.primary, tokens.secondary, tokens.accent]
    for index, series_item in enumerate(chart.series):
        color = palette[index % len(palette)]
        if chart_type in {"pie", "doughnut"}:
            for point_index, point in enumerate(series_item.points):
                point_color = palette[point_index % len(palette)]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(point_color)
                point.format.line.color.rgb = _rgb(point_color)
        else:
            series_item.format.fill.solid()
            series_item.format.fill.fore_color.rgb = _rgb(color)
            series_item.format.line.color.rgb = _rgb(color)
    chart_shape = next(shape for shape in reversed(slide.shapes) if getattr(shape, "has_chart", False))
    _name(chart_shape, element)


def _render_table(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    columns = list(element.data.get("columns", []))
    rows = list(element.data.get("rows", []))
    if not columns:
        return
    table_shape = slide.shapes.add_table(len(rows) + 1, len(columns), *_coords(prs, element))
    _name(table_shape, element)
    table = table_shape.table
    for col_index, heading in enumerate(columns):
        cell = table.cell(0, col_index)
        cell.text = str(heading)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(tokens.primary)
    for row_index, row in enumerate(rows, start=1):
        values = row if isinstance(row, list) else [row.get(column, "") for column in columns]
        for col_index, value in enumerate(values[: len(columns)]):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(tokens.background if row_index % 2 else _mix(tokens.background, tokens.primary, 0.06))
    for row_index in range(len(rows) + 1):
        for col_index in range(len(columns)):
            cell = table.cell(row_index, col_index)
            frame = cell.text_frame
            frame.margin_left = Inches(0.16)
            frame.margin_right = Inches(0.16)
            frame.margin_top = Inches(0.08)
            frame.margin_bottom = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            heading_key = str(columns[col_index]).casefold().replace(" ", "")
            centered_column = any(
                keyword in heading_key
                for keyword in ("種別", "区分", "対応", "可否", "状態", "status", "type", "数値", "件数", "比率")
            )
            alignment = PP_ALIGN.CENTER if row_index == 0 or centered_column else PP_ALIGN.LEFT
            for paragraph in frame.paragraphs:
                paragraph.font.name = tokens.body_font
                paragraph.font.size = Pt(11.5 * _font_scale(prs))
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = _rgb(_contrast_text(tokens.primary) if row_index == 0 else tokens.text)
                paragraph.alignment = alignment


def _render_process_step(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, index: int) -> None:
    fill = _mix(tokens.background, tokens.primary, 0.07 + min(index, 3) * 0.025)
    shape = _shape_box(slide, prs, element, fill, _mix(tokens.background, tokens.primary, 0.28))
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0.13, vertical=MSO_ANCHOR.MIDDLE)
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"{element.data.get('number', index + 1):02d}"
    _set_font(run, tokens.heading_font, 18, tokens.primary, True, _font_scale(prs))
    p = frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(5)
    run = p.add_run()
    run.text = str(element.data.get("heading", ""))
    _set_font(run, tokens.heading_font, 15, tokens.text, True, _font_scale(prs))
    p = frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(element.data.get("body", ""))
    _set_font(run, tokens.body_font, 10.5, tokens.muted_text, scale=_font_scale(prs))


def _render_arrow(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, *_coords(prs, element))
    _name(shape, element)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(tokens.primary)
    shape.line.color.rgb = _rgb(tokens.primary)


def _add_image_cover(slide: Any, prs: Presentation, element: PlannedElement, image_path: Path) -> Any:
    left, top, width, height = _coords(prs, element)
    with Image.open(image_path) as image:
        image_ratio = image.width / max(1, image.height)
    frame_ratio = width / max(1, height)
    picture = slide.shapes.add_picture(str(image_path), left, top, width, height)
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < frame_ratio:
        crop = (1 - image_ratio / frame_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    _name(picture, element)
    return picture


def _render_image(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, asset_paths: dict[str, Path]) -> None:
    asset_id = str(element.data.get("asset_id", ""))
    path_value = asset_paths.get(asset_id)
    if path_value and path_value.is_file():
        _add_image_cover(slide, prs, element, path_value)
        return
    background = _shape_box(slide, prs, element, _mix(tokens.primary, tokens.background, 0.12), tokens.primary)
    background.name = f"SHN_CONTENT:{element.element_id}:UNRESOLVED_ASSET"
    for index, (x, y, w, h, color) in enumerate(
        [
            (element.x + 0.05, element.y + 0.08, element.w * 0.48, element.h * 0.33, tokens.primary),
            (element.x + element.w * 0.42, element.y + element.h * 0.44, element.w * 0.42, element.h * 0.35, tokens.secondary),
            (element.x + element.w * 0.16, element.y + element.h * 0.56, element.w * 0.22, element.h * 0.18, tokens.accent),
        ]
    ):
        decoration = PlannedElement(f"{element.element_id}-fallback-{index}", "decoration", x, y, w, h, "asset_fallback", {})
        _shape_box(slide, prs, decoration, color, color)


def _render_callout(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    heading = str(element.data.get("heading", ""))
    body = str(element.data.get("body", element.data.get("text", "")))
    shape = _add_rich_box(slide, prs, element, tokens, fill=_mix(tokens.background, tokens.accent, 0.07), heading=heading, body=body, eyebrow=str(element.data.get("eyebrow", "KEY POINT")))
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def _render_closing_panel(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    panel = _shape_box(slide, prs, element, _mix(tokens.background, tokens.primary, 0.055), _mix(tokens.background, tokens.primary, 0.24))
    panel.line.width = Pt(1.25 * _font_scale(prs))


def _render_action_list(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any) -> None:
    items = [str(item) for item in element.data.get("items", []) if str(item).strip()]
    text = "   •   ".join(items)
    _add_text(slide, prs, element, text, tokens.body_font, 12, tokens.primary, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def _copy_text_chrome(slide: Any, source_shape: Any, prs: Presentation, chrome: ChromeElement, tokens: Any) -> None:
    text = source_shape.text.strip()
    if not text or re.fullmatch(r"[\d\s./-]+", text):
        return
    element = PlannedElement(f"chrome-{chrome.source_slide_index}-{chrome.source_shape_index}", "chrome", chrome.x, chrome.y, chrome.w, chrome.h, "chrome", {})
    shape = slide.shapes.add_textbox(*_coords(prs, element))
    shape.name = f"SHN_CHROME:{element.element_id}"
    frame = shape.text_frame
    _configure_text_frame(frame, margin=0)
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    source_run = None
    for paragraph in source_shape.text_frame.paragraphs:
        if paragraph.runs:
            source_run = paragraph.runs[0]
            break
    name = source_run.font.name if source_run and source_run.font.name else tokens.body_font
    size = source_run.font.size.pt if source_run and source_run.font.size else 9
    color = tokens.muted_text
    if source_run and source_run.font.color.type is not None:
        try:
            color = str(source_run.font.color.rgb)
        except (AttributeError, TypeError):
            pass
    _set_font(run, name, size, color, bool(source_run.font.bold) if source_run else False)


def _copy_chrome(slide: Any, prs: Presentation, source_prs: Presentation, chrome_items: list[ChromeElement], tokens: Any) -> None:
    for chrome in chrome_items:
        if chrome.source_slide_index >= len(source_prs.slides):
            continue
        source_slide = source_prs.slides[chrome.source_slide_index]
        if chrome.source_shape_index >= len(source_slide.shapes):
            continue
        source_shape = source_slide.shapes[chrome.source_shape_index]
        if source_shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = slide.shapes.add_picture(BytesIO(source_shape.image.blob), round(prs.slide_width * chrome.x), round(prs.slide_height * chrome.y), round(prs.slide_width * chrome.w), round(prs.slide_height * chrome.h))
            picture.crop_left = source_shape.crop_left
            picture.crop_right = source_shape.crop_right
            picture.crop_top = source_shape.crop_top
            picture.crop_bottom = source_shape.crop_bottom
            picture.name = f"SHN_CHROME:picture-{chrome.source_slide_index}-{chrome.source_shape_index}"
        elif getattr(source_shape, "has_text_frame", False):
            _copy_text_chrome(slide, source_shape, prs, chrome, tokens)
        elif source_shape.shape_type != MSO_SHAPE_TYPE.GROUP:
            if source_shape.element.xpath(".//a:blip"):
                continue
            cloned = deepcopy(source_shape.element)
            slide.shapes._spTree.insert_element_before(cloned, "p:extLst")
            for node in cloned.xpath(".//p:cNvPr"):
                node.set("name", f"SHN_CHROME:{node.get('name', 'shape')}")


def _clear_slide_placeholders(slide: Any) -> None:
    for shape in list(slide.shapes):
        if getattr(shape, "is_placeholder", False):
            shape.element.getparent().remove(shape.element)


def _remove_slide(prs: Presentation, target: Any) -> None:
    for index, slide in enumerate(prs.slides):
        if slide.slide_id != target.slide_id:
            continue
        slide_id = prs.slides._sldIdLst[index]
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        del prs.slides._sldIdLst[index]
        return


def _remove_shape(shape: Any) -> None:
    parent = shape.element.getparent()
    if parent is not None:
        parent.remove(shape.element)


def _remove_drop_shadows(slide: Any) -> None:
    for effect_name in ("outerShdw", "prstShdw"):
        for effect in list(slide._element.xpath(f".//a:{effect_name}")):
            parent = effect.getparent()
            if parent is not None:
                parent.remove(effect)


def _render_element(slide: Any, prs: Presentation, element: PlannedElement, tokens: Any, asset_paths: dict[str, Path], counters: dict[str, int]) -> None:
    kind = element.kind
    if kind == "cover_panel":
        _render_cover_panel(slide, prs, element, tokens)
    elif kind == "section_panel":
        _render_section_panel(slide, prs, element, tokens)
    elif kind == "statement_panel":
        _render_statement_panel(slide, prs, element, tokens)
    elif kind == "closing_panel":
        _render_closing_panel(slide, prs, element, tokens)
    elif kind == "text" and element.role in {"cover_title", "cover_subtitle", "section_title", "section_message", "statement", "statement_label", "closing_title", "closing_body"}:
        _render_special_text(slide, prs, element, tokens)
    elif kind == "text" and element.role == "title":
        _render_title(slide, prs, element, tokens)
    elif kind == "text" and element.role == "message":
        _render_message(slide, prs, element, tokens)
    elif kind == "body":
        _render_body(slide, prs, element, tokens)
    elif kind == "card":
        index = counters.get(kind, 0)
        _render_card(slide, prs, element, tokens, index)
        counters[kind] = index + 1
    elif kind == "comparison_column":
        index = counters.get(kind, 0)
        _render_comparison(slide, prs, element, tokens, index)
        counters[kind] = index + 1
    elif kind == "column_box":
        index = counters.get(kind, 0)
        _render_column_box(slide, prs, element, tokens, index)
        counters[kind] = index + 1
    elif kind == "metric":
        index = counters.get(kind, 0)
        _render_metric(slide, prs, element, tokens, index)
        counters[kind] = index + 1
    elif kind in {"bar_chart", "native_chart"}:
        _render_chart(slide, prs, element, tokens)
    elif kind == "table":
        _render_table(slide, prs, element, tokens)
    elif kind == "process_step":
        index = counters.get(kind, 0)
        _render_process_step(slide, prs, element, tokens, index)
        counters[kind] = index + 1
    elif kind == "process_arrow":
        _render_arrow(slide, prs, element, tokens)
    elif kind == "image":
        _render_image(slide, prs, element, tokens, asset_paths)
    elif kind == "callout":
        _render_callout(slide, prs, element, tokens)
    elif kind == "action_list":
        _render_action_list(slide, prs, element, tokens)


def _render_elements(slide: Any, prs: Presentation, elements: list[PlannedElement], tokens: Any, asset_paths: dict[str, Path]) -> None:
    """Render connectors first so later nodes and labels stay above them."""

    counters: dict[str, int] = {}
    ordered = [element for element in elements if element.kind == "process_arrow"] + [
        element for element in elements if element.kind != "process_arrow"
    ]
    for element in ordered:
        _render_element(slide, prs, element, tokens, asset_paths, counters)


def compose_deck(
    template_path: str | Path,
    deck: DeckSpec,
    plans: list[LayoutPlan],
    template_system: TemplateSystem,
    output_path: str | Path,
    asset_paths: dict[str, Path] | None = None,
    composition_mode: str = "reconstruct",
    source_slide_report: dict[str, Any] | None = None,
    source_evidence: list[SourceSlideEvidence] | None = None,
    allow_shadows: bool = False,
    decorativeness: str = "medium",
    source_notes: dict[str, str] | None = None,
) -> Path:
    if decorativeness not in {"low", "medium", "high"}:
        raise ValueError(f"unsupported decorativeness: {decorativeness}")
    source_prs = Presentation(str(template_path))
    prs = Presentation(str(template_path))
    original_slides = list(prs.slides)
    asset_paths = asset_paths or {}
    best_layout = max(template_system.layouts, key=lambda item: item.score_as_canvas) if template_system.layouts else None
    layout_index = best_layout.index if best_layout and best_layout.index < len(prs.slide_layouts) else 0
    slide_layout = prs.slide_layouts[layout_index]
    tokens = template_system.tokens
    source_notes = source_notes or {}
    if composition_mode == "source-slides":
        if source_slide_report is None or source_evidence is None:
            raise ValueError("source-slides mode requires a source slide report and evidence")
        selections = {item["slide_id"]: item for item in source_slide_report.get("selections", [])}
        evidence_by_index = {item.slide_index: item for item in source_evidence}
        for sequence_number, plan in enumerate(plans, start=1):
            selection = selections.get(plan.slide_id)
            if selection is None:
                raise ValueError(f"source slide selection missing for {plan.slide_id}")
            source_index = int(selection["source_slide_index"])
            if source_index >= len(original_slides) or source_index not in evidence_by_index:
                raise ValueError(f"source slide index {source_index} is invalid for {plan.slide_id}")
            slide = clone_slide(prs, original_slides[source_index])
            native_elements = populate_source_slide(
                slide,
                prs,
                evidence_by_index[source_index],
                plan,
                sequence_number,
                asset_paths,
                decorativeness,
            )
            _render_elements(slide, prs, native_elements, tokens, asset_paths)
            if not allow_shadows:
                _remove_drop_shadows(slide)
            if plan.slide_id in source_notes:
                apply_source_notes(slide, source_notes[plan.slide_id])
            slide._element.cSld.set("name", plan.slide_id)
    elif composition_mode == "reconstruct":
        for plan in plans:
            slide = prs.slides.add_slide(slide_layout)
            _clear_slide_placeholders(slide)
            _copy_chrome(slide, prs, source_prs, template_system.chrome, tokens)
            _render_elements(slide, prs, plan.elements, tokens, asset_paths)
            if decorativeness == "low":
                decorative_tokens = ("-accent", "-marker", "-number", "-layer-", "-quote")
                for shape in list(slide.shapes):
                    name = str(shape.name)
                    if name.startswith("SHN_CONTENT:") and any(token in name for token in decorative_tokens):
                        _remove_shape(shape)
            if not allow_shadows:
                _remove_drop_shadows(slide)
            if plan.slide_id in source_notes:
                apply_source_notes(slide, source_notes[plan.slide_id])
            slide._element.cSld.set("name", plan.slide_id)
    else:
        raise ValueError(f"unsupported composition mode: {composition_mode}")
    for slide in original_slides:
        _remove_slide(prs, slide)
    prs.core_properties.title = deck.title
    prs.core_properties.subject = "Generated by SlideHarness Next greenfield MVP"
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    return target


def extract_template_preview_image(template_path: str | Path, output_dir: str | Path) -> Path | None:
    prs = Presentation(str(template_path))
    candidates: list[tuple[int, Any]] = []
    slide_area = prs.slide_width * prs.slide_height
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            area = shape.width * shape.height
            ratio = area / slide_area
            if 0.08 <= ratio <= 0.75:
                candidates.append((area, shape))
    if not candidates:
        return None
    _, shape = max(candidates, key=lambda item: item[0])
    suffix = f".{shape.image.ext or 'png'}"
    target = Path(output_dir) / f"template-preview{suffix}"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(shape.image.blob)
    return target
