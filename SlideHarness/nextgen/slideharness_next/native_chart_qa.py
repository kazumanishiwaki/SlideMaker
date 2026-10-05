from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation

from .charts import choose_chart_type, validate_chart_content


def inspect_native_charts(pptx_path: str | Path, plans: list[Any]) -> dict[str, Any]:
    prs = Presentation(str(pptx_path))
    reports: list[dict[str, Any]] = []
    findings: list[dict[str, Any]] = []
    for index, plan in enumerate(plans):
        if plan.recipe not in {"native_bar_chart", "native_chart"}:
            continue
        slide = prs.slides[index]
        chart_shapes = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
        element = next((item for item in plan.elements if item.kind in {"native_chart", "bar_chart"}), None)
        chart_data = dict(element.data) if element is not None else {}
        expected_type = choose_chart_type(chart_data)
        errors = validate_chart_content(chart_data)
        if len(chart_shapes) != 1:
            findings.append({"category": "Technical", "code": "native_chart_count", "severity": "blocking", "slide_id": plan.slide_id, "message": f"expected one editable chart, found {len(chart_shapes)}"})
        if errors:
            findings.append({"category": "Content", "code": "native_chart_contract", "severity": "blocking", "slide_id": plan.slide_id, "message": "; ".join(errors)})
        actual = chart_shapes[0].chart if chart_shapes else None
        reports.append({
            "slide_id": plan.slide_id,
            "expected_type": expected_type,
            "powerpoint_chart_type": str(actual.chart_type) if actual is not None else None,
            "series_count": len(actual.series) if actual is not None else 0,
            "category_count": len(chart_data.get("categories") or []),
            "editable": actual is not None,
        })
    if not reports:
        return {"schema_version": "0.4.0", "status": "SKIP", "reason": "no native charts", "charts": [], "findings": []}
    return {
        "schema_version": "0.4.0",
        "status": "FAIL" if findings else "PASS",
        "charts": reports,
        "findings": findings,
        "summary": {"chart_count": len(reports), "blocking_count": len(findings)},
    }

