from __future__ import annotations

from pathlib import Path
from typing import Any
import re

from pptx import Presentation

from .models import DeckSpec, SourceBundle


URL_PATTERN = re.compile(r"https?://[^\s<>()\[\]{}\"']+", re.IGNORECASE)


def _urls(text: str) -> list[str]:
    return list(
        dict.fromkeys(
            match.group(0).rstrip(".,;:!?。．、）」』】")
            for match in URL_PATTERN.finditer(text)
        )
    )


def _asset_citations(content: Any) -> list[dict[str, str]]:
    citations: list[dict[str, str]] = []

    def visit(value: Any) -> None:
        if isinstance(value, dict):
            asset_id = str(value.get("asset_id") or "").strip()
            raw_urls = value.get("source_urls")
            urls = [str(item).strip() for item in raw_urls] if isinstance(raw_urls, list) else []
            if value.get("source_url"):
                urls.insert(0, str(value["source_url"]).strip())
            note = str(value.get("source_note") or value.get("source_citation") or "").strip()
            for url in dict.fromkeys(url for url in urls if url):
                citations.append({"asset_id": asset_id or "external-asset", "url": url, "note": note})
            for nested in value.values():
                visit(nested)
        elif isinstance(value, list):
            for nested in value:
                visit(nested)

    visit(content)
    unique: dict[tuple[str, str], dict[str, str]] = {}
    for item in citations:
        unique[(item["asset_id"], item["url"])] = item
    return list(unique.values())


def build_source_notes(deck: DeckSpec, bundle: SourceBundle) -> tuple[dict[str, str], dict[str, Any]]:
    """Build compact, path-free source ledgers for PowerPoint speaker notes."""

    files = {item.source_id: item for item in bundle.files}
    units = {item.unit_id: item for item in bundle.units}
    notes: dict[str, str] = {}
    slides: list[dict[str, Any]] = []
    for slide in deck.slides:
        lines = ["[Sources]"]
        records: list[dict[str, Any]] = []
        for unit_id in slide.source_unit_ids:
            unit = units.get(unit_id)
            if unit is None:
                lines.append(f"- {unit_id} | unresolved source unit")
                records.append({"unit_id": unit_id, "status": "unresolved"})
                continue
            source = files.get(unit.source_id)
            filename = source.filename if source else unit.source_id
            section = " > ".join(unit.section_path) if unit.section_path else "(document root)"
            source_urls = _urls(unit.source_text)
            suffix = f" | {' | '.join(source_urls)}" if source_urls else ""
            lines.append(f"- {unit.unit_id} | {filename} | {section}{suffix}")
            records.append(
                {
                    "unit_id": unit.unit_id,
                    "source_id": unit.source_id,
                    "filename": filename,
                    "section_path": unit.section_path,
                    "urls": source_urls,
                    "status": "resolved",
                }
            )
        asset_records = _asset_citations(slide.content)
        for item in asset_records:
            note = f" | {item['note']}" if item["note"] else ""
            lines.append(f"- Asset {item['asset_id']} | {item['url']}{note}")
        notes[slide.slide_id] = "\n".join(lines)
        slides.append({"slide_id": slide.slide_id, "source_unit_ids": slide.source_unit_ids, "records": records, "asset_records": asset_records})
    return notes, {"schema_version": "0.11-mvp.4", "slides": slides}


def apply_source_notes(slide: Any, text: str) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.text = text


def _notes_text(slide: Any) -> str:
    try:
        return slide.notes_slide.notes_text_frame.text or ""
    except (AttributeError, ValueError):
        return ""


def inspect_provenance(
    pptx_path: str | Path,
    deck: DeckSpec,
    bundle: SourceBundle | None,
) -> dict[str, Any]:
    if bundle is None:
        return {
            "schema_version": "0.11-mvp.4",
            "status": "SKIP",
            "reason": "source bundle is unavailable",
            "summary": {"slide_count": len(deck.slides), "blocking_count": 0, "finding_count": 0},
            "slides": [],
            "findings": [],
        }

    prs = Presentation(str(pptx_path))
    valid_ids = {item.unit_id for item in bundle.units}
    units = {item.unit_id: item for item in bundle.units}
    findings: list[dict[str, Any]] = []
    slides: list[dict[str, Any]] = []
    for index, slide_spec in enumerate(deck.slides):
        slide = prs.slides[index] if index < len(prs.slides) else None
        notes = _notes_text(slide) if slide is not None else ""
        missing_ids = [unit_id for unit_id in slide_spec.source_unit_ids if unit_id not in notes]
        unknown_ids = [unit_id for unit_id in slide_spec.source_unit_ids if unit_id not in valid_ids]
        expected_urls = list(
            dict.fromkeys(
                url
                for unit_id in slide_spec.source_unit_ids
                if unit_id in units
                for url in _urls(units[unit_id].source_text)
            )
        )
        asset_records = _asset_citations(slide_spec.content)
        expected_asset_ids = [item["asset_id"] for item in asset_records]
        expected_urls.extend(item["url"] for item in asset_records if item["url"] not in expected_urls)
        missing_urls = [url for url in expected_urls if url not in notes]
        missing_asset_ids = [asset_id for asset_id in expected_asset_ids if asset_id not in notes]
        if slide is None:
            findings.append(
                {
                    "category": "Technical",
                    "code": "provenance_slide_missing",
                    "severity": "blocking",
                    "slide_id": slide_spec.slide_id,
                    "message": "the expected slide is missing from the PPTX",
                }
            )
        elif not notes.startswith("[Sources]"):
            findings.append(
                {
                    "category": "Content",
                    "code": "sources_block_missing",
                    "severity": "blocking",
                    "slide_id": slide_spec.slide_id,
                    "message": "speaker notes do not start with a [Sources] block",
                }
            )
        if missing_ids:
            findings.append(
                {
                    "category": "Content",
                    "code": "source_note_unit_missing",
                    "severity": "blocking",
                    "slide_id": slide_spec.slide_id,
                    "message": f"speaker notes omit {len(missing_ids)} referenced source unit(s)",
                    "items": missing_ids,
                }
            )
        if unknown_ids:
            findings.append(
                {
                    "category": "Content",
                    "code": "source_note_unit_unknown",
                    "severity": "blocking",
                    "slide_id": slide_spec.slide_id,
                    "message": f"speaker notes reference {len(unknown_ids)} unknown source unit(s)",
                    "items": unknown_ids,
                }
            )
        if missing_urls:
            findings.append(
                {
                    "category": "Content",
                    "code": "source_note_url_missing",
                    "severity": "blocking",
                    "slide_id": slide_spec.slide_id,
                    "message": f"speaker notes omit {len(missing_urls)} source URL(s)",
                    "items": missing_urls,
                }
            )
        if missing_asset_ids:
            findings.append(
                {
                    "category": "Content",
                    "code": "source_note_asset_missing",
                    "severity": "blocking",
                    "slide_id": slide_spec.slide_id,
                    "message": f"speaker notes omit {len(missing_asset_ids)} externally sourced asset ID(s)",
                    "items": missing_asset_ids,
                }
            )
        slides.append(
            {
                "slide_id": slide_spec.slide_id,
                "has_sources_block": notes.startswith("[Sources]"),
                "expected_source_unit_count": len(slide_spec.source_unit_ids),
                "missing_source_unit_ids": missing_ids,
                "unknown_source_unit_ids": unknown_ids,
                "expected_source_urls": expected_urls,
                "missing_source_urls": missing_urls,
                "expected_external_asset_ids": expected_asset_ids,
                "missing_external_asset_ids": missing_asset_ids,
            }
        )
    blocking = [item for item in findings if item["severity"] == "blocking"]
    return {
        "schema_version": "0.11-mvp.4",
        "status": "PASS" if not blocking else "FAIL",
        "summary": {"slide_count": len(deck.slides), "blocking_count": len(blocking), "finding_count": len(findings)},
        "slides": slides,
        "findings": findings,
    }
