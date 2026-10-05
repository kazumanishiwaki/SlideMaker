from __future__ import annotations

from collections import Counter
from copy import deepcopy
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any
import math
import re

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

from .models import LayoutPlan, PlannedElement


REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass(frozen=True)
class TextSlotEvidence:
    shape_index: int
    x: float
    y: float
    w: float
    h: float
    max_font_pt: float
    text: str
    normalized_text: str


@dataclass(frozen=True)
class SourceSlideEvidence:
    slide_index: int
    title: str
    shape_count: int
    picture_count: int
    photo_like_picture_indices: list[int]
    table_shape_indices: list[int]
    chart_shape_indices: list[int]
    protected_text_indices: list[int]
    title_shape_index: int | None
    message_shape_index: int | None
    text_slots: list[TextSlotEvidence] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _normalize(text: str) -> str:
    return " ".join(text.split()).casefold()


def _max_font(shape: Any) -> float:
    sizes = [run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size]
    return max(sizes) if sizes else 0.0


def _picture_is_photo_like(shape: Any, area_ratio: float) -> bool:
    if area_ratio < 0.025:
        return False
    try:
        with Image.open(BytesIO(shape.image.blob)).convert("RGB") as image:
            preview = image.copy()
            preview.thumbnail((64, 64))
            colors = preview.getcolors(maxcolors=64 * 64) or []
            dominant_ratio = max((count for count, _ in colors), default=0) / max(1, preview.width * preview.height)
            histogram = preview.convert("L").histogram()
            total = sum(histogram)
            entropy = -sum((count / total) * math.log2(count / total) for count in histogram if count)
            stddev = sum(ImageStat.Stat(preview).stddev) / 3
            # Canva-style decks often embed vector ornaments as raster pictures.
            # Entropy and dominant-color checks avoid discarding those geometric assets.
            return bool(entropy >= 4.5 and dominant_ratio <= 0.45 and stddev >= 32)
    except Exception:
        return False


def _repeated_texts(prs: Presentation) -> set[str]:
    counts: Counter[str] = Counter()
    for slide in prs.slides:
        seen: set[str] = set()
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            value = _normalize(shape.text)
            if value and value not in seen:
                counts[value] += 1
                seen.add(value)
    threshold = max(2, math.ceil(len(prs.slides) * 0.5))
    return {value for value, count in counts.items() if count >= threshold}


def analyze_source_slides(template_path: str | Path) -> list[SourceSlideEvidence]:
    prs = Presentation(str(template_path))
    repeated = _repeated_texts(prs)
    evidence: list[SourceSlideEvidence] = []
    slide_area = prs.slide_width * prs.slide_height
    for slide_index, slide in enumerate(prs.slides):
        text_slots: list[TextSlotEvidence] = []
        pictures: list[int] = []
        photo_like: list[int] = []
        tables: list[int] = []
        charts: list[int] = []
        protected: list[int] = []
        for shape_index, shape in enumerate(slide.shapes):
            area_ratio = (shape.width * shape.height) / slide_area
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures.append(shape_index)
                if _picture_is_photo_like(shape, area_ratio):
                    photo_like.append(shape_index)
            if getattr(shape, "has_table", False):
                tables.append(shape_index)
            if getattr(shape, "has_chart", False):
                charts.append(shape_index)
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            value = _normalize(shape.text)
            slot = TextSlotEvidence(
                shape_index,
                round(shape.left / prs.slide_width, 4),
                round(shape.top / prs.slide_height, 4),
                round(shape.width / prs.slide_width, 4),
                round(shape.height / prs.slide_height, 4),
                round(_max_font(shape), 2),
                shape.text.strip(),
                value,
            )
            text_slots.append(slot)
            area = slot.w * slot.h
            if value in repeated and (area <= 0.035 or slot.y <= 0.16 or slot.y + slot.h >= 0.9):
                protected.append(shape_index)
        candidates = [slot for slot in text_slots if slot.shape_index not in protected and not re.fullmatch(r"[\d\s./-]+", slot.text)]
        if slide_index not in {0, len(prs.slides) - 1}:
            top_candidates = [slot for slot in candidates if slot.y <= 0.17 and slot.x >= 0.08]
        else:
            top_candidates = candidates
        semantic_candidates = top_candidates or candidates
        title_slot = max(semantic_candidates, key=lambda slot: (slot.max_font_pt, slot.w), default=None)
        message_candidates = [slot for slot in semantic_candidates if title_slot is None or slot.shape_index != title_slot.shape_index]
        message_slot = max(message_candidates, key=lambda slot: (slot.w, -slot.y), default=None)
        title_text = title_slot.text if title_slot else ""
        evidence.append(
            SourceSlideEvidence(
                slide_index=slide_index,
                title=title_text,
                shape_count=len(slide.shapes),
                picture_count=len(pictures),
                photo_like_picture_indices=photo_like,
                table_shape_indices=tables,
                chart_shape_indices=charts,
                protected_text_indices=protected,
                title_shape_index=title_slot.shape_index if title_slot else None,
                message_shape_index=message_slot.shape_index if message_slot else None,
                text_slots=text_slots,
            )
        )
    return evidence


def _content_anchors(item: SourceSlideEvidence) -> list[TextSlotEvidence]:
    header = {item.title_shape_index, item.message_shape_index, *item.protected_text_indices}
    output = []
    for slot in item.text_slots:
        if slot.shape_index in header or slot.y < 0.18:
            continue
        value = slot.text.strip()
        if re.fullmatch(r"(?:0?[1-9]|[1-9]\d|20XX)", value, flags=re.IGNORECASE):
            output.append(slot)
    return output


def _cluster_count(values: list[float], tolerance: float = 0.09) -> int:
    clusters: list[float] = []
    for value in sorted(values):
        if not clusters or abs(value - clusters[-1]) > tolerance:
            clusters.append(value)
        else:
            clusters[-1] = (clusters[-1] + value) / 2
    return len(clusters)


def _recipe_score(item: SourceSlideEvidence, recipe: str, total_slides: int) -> tuple[float, list[str]]:
    title = _normalize(item.title)
    anchors = _content_anchors(item)
    x_clusters = _cluster_count([slot.x + slot.w / 2 for slot in anchors])
    y_clusters = _cluster_count([slot.y + slot.h / 2 for slot in anchors])
    score = 0.0
    reasons: list[str] = []

    def add(value: float, reason: str) -> None:
        nonlocal score
        score += value
        reasons.append(reason)

    if recipe == "cover":
        if item.slide_index == 0:
            add(120, "first_slide")
        if item.title_shape_index is not None:
            add(10, "large_title_available")
    elif recipe in {"section", "statement"}:
        if item.slide_index == 1:
            add(100, "section_like_second_slide")
        if len(item.text_slots) <= 4:
            add(25, "few_large_text_slots")
        if any(keyword in title for keyword in ("qna", "section", "タイトル")):
            add(15, "section_title_evidence")
    elif recipe == "closing_cta":
        if item.slide_index == total_slides - 1:
            add(120, "last_slide")
        if any(keyword in title for keyword in ("thank", "closing", "qna")):
            add(25, "closing_title_evidence")
    elif recipe == "table":
        if item.table_shape_indices:
            add(120, "native_table_frame")
        if any(keyword in title for keyword in ("概要", "table", "一覧")):
            add(20, "table_title_evidence")
    elif recipe == "three_column":
        if x_clusters == 3:
            add(85, "three_horizontal_anchor_groups")
        if any(keyword in title for keyword in ("3段", "3項目", "three")):
            add(50, "three_column_title_evidence")
    elif recipe == "card_grid_2x2":
        if len(anchors) == 4 and x_clusters == 2 and y_clusters == 2:
            add(110, "two_by_two_anchor_grid")
        elif len(anchors) == 4:
            add(25, "four_anchor_groups")
        if any(keyword in title for keyword in ("4段", "2x2", "2×2")):
            add(55, "four_block_title_evidence")
    elif recipe == "process":
        if len(anchors) in {3, 4, 5} and x_clusters == len(anchors):
            add(85, "horizontal_sequence_anchors")
        if any(keyword in title for keyword in ("時間", "経過", "flow", "process", "timeline")):
            add(65, "sequence_title_evidence")
    elif recipe == "kpi_grid":
        if len(anchors) >= 4:
            add(35, "multiple_numeric_slots")
        if any(keyword in title for keyword in ("ハイライト", "kpi", "業績")):
            add(80, "kpi_title_evidence")
    elif recipe in {"native_bar_chart", "native_chart"}:
        if any(keyword in title for keyword in ("活動", "グラフ", "chart", "売上")):
            add(70, "chart_title_evidence")
        if item.picture_count >= 4:
            add(10, "visual_data_components")
    elif recipe in {"two_column_equal", "comparison"}:
        if x_clusters == 2:
            add(65, "two_horizontal_groups")
        if any(keyword in title for keyword in ("比較", "ベン", "2項目", "comparison")):
            add(60, "comparison_title_evidence")
    elif recipe in {"title_body", "title_body_callout", "image_text", "generated_visual"}:
        if any(keyword in title for keyword in ("長文", "text", "内容")):
            add(60, "body_text_title_evidence")
        if item.photo_like_picture_indices and recipe in {"image_text", "generated_visual"}:
            add(55, "replaceable_photo_frame")
        if len(item.text_slots) >= 4:
            add(15, "body_text_slots")
    else:
        add(5, "generic_source_slide")
    score -= abs(item.shape_count - 24) * 0.05
    return round(score, 3), reasons


def select_source_slides(evidence: list[SourceSlideEvidence], plans: list[LayoutPlan]) -> dict[str, Any]:
    selections: list[dict[str, Any]] = []
    used: Counter[int] = Counter()
    total = len(evidence)
    for plan in plans:
        candidates = []
        for item in evidence:
            score, reasons = _recipe_score(item, plan.recipe, total)
            reuse_penalty = min(18, used[item.slide_index] * 6)
            candidates.append((score - reuse_penalty, score, item, reasons, reuse_penalty))
        effective, raw_score, selected, reasons, penalty = max(candidates, key=lambda row: row[0])
        used[selected.slide_index] += 1
        selections.append(
            {
                "slide_id": plan.slide_id,
                "recipe": plan.recipe,
                "source_slide_index": selected.slide_index,
                "source_slide_number": selected.slide_index + 1,
                "source_title": selected.title,
                "raw_score": raw_score,
                "effective_score": round(effective, 3),
                "reuse_penalty": penalty,
                "reasons": reasons,
            }
        )
    return {
        "schema_version": "0.11-mvp.4",
        "mode": "source_slide_adaptive",
        "selections": selections,
        "source_evidence": [item.to_dict() for item in evidence],
    }


def clone_slide(prs: Presentation, source_slide: Any) -> Any:
    destination = prs.slides.add_slide(source_slide.slide_layout)
    for child in list(destination._element):
        destination._element.remove(child)
    for child in source_slide._element:
        destination._element.append(deepcopy(child))
    destination._element.attrib.clear()
    destination._element.attrib.update(source_slide._element.attrib)
    # python-pptx caches Slide.shapes while add_slide() clones placeholders.
    # The OOXML tree above is now authoritative, so force a fresh SlideShapes view.
    destination.__dict__.pop("shapes", None)
    relationship_map: dict[str, str] = {}
    for rel_id, relationship in source_slide.part.rels.items():
        kind = relationship.reltype.rsplit("/", 1)[-1]
        if kind == "notesSlide":
            continue
        if kind == "slideLayout":
            for destination_id, destination_rel in destination.part.rels.items():
                if destination_rel.reltype == relationship.reltype:
                    relationship_map[rel_id] = destination_id
                    break
            continue
        target = relationship.target_ref if relationship.is_external else relationship.target_part
        relationship_map[rel_id] = destination.part.relate_to(target, relationship.reltype, relationship.is_external)
    for node in destination._element.iter():
        for attribute, value in list(node.attrib.items()):
            if attribute.startswith(f"{{{REL_NS}}}") and value in relationship_map:
                node.set(attribute, relationship_map[value])
    return destination


def _set_shape_text(shape: Any, text: str, name: str) -> None:
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    paragraphs = list(frame.paragraphs)
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.text = text
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def _remove_shape(shape: Any) -> None:
    parent = shape.element.getparent()
    if parent is not None:
        parent.remove(shape.element)


def _group_slots(evidence: SourceSlideEvidence, recipe: str) -> list[list[TextSlotEvidence]]:
    anchors = _content_anchors(evidence)
    header = {evidence.title_shape_index, evidence.message_shape_index, *evidence.protected_text_indices}
    content_slots = [slot for slot in evidence.text_slots if slot.shape_index not in header and slot.y >= 0.18]
    if not anchors:
        return []
    if recipe == "card_grid_2x2":
        anchors = sorted(anchors, key=lambda slot: (slot.y, slot.x))
    else:
        anchors = sorted(anchors, key=lambda slot: (slot.x, slot.y))
    groups: list[list[TextSlotEvidence]] = [[anchor] for anchor in anchors]
    for slot in content_slots:
        if any(slot.shape_index == anchor.shape_index for anchor in anchors):
            continue
        center_x, center_y = slot.x + slot.w / 2, slot.y + slot.h / 2
        if recipe in {"three_column", "process"}:
            distances = [abs(center_x - (anchor.x + anchor.w / 2)) for anchor in anchors]
        else:
            distances = [abs(center_x - (anchor.x + anchor.w / 2)) + abs(center_y - (anchor.y + anchor.h / 2)) * 1.2 for anchor in anchors]
        groups[distances.index(min(distances))].append(slot)
    return groups


def _semantic_items(plan: LayoutPlan) -> list[dict[str, str]]:
    def deduplicate(heading: str, body: str) -> str:
        semantic_heading = re.sub(r"[\s。．.!！?？、,]", "", heading).casefold()
        semantic_body = re.sub(r"[\s。．.!！?？、,]", "", body).casefold()
        return "" if semantic_heading and semantic_heading == semantic_body else body

    output: list[dict[str, str]] = []
    for element in plan.elements:
        if element.kind in {"card", "column_box", "comparison_column"}:
            heading = str(element.data.get("heading", ""))
            body = str(element.data.get("body", "\n".join(element.data.get("items", []))))
            output.append({"number": str(len(output) + 1).zfill(2), "heading": heading, "body": deduplicate(heading, body)})
        elif element.kind == "process_step":
            heading = str(element.data.get("heading", ""))
            body = str(element.data.get("body", ""))
            output.append({"number": str(element.data.get("number", len(output) + 1)).zfill(2), "heading": heading, "body": deduplicate(heading, body)})
        elif element.kind == "metric":
            output.append({"number": str(len(output) + 1).zfill(2), "heading": str(element.data.get("label", "")), "body": str(element.data.get("value", ""))})
    return output


def _fill_groups(slide: Any, evidence: SourceSlideEvidence, plan: LayoutPlan, assigned: set[int]) -> bool:
    groups = _group_slots(evidence, plan.recipe)
    items = _semantic_items(plan)
    if len(groups) < len(items):
        return False
    shapes = list(slide.shapes)
    component_name = "source-step" if plan.recipe == "process" else "source-item"
    for group_index, item in enumerate(items):
        group = groups[group_index]
        anchor = group[0]
        _set_shape_text(shapes[anchor.shape_index], item["number"], f"SHN_CONTENT:{plan.slide_id}:{component_name}-{group_index + 1}-number")
        assigned.add(anchor.shape_index)
        remaining = sorted(group[1:], key=lambda slot: (slot.y, slot.x))
        if not remaining:
            continue
        heading_slot = max(remaining, key=lambda slot: (slot.max_font_pt, -slot.y))
        _set_shape_text(shapes[heading_slot.shape_index], item["heading"], f"SHN_CONTENT:{plan.slide_id}:{component_name}-{group_index + 1}-heading")
        assigned.add(heading_slot.shape_index)
        body_slots = [slot for slot in remaining if slot.shape_index != heading_slot.shape_index]
        if body_slots:
            body_slot = sorted(body_slots, key=lambda slot: (slot.y, slot.x))[0]
            body_name = f"SHN_CONTENT:{plan.slide_id}:{component_name}-{group_index + 1}-body" if item["body"] else f"SHN_SOURCE_CLEARED:{plan.slide_id}:{body_slot.shape_index}"
            _set_shape_text(shapes[body_slot.shape_index], item["body"], body_name)
            assigned.add(body_slot.shape_index)
    return True


def _replace_picture(slide: Any, shape: Any, path: Path, slide_id: str) -> None:
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    _remove_shape(shape)
    picture = slide.shapes.add_picture(str(path), left, top, width, height)
    with Image.open(path) as image:
        image_ratio = image.width / max(1, image.height)
    frame_ratio = width / max(1, height)
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < frame_ratio:
        crop = (1 - image_ratio / frame_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    picture.name = f"SHN_CONTENT:{slide_id}:source-image"


def _compact_process_after_photo_removal(slide: Any, prs: Presentation) -> None:
    """Close the semantic hole left when a stock-photo timeline becomes text-only."""

    # The source timeline reserves a photo strip above its cards. Removing the
    # photos should narrow that gap, but a large upward jump leaves the bottom
    # third visually empty. A restrained shift keeps the process centered and
    # distributes the recovered whitespace above and below the card row.
    shift = round(prs.slide_height * 0.08)
    for shape in list(slide.shapes):
        top = shape.top / prs.slide_height
        bottom = (shape.top + shape.height) / prs.slide_height
        if top < 0.36 or bottom > 0.94:
            continue
        shape.top = max(0, shape.top - shift)


def _expand_closing_cta(slide: Any, prs: Presentation) -> None:
    """Give long closing actions enough fixed-size room without auto-shrink."""

    labels = [
        shape
        for shape in slide.shapes
        if str(shape.name).endswith((":source-body", ":source-actions"))
        and getattr(shape, "has_text_frame", False)
        and shape.text_frame.text.strip()
    ]
    if not labels:
        return
    labels.sort(key=lambda shape: shape.left)
    start_x = round(prs.slide_width * 0.062)
    gap = round(prs.slide_width * 0.025)
    total_width = round(prs.slide_width * 0.66)
    container_width = (total_width - gap * (len(labels) - 1)) // len(labels)
    container_top = round(prs.slide_height * 0.80)
    container_height = round(prs.slide_height * 0.075)
    inset_x = round(prs.slide_width * 0.009)
    inset_y = round(prs.slide_height * 0.010)

    pictures = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and 0.74 <= shape.top / prs.slide_height <= 0.92
    ]
    used_picture_ids: set[int] = set()
    for index, label in enumerate(labels):
        left = start_x + index * (container_width + gap)
        nearest = min(
            (picture for picture in pictures if id(picture) not in used_picture_ids),
            key=lambda picture: abs((picture.left + picture.width / 2) - (label.left + label.width / 2)),
            default=None,
        )
        if nearest is not None:
            nearest.left = left
            nearest.top = container_top
            nearest.width = container_width
            nearest.height = container_height
            used_picture_ids.add(id(nearest))
        label.left = left + inset_x
        label.top = container_top + inset_y
        label.width = container_width - 2 * inset_x
        label.height = container_height - 2 * inset_y
        label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def populate_source_slide(
    slide: Any,
    prs: Presentation,
    evidence: SourceSlideEvidence,
    plan: LayoutPlan,
    sequence_number: int,
    asset_paths: dict[str, Path],
    decorativeness: str = "medium",
) -> list[PlannedElement]:
    if decorativeness not in {"low", "medium", "high"}:
        raise ValueError(f"unsupported decorativeness: {decorativeness}")
    shapes = list(slide.shapes)
    assigned: set[int] = set()
    protected = set(evidence.protected_text_indices)
    if evidence.title_shape_index is not None and evidence.title_shape_index < len(shapes):
        title_element = next((element for element in plan.elements if element.role in {"title", "cover_title", "section_title", "closing_title", "statement_label"}), None)
        title = str(title_element.data.get("text", "")) if title_element else ""
        _set_shape_text(shapes[evidence.title_shape_index], title, f"SHN_CONTENT:{plan.slide_id}:source-title")
        assigned.add(evidence.title_shape_index)
    if plan.recipe != "closing_cta" and evidence.message_shape_index is not None and evidence.message_shape_index < len(shapes):
        message_element = next((element for element in plan.elements if element.role in {"message", "cover_subtitle", "section_message", "closing_body", "statement"}), None)
        message = str(message_element.data.get("text", "")) if message_element else ""
        _set_shape_text(shapes[evidence.message_shape_index], message, f"SHN_CONTENT:{plan.slide_id}:source-message")
        assigned.add(evidence.message_shape_index)
    group_handled = False
    if plan.recipe in {"three_column", "card_grid_2x2", "process", "kpi_grid", "two_column_equal", "comparison"}:
        group_handled = _fill_groups(slide, evidence, plan, assigned)
    elif plan.recipe == "title_body":
        body_element = next((element for element in plan.elements if element.kind == "body"), None)
        candidates = [slot for slot in evidence.text_slots if slot.y >= 0.18 and slot.shape_index not in protected]
        if body_element and candidates:
            target = max(candidates, key=lambda slot: slot.w * slot.h)
            body = "\n".join(f"• {item}" for item in body_element.data.get("items", []))
            _set_shape_text(shapes[target.shape_index], body, f"SHN_CONTENT:{plan.slide_id}:source-body")
            assigned.add(target.shape_index)
            group_handled = True
    elif plan.recipe == "closing_cta":
        candidates = sorted([slot for slot in evidence.text_slots if slot.y >= 0.18 and slot.shape_index not in protected and slot.shape_index not in assigned], key=lambda slot: (slot.y, slot.x))
        body_element = next((element for element in plan.elements if element.role == "closing_body"), None)
        action_element = next((element for element in plan.elements if element.kind == "action_list"), None)
        if candidates and body_element:
            _set_shape_text(shapes[candidates[0].shape_index], str(body_element.data.get("text", "")), f"SHN_CONTENT:{plan.slide_id}:source-body")
            assigned.add(candidates[0].shape_index)
        if len(candidates) > 1 and action_element:
            _set_shape_text(shapes[candidates[1].shape_index], " / ".join(str(item) for item in action_element.data.get("items", [])), f"SHN_CONTENT:{plan.slide_id}:source-actions")
            assigned.add(candidates[1].shape_index)
        group_handled = True
    elif plan.recipe in {"cover", "section", "statement"}:
        group_handled = True

    native_elements: list[PlannedElement] = []
    if plan.recipe == "table":
        table_element = next((element for element in plan.elements if element.kind == "table"), None)
        if table_element and evidence.table_shape_indices:
            source_table = shapes[evidence.table_shape_indices[0]]
            native_elements.append(
                PlannedElement(
                    table_element.element_id,
                    table_element.kind,
                    table_element.x,
                    table_element.y,
                    table_element.w,
                    table_element.h,
                    table_element.role,
                    table_element.data,
                )
            )
            _remove_shape(source_table)
            group_handled = True
    if plan.recipe in {"image_text", "generated_visual"}:
        image_element = next((element for element in plan.elements if element.kind == "image"), None)
        asset_path = asset_paths.get(str(image_element.data.get("asset_id", ""))) if image_element else None
        photo_shapes = [shapes[index] for index in evidence.photo_like_picture_indices if index < len(shapes)]
        if asset_path and photo_shapes:
            target = max(photo_shapes, key=lambda shape: shape.width * shape.height)
            _replace_picture(slide, target, asset_path, plan.slide_id)
            group_handled = True

    for slot in evidence.text_slots:
        if slot.shape_index in protected or slot.shape_index in assigned or slot.shape_index >= len(shapes):
            continue
        shape = shapes[slot.shape_index]
        if getattr(shape, "has_text_frame", False):
            if slot.y < 0.16 and re.fullmatch(r"\d+", slot.text.strip()):
                _set_shape_text(shape, f"{sequence_number:02d}", f"SHN_CONTENT:{plan.slide_id}:source-page-number")
                assigned.add(slot.shape_index)
            else:
                _set_shape_text(shape, "", f"SHN_SOURCE_CLEARED:{plan.slide_id}:{slot.shape_index}")
    for index in evidence.photo_like_picture_indices:
        if index >= len(shapes):
            continue
        shape = shapes[index]
        if shape.element.getparent() is not None and not str(shape.name).startswith("SHN_CONTENT:"):
            _remove_shape(shape)
    if plan.recipe == "process" and evidence.photo_like_picture_indices:
        _compact_process_after_photo_removal(slide, prs)
    if plan.recipe == "closing_cta":
        _expand_closing_cta(slide, prs)
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    assigned_centers = []
    for index in assigned:
        if index >= len(shapes) or shapes[index].element.getparent() is None:
            continue
        item = shapes[index]
        assigned_centers.append((item.left + item.width / 2, item.top + item.height / 2))
    cleared_centers = []
    for slot in evidence.text_slots:
        if slot.shape_index in protected or slot.shape_index in assigned or slot.y < 0.18 or slot.shape_index >= len(shapes):
            continue
        item = shapes[slot.shape_index]
        if item.element.getparent() is not None:
            cleared_centers.append((item.left + item.width / 2, item.top + item.height / 2))
    for index, shape in enumerate(list(slide.shapes)):
        if str(shape.name).startswith(("SHN_CONTENT:", "SHN_SOURCE_CLEARED:")):
            continue
        x = shape.left / slide_width
        y = shape.top / slide_height
        area = (shape.width * shape.height) / (slide_width * slide_height)
        overflows = x < -0.001 or y < -0.001 or x + shape.width / slide_width > 1.001 or y + shape.height / slide_height > 1.001
        contains_assigned_text = any(
            shape.left <= center_x <= shape.left + shape.width and shape.top <= center_y <= shape.top + shape.height
            for center_x, center_y in assigned_centers
        )
        contains_only_cleared_text = not contains_assigned_text and any(
            shape.left <= center_x <= shape.left + shape.width and shape.top <= center_y <= shape.top + shape.height
            for center_x, center_y in cleared_centers
        )
        structural = bool(
            getattr(shape, "has_text_frame", False)
            or getattr(shape, "has_table", False)
            or getattr(shape, "has_chart", False)
            or contains_assigned_text
        )
        if y >= 0.16 and 0.002 <= area <= 0.45 and not overflows and not structural and contains_only_cleared_text:
            _remove_shape(shape)
            continue
        if decorativeness == "low" and y >= 0.16 and 0.002 <= area <= 0.45 and not overflows and not structural:
            _remove_shape(shape)
            continue
        if y >= 0.16 and 0.002 <= area <= 0.45 and not overflows:
            shape.name = f"SHN_CONTENT:{plan.slide_id}:SOURCE_COMPONENT:{index}"
        else:
            shape.name = f"SHN_SOURCE_CHROME:{plan.slide_id}:{index}"
    if not group_handled:
        native_elements.extend(plan.elements)
    return native_elements


def template_fidelity_qa(pptx_path: str | Path, selection_report: dict[str, Any]) -> dict[str, Any]:
    prs = Presentation(str(pptx_path))
    findings: list[dict[str, Any]] = []
    slides: list[dict[str, Any]] = []
    selections = selection_report.get("selections", [])
    for index, slide in enumerate(prs.slides):
        selection = selections[index] if index < len(selections) else {}
        source_components = sum("SOURCE_COMPONENT" in str(shape.name) for shape in slide.shapes)
        source_chrome = sum(str(shape.name).startswith("SHN_SOURCE_CHROME:") for shape in slide.shapes)
        if source_components + source_chrome == 0:
            findings.append({"category": "Design", "code": "source_decoration_missing", "severity": "blocking", "slide_id": selection.get("slide_id"), "message": "source-slide mode preserved no source decoration"})
        slides.append(
            {
                "slide_id": selection.get("slide_id"),
                "source_slide_number": selection.get("source_slide_number"),
                "source_component_count": source_components,
                "source_chrome_count": source_chrome,
            }
        )
    return {
        "schema_version": "0.11-mvp.3",
        "status": "PASS" if not findings else "FAIL",
        "summary": {"slide_count": len(prs.slides), "finding_count": len(findings)},
        "slides": slides,
        "findings": findings,
    }
