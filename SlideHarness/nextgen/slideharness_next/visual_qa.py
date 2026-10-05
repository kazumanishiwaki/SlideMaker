from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from xml.etree import ElementTree as ET
import re
import shutil
import subprocess

from PIL import Image, ImageFilter, ImageStat
from pptx import Presentation

from .models import LayoutPlan


SPARSE_RECIPES = {"cover", "section", "statement", "closing_cta"}
TEXT_OVERLAP_LIMIT = 0.45
RENDERED_TEXT_BLOCKING_PT = 5.5
RENDERED_TEXT_REVIEW_PT = 7.5


@dataclass(frozen=True)
class VisualFinding:
    category: str
    code: str
    severity: str
    slide_id: str | None
    message: str
    evidence: dict[str, Any] | None = None

    def to_dict(self) -> dict[str, Any]:
        result: dict[str, Any] = {
            "category": self.category,
            "code": self.code,
            "severity": self.severity,
            "slide_id": self.slide_id,
            "message": self.message,
        }
        if self.evidence:
            result["evidence"] = self.evidence
        return result


def _content_text_shapes(slide: Any) -> list[Any]:
    return [
        shape
        for shape in slide.shapes
        if str(shape.name).startswith("SHN_CONTENT:")
        and getattr(shape, "has_text_frame", False)
        and shape.text_frame.text.strip()
    ]


def _semantic_root(name: str) -> str:
    value = re.sub(r"-(?:number|heading|body|marker|accent)$", "", name)
    value = re.sub(r":source-(?:item|step)-(\d+)-(?:number|heading|body)$", r":source-component-\1", value)
    return value


def _expected_text_overlap(left_name: str, right_name: str) -> bool:
    names = {left_name, right_name}
    if any(name.endswith(":source-title") for name in names) and any(name.endswith(":source-message") for name in names):
        return True
    return _semantic_root(left_name) == _semantic_root(right_name)


def find_text_overlaps(pptx_path: str | Path, plans: list[LayoutPlan]) -> list[VisualFinding]:
    """Return only high-confidence collisions between generated text containers."""

    prs = Presentation(str(pptx_path))
    findings: list[VisualFinding] = []
    for slide_index, slide in enumerate(prs.slides):
        slide_id = plans[slide_index].slide_id if slide_index < len(plans) else f"S{slide_index + 1:03d}"
        shapes = _content_text_shapes(slide)
        for left_index, left in enumerate(shapes):
            for right in shapes[left_index + 1 :]:
                if _expected_text_overlap(str(left.name), str(right.name)):
                    continue
                overlap_width = max(0, min(left.left + left.width, right.left + right.width) - max(left.left, right.left))
                overlap_height = max(0, min(left.top + left.height, right.top + right.height) - max(left.top, right.top))
                intersection = overlap_width * overlap_height
                if not intersection:
                    continue
                smaller_area = min(left.width * left.height, right.width * right.height)
                overlap_ratio = intersection / max(1, smaller_area)
                if overlap_ratio < TEXT_OVERLAP_LIMIT:
                    continue
                findings.append(
                    VisualFinding(
                        "Design",
                        "text_container_overlap",
                        "blocking",
                        slide_id,
                        f"text containers overlap by {overlap_ratio:.0%}: {left.name} and {right.name}",
                        {
                            "left_shape": str(left.name),
                            "right_shape": str(right.name),
                            "overlap_ratio": round(overlap_ratio, 4),
                        },
                    )
                )
    return findings


def _render_metrics(image_path: Path) -> dict[str, Any]:
    with Image.open(image_path).convert("RGB") as source:
        working = source.resize((320, 180))
        grayscale = working.convert("L")
        color_stddev = round(sum(ImageStat.Stat(working).stddev) / 3, 3)
        quantized = working.quantize(colors=16)
        histogram = quantized.histogram()
        dominant_color_ratio = max(histogram) / max(1, sum(histogram))
        edge = grayscale.filter(ImageFilter.FIND_EDGES).crop((8, 5, 312, 175))

    cell = 4
    columns = edge.width // cell
    rows = edge.height // cell
    pixels = edge.load()
    row_densities: list[float] = []
    occupied_total = 0
    for row in range(rows):
        occupied = 0
        for column in range(columns):
            edge_pixels = sum(
                pixels[column * cell + x, row * cell + y] >= 12
                for y in range(cell)
                for x in range(cell)
            )
            occupied += int(edge_pixels >= 2)
        occupied_total += occupied
        row_densities.append(occupied / columns)

    longest_start = 0
    longest_length = 0
    current_start = 0
    current_length = 0
    for index, density in enumerate(row_densities):
        if density < 0.025:
            if current_length == 0:
                current_start = index
            current_length += 1
            if current_length > longest_length:
                longest_start = current_start
                longest_length = current_length
        else:
            current_length = 0
    return {
        "edge_cell_density": round(occupied_total / max(1, rows * columns), 4),
        "quiet_horizontal_band": round(longest_length / max(1, rows), 4),
        "quiet_band_start": round(longest_start / max(1, rows), 4),
        "quiet_band_end": round((longest_start + longest_length) / max(1, rows), 4),
        "dominant_color_ratio": round(dominant_color_ratio, 4),
        "color_stddev": color_stddev,
    }


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _parse_bbox_pages(path: Path) -> list[list[dict[str, Any]]]:
    root = ET.parse(path).getroot()
    pages: list[list[dict[str, Any]]] = []
    for page in (node for node in root.iter() if _local_name(node.tag) == "page"):
        words: list[dict[str, Any]] = []
        lines = [node for node in page.iter() if _local_name(node.tag) == "line"]
        for line_index, line in enumerate(lines):
            for word in (node for node in line.iter() if _local_name(node.tag) == "word"):
                text = "".join(word.itertext()).strip()
                if not text or not re.search(r"[A-Za-z0-9\u3040-\u30ff\u3400-\u9fff]", text):
                    continue
                try:
                    height = float(word.attrib["yMax"]) - float(word.attrib["yMin"])
                except (KeyError, ValueError):
                    continue
                words.append({"text": text, "height_pt": round(height, 3), "line_index": line_index})
        pages.append(words)
    return pages


def _normalize_rendered_text(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9\u3040-\u30ff\u3400-\u9fff]", "", text).casefold()


def _normalize_font_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.split("+", 1)[-1].casefold())


def _expected_fonts(pptx_path: str | Path) -> Counter[str]:
    prs = Presentation(str(pptx_path))
    fonts: Counter[str] = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = []
            if getattr(shape, "has_text_frame", False):
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False):
                frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += 1
    return fonts


def _rendered_font_evidence(pdf_path: Path) -> dict[str, Any]:
    pdffonts = shutil.which("pdffonts")
    if not pdffonts:
        return {"status": "SKIP", "reason": "pdffonts is unavailable", "fonts": []}
    result = subprocess.run([pdffonts, str(pdf_path)], capture_output=True, text=True, timeout=90)
    if result.returncode != 0:
        return {"status": "SKIP", "reason": result.stderr or result.stdout or "pdffonts failed", "fonts": []}
    fonts: list[str] = []
    for line in result.stdout.splitlines():
        if not line.strip() or line.startswith(("name ", "---")):
            continue
        fonts.append(line.split()[0])
    return {"status": "PASS", "fonts": list(dict.fromkeys(fonts))}


def _missing_major_fonts(expected: Counter[str], rendered: list[str]) -> list[dict[str, Any]]:
    rendered_names = {_normalize_font_name(name) for name in rendered}
    aliases = {
        "arial": {"arial", "liberationsans"},
        "calibri": {"calibri", "carlito"},
        "timesnewroman": {"timesnewroman", "liberationserif"},
        "couriernew": {"couriernew", "liberationmono"},
    }
    missing: list[dict[str, Any]] = []
    for name, count in expected.items():
        if count < 2:
            continue
        normalized = _normalize_font_name(name)
        accepted = aliases.get(normalized, {normalized})
        if accepted.isdisjoint(rendered_names):
            missing.append({"font": name, "run_count": count})
    return missing


def _rendered_text_coverage(expected: str, rendered: str) -> tuple[float, float]:
    expected_counts = Counter(expected)
    rendered_counts = Counter(rendered)
    character_coverage = sum(min(count, rendered_counts[character]) for character, count in expected_counts.items()) / max(1, len(expected))
    chunks = [expected[index : index + 6] for index in range(0, len(expected), 6) if len(expected[index : index + 6]) >= 3]
    chunk_coverage = sum(chunk in rendered for chunk in chunks) / max(1, len(chunks))
    return character_coverage, chunk_coverage


def _render_expectations(pptx_path: str | Path) -> list[dict[str, Any]]:
    prs = Presentation(str(pptx_path))
    pages: list[dict[str, Any]] = []
    for slide in prs.slides:
        segments: list[dict[str, str]] = []
        titles: list[str] = []
        semantic_top = 1.0
        semantic_bottom = 0.0
        for shape in slide.shapes:
            name = str(shape.name)
            if not name.startswith("SHN_CONTENT:"):
                continue
            semantic_top = min(semantic_top, max(0.0, shape.top / prs.slide_height))
            semantic_bottom = max(semantic_bottom, min(1.0, (shape.top + shape.height) / prs.slide_height))
            texts: list[str] = []
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells if cell.text.strip())
            for text in texts:
                normalized = _normalize_rendered_text(text)
                if len(normalized) < 4:
                    continue
                segments.append({"shape": name, "text": text, "normalized": normalized})
                if name.endswith(":source-title") or name.endswith("-title"):
                    titles.append(text)
        pages.append(
            {
                "segments": segments,
                "titles": titles,
                "semantic_top": semantic_top if semantic_bottom else None,
                "semantic_bottom": semantic_bottom if semantic_bottom else None,
            }
        )
    return pages


def _rendered_text_pages(pdf_path: Path) -> tuple[list[list[dict[str, Any]]], dict[str, Any]]:
    pdftotext = shutil.which("pdftotext")
    if not pdftotext:
        return [], {"status": "FAIL", "reason": "pdftotext is unavailable"}
    with TemporaryDirectory(prefix="slideharness-next-bbox-") as directory:
        bbox_path = Path(directory) / "bbox.xhtml"
        result = subprocess.run(
            [pdftotext, "-bbox-layout", str(pdf_path), str(bbox_path)],
            capture_output=True,
            text=True,
            timeout=90,
        )
        if result.returncode != 0 or not bbox_path.is_file():
            return [], {"status": "FAIL", "reason": result.stderr or result.stdout or "bbox extraction failed"}
        try:
            pages = _parse_bbox_pages(bbox_path)
        except (ET.ParseError, OSError) as error:
            return [], {"status": "FAIL", "reason": f"bbox parse failed: {error}"}
    return pages, {"status": "PASS", "page_count": len(pages)}


def inspect_visual(
    pptx_path: str | Path,
    plans: list[LayoutPlan],
    render_report: dict[str, Any],
    information_density: str = "medium",
) -> dict[str, Any]:
    """Inspect rendered pages plus generated text geometry.

    The blocking thresholds intentionally target obvious breakage. Borderline
    whitespace and small rendered text are recorded as review findings so the
    full-size human pass has concrete pages and measurements to inspect.
    """

    if information_density not in {"low", "medium", "high"}:
        raise ValueError(f"unsupported information density: {information_density}")
    if render_report.get("status") != "PASS":
        return {
            "schema_version": "0.11-mvp.4",
            "status": "SKIP",
            "reason": "visual QA requires a successful PDF/PNG render",
            "summary": {"slide_count": 0, "blocking_count": 0, "review_count": 0, "finding_count": 0},
            "slides": [],
            "findings": [],
        }

    findings = find_text_overlaps(pptx_path, plans)
    expectations = _render_expectations(pptx_path)
    images = [Path(value) for value in render_report.get("images", [])]
    pdf_value = render_report.get("pdf")
    pdf_path = Path(pdf_value) if pdf_value else None
    missing_images = [str(path) for path in images if not path.is_file()]
    if not images or missing_images:
        findings.append(
            VisualFinding(
                "Technical",
                "rendered_page_missing",
                "blocking",
                None,
                "visual QA received no readable rendered slide images",
                {"missing_images": missing_images},
            )
        )
    if len(images) != len(plans):
        findings.append(
            VisualFinding(
                "Technical",
                "rendered_page_count_mismatch",
                "blocking",
                None,
                f"expected {len(plans)} rendered pages, found {len(images)}",
            )
        )

    text_pages: list[list[dict[str, Any]]] = []
    text_extraction: dict[str, Any]
    if pdf_path and pdf_path.is_file():
        text_pages, text_extraction = _rendered_text_pages(pdf_path)
    else:
        text_extraction = {"status": "FAIL", "reason": "rendered PDF is missing"}
    if text_extraction["status"] != "PASS":
        findings.append(
            VisualFinding(
                "Technical",
                "rendered_text_inspection_unavailable",
                "blocking",
                None,
                str(text_extraction.get("reason", "rendered text inspection failed")),
            )
        )

    expected_fonts = _expected_fonts(pptx_path)
    font_evidence = _rendered_font_evidence(pdf_path) if pdf_path and pdf_path.is_file() else {"status": "SKIP", "reason": "rendered PDF is missing", "fonts": []}
    missing_fonts = _missing_major_fonts(expected_fonts, font_evidence.get("fonts", [])) if font_evidence["status"] == "PASS" else []
    if missing_fonts:
        findings.append(
            VisualFinding(
                "Design",
                "font_substitution_review",
                "review",
                None,
                f"review {len(missing_fonts)} frequently used font family/style(s) absent from the rendered PDF",
                {"missing_fonts": missing_fonts},
            )
        )
    elif len(text_pages) != len(plans):
        findings.append(
            VisualFinding(
                "Technical",
                "rendered_text_page_count_mismatch",
                "blocking",
                None,
                f"expected text evidence for {len(plans)} pages, found {len(text_pages)}",
            )
        )

    slide_reports: list[dict[str, Any]] = []
    density_minimums = {"low": 0.035, "medium": 0.05, "high": 0.06}
    sparse_minimums = {"low": 0.015, "medium": 0.02, "high": 0.025}
    for index, plan in enumerate(plans):
        slide_id = plan.slide_id
        metrics = _render_metrics(images[index]) if index < len(images) and images[index].is_file() else None
        words = text_pages[index] if index < len(text_pages) else []
        expected = expectations[index] if index < len(expectations) else {"segments": [], "titles": []}
        if metrics:
            effective_quiet_band = metrics["quiet_horizontal_band"]
            semantic_top = expected.get("semantic_top")
            semantic_bottom = expected.get("semantic_bottom")
            if metrics["quiet_band_end"] >= 0.99 and semantic_bottom is not None and semantic_bottom > metrics["quiet_band_start"]:
                effective_quiet_band = 1.0 - float(semantic_bottom)
            elif metrics["quiet_band_start"] <= 0.01 and semantic_top is not None and semantic_top < metrics["quiet_band_end"]:
                effective_quiet_band = float(semantic_top)
            metrics["semantic_content_top"] = round(float(semantic_top), 4) if semantic_top is not None else None
            metrics["semantic_content_bottom"] = round(float(semantic_bottom), 4) if semantic_bottom is not None else None
            metrics["effective_quiet_horizontal_band"] = round(max(0.0, effective_quiet_band), 4)
        rendered_text = _normalize_rendered_text("".join(str(word["text"]) for word in words))
        line_texts: dict[int, str] = {}
        for word in words:
            line_index = int(word.get("line_index", -1))
            line_texts[line_index] = line_texts.get(line_index, "") + str(word["text"])
        normalized_lines = [_normalize_rendered_text(text) for text in line_texts.values()]
        missing_segments: list[dict[str, Any]] = []
        for item in expected["segments"]:
            character_coverage, chunk_coverage = _rendered_text_coverage(item["normalized"], rendered_text)
            if character_coverage < 0.95 or chunk_coverage < 0.65:
                missing_segments.append({**item, "character_coverage": character_coverage, "chunk_coverage": chunk_coverage})
        word_heights = [float(word["height_pt"]) for word in words]
        minimum_text_height = min(word_heights) if word_heights else None
        small_words = [word for word in words if float(word["height_pt"]) < RENDERED_TEXT_REVIEW_PT]
        sparse_recipe = plan.recipe in SPARSE_RECIPES
        minimum_density = sparse_minimums[information_density] if sparse_recipe else density_minimums[information_density]

        if metrics and metrics["edge_cell_density"] < minimum_density:
            findings.append(
                VisualFinding(
                    "Design",
                    "rendered_slide_too_sparse",
                    "blocking",
                    slide_id,
                    f"rendered edge density {metrics['edge_cell_density']:.1%} is below {minimum_density:.1%}",
                    metrics,
                )
            )
        if words and missing_segments:
            findings.append(
                VisualFinding(
                    "Design",
                    "rendered_text_missing",
                    "blocking",
                    slide_id,
                    f"{len(missing_segments)} generated text container(s) are absent or clipped in the rendered PDF",
                    {
                        "examples": [
                            {
                                "shape": item["shape"],
                                "text": item["text"][:120],
                                "character_coverage": round(item["character_coverage"], 3),
                                "chunk_coverage": round(item["chunk_coverage"], 3),
                            }
                            for item in missing_segments[:5]
                        ]
                    },
                )
            )
        for title in expected["titles"]:
            normalized_title = _normalize_rendered_text(title)
            if "\n" in title or len(normalized_title) < 4 or normalized_title not in rendered_text:
                continue
            if not any(normalized_title in line for line in normalized_lines):
                findings.append(
                    VisualFinding(
                        "Design",
                        "single_line_title_wrapped",
                        "blocking",
                        slide_id,
                        "a generated title intended for one line wrapped in the rendered PDF",
                        {"title": title},
                    )
                )
        if metrics and not sparse_recipe and metrics["effective_quiet_horizontal_band"] >= 0.50:
            findings.append(
                VisualFinding(
                    "Design",
                    "extreme_rendered_whitespace",
                    "blocking",
                    slide_id,
                    f"a nearly empty horizontal band spans {metrics['effective_quiet_horizontal_band']:.0%} of slide height",
                    metrics,
                )
            )
        elif metrics and not sparse_recipe and metrics["effective_quiet_horizontal_band"] >= 0.24:
            findings.append(
                VisualFinding(
                    "Design",
                    "rendered_whitespace_review",
                    "review",
                    slide_id,
                    f"review a nearly empty horizontal band spanning {metrics['effective_quiet_horizontal_band']:.0%} of slide height",
                    metrics,
                )
            )
        if minimum_text_height is not None and minimum_text_height < RENDERED_TEXT_BLOCKING_PT:
            findings.append(
                VisualFinding(
                    "Design",
                    "rendered_text_too_small",
                    "blocking",
                    slide_id,
                    f"rendered text includes a word only {minimum_text_height:.1f} pt high",
                    {"minimum_word_height_pt": round(minimum_text_height, 3), "examples": small_words[:5]},
                )
            )
        elif minimum_text_height is not None and minimum_text_height < RENDERED_TEXT_REVIEW_PT:
            findings.append(
                VisualFinding(
                    "Design",
                    "rendered_text_size_review",
                    "review",
                    slide_id,
                    f"review rendered text with a minimum word height of {minimum_text_height:.1f} pt",
                    {"minimum_word_height_pt": round(minimum_text_height, 3), "examples": small_words[:5]},
                )
            )

        slide_reports.append(
            {
                "slide_id": slide_id,
                "recipe": plan.recipe,
                "render_metrics": metrics,
                "rendered_word_count": len(words),
                "minimum_rendered_word_height_pt": round(minimum_text_height, 3) if minimum_text_height is not None else None,
                "text_overlap_count": sum(1 for finding in findings if finding.slide_id == slide_id and finding.code == "text_container_overlap"),
            }
        )

    blocking = [finding for finding in findings if finding.severity == "blocking"]
    review = [finding for finding in findings if finding.severity == "review"]
    return {
        "schema_version": "0.11-mvp.4",
        "status": "PASS" if not blocking else "FAIL",
        "summary": {
            "slide_count": len(plans),
            "blocking_count": len(blocking),
            "review_count": len(review),
            "finding_count": len(findings),
        },
        "rendered_text_extraction": text_extraction,
        "rendered_font_evidence": {
            **font_evidence,
            "expected_fonts": dict(expected_fonts),
            "missing_major_fonts": missing_fonts,
        },
        "slides": slide_reports,
        "findings": [finding.to_dict() for finding in findings],
    }
