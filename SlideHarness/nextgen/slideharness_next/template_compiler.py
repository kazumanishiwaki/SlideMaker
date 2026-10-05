from __future__ import annotations

from collections import Counter, defaultdict
from hashlib import sha1, sha256
from pathlib import Path
from typing import Any
from zipfile import ZipFile
import math
import re
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from .models import ChromeElement, LayoutEvidence, TemplateSystem, ThemeTokens


EMU_PER_INCH = 914400
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _hash_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _theme_values(path: Path) -> tuple[dict[str, str], dict[str, str]]:
    colors: dict[str, str] = {}
    fonts: dict[str, str] = {}
    with ZipFile(path) as archive:
        theme_names = sorted(name for name in archive.namelist() if name.startswith("ppt/theme/theme") and name.endswith(".xml"))
        if not theme_names:
            return colors, fonts
        root = ET.fromstring(archive.read(theme_names[0]))
    ns = {"a": A_NS}
    scheme = root.find(".//a:themeElements/a:clrScheme", ns)
    if scheme is not None:
        for item in list(scheme):
            name = item.tag.rsplit("}", 1)[-1]
            child = next(iter(item), None)
            if child is None:
                continue
            value = child.attrib.get("lastClr") or child.attrib.get("val")
            if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                colors[name] = value.upper()
    font_scheme = root.find(".//a:themeElements/a:fontScheme", ns)
    if font_scheme is not None:
        for label, xpath in (("heading", "a:majorFont"), ("body", "a:minorFont")):
            group = font_scheme.find(xpath, ns)
            if group is None:
                continue
            east_asian = group.find("a:ea", ns)
            latin = group.find("a:latin", ns)
            candidate = ""
            if east_asian is not None:
                candidate = east_asian.attrib.get("typeface", "")
            if not candidate and latin is not None:
                candidate = latin.attrib.get("typeface", "")
            if candidate:
                fonts[label] = candidate
    return colors, fonts


def _observed_fonts(prs: Presentation) -> tuple[str | None, str | None]:
    headings: Counter[str] = Counter()
    bodies: Counter[str] = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    name = run.font.name
                    size = run.font.size.pt if run.font.size else None
                    if not name:
                        continue
                    if size and size >= 24:
                        headings[name] += 1
                    else:
                        bodies[name] += 1
    return (headings.most_common(1)[0][0] if headings else None, bodies.most_common(1)[0][0] if bodies else None)


def _tokens(path: Path, prs: Presentation) -> ThemeTokens:
    colors, fonts = _theme_values(path)
    observed_heading, observed_body = _observed_fonts(prs)
    return ThemeTokens(
        primary=colors.get("accent1", "2F5597"),
        secondary=colors.get("accent2", "70AD47"),
        accent=colors.get("accent3", colors.get("accent2", "ED7D31")),
        background=colors.get("lt1", "FFFFFF"),
        surface=colors.get("lt2", "F4F6F8"),
        text=colors.get("dk1", "1F2937"),
        muted_text=colors.get("dk2", "667085"),
        heading_font=observed_heading or fonts.get("heading", "Aptos Display"),
        body_font=observed_body or fonts.get("body", "Aptos"),
    )


def _shape_kind(shape: Any) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if getattr(shape, "has_text_frame", False):
        return "text"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    return "shape"


def _rounded_signature(shape: Any, width: int, height: int) -> tuple[Any, ...]:
    fingerprint = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        fingerprint = sha1(shape.image.blob).hexdigest()[:12]
    elif getattr(shape, "has_text_frame", False):
        fingerprint = " ".join(shape.text.split()).casefold()
    return (
        _shape_kind(shape),
        round(shape.left / width, 2),
        round(shape.top / height, 2),
        round(shape.width / width, 2),
        round(shape.height / height, 2),
        fingerprint,
    )


def _is_edge_or_background(signature: tuple[Any, ...]) -> bool:
    _, x, y, w, h, _ = signature
    area = w * h
    return bool(area >= 0.78 or y <= 0.08 or y + h >= 0.92 or x <= 0.025 or x + w >= 0.975)


def _chrome(prs: Presentation) -> list[ChromeElement]:
    slide_count = len(prs.slides)
    if slide_count < 2:
        return []
    occurrences: dict[tuple[Any, ...], list[tuple[int, int]]] = defaultdict(list)
    for slide_index, slide in enumerate(prs.slides):
        seen_on_slide: set[tuple[Any, ...]] = set()
        for shape_index, shape in enumerate(slide.shapes):
            signature = _rounded_signature(shape, prs.slide_width, prs.slide_height)
            if signature in seen_on_slide or not _is_edge_or_background(signature):
                continue
            seen_on_slide.add(signature)
            occurrences[signature].append((slide_index, shape_index))
    threshold = max(2, math.ceil(slide_count * 0.55))
    chrome: list[ChromeElement] = []
    for signature, positions in occurrences.items():
        if len(positions) < threshold:
            continue
        kind, x, y, w, h, _ = signature
        source_slide, source_shape = positions[0]
        chrome.append(
            ChromeElement(
                source_slide_index=source_slide,
                source_shape_index=source_shape,
                kind=kind,
                x=x,
                y=y,
                w=w,
                h=h,
                occurrence_ratio=round(len(positions) / slide_count, 3),
            )
        )
    return sorted(chrome, key=lambda item: (item.y, item.x, item.kind))


def _layout_evidence(prs: Presentation) -> list[LayoutEvidence]:
    output: list[LayoutEvidence] = []
    for index, layout in enumerate(prs.slide_layouts):
        placeholders = [shape for shape in layout.shapes if getattr(shape, "is_placeholder", False)]
        title_slots = 0
        body_slots = 0
        for shape in placeholders:
            try:
                placeholder_type = shape.placeholder_format.type
            except (AttributeError, ValueError):
                continue
            if placeholder_type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}:
                title_slots += 1
            elif placeholder_type in {
                PP_PLACEHOLDER.BODY,
                PP_PLACEHOLDER.OBJECT,
                PP_PLACEHOLDER.SUBTITLE,
                PP_PLACEHOLDER.VERTICAL_BODY,
                PP_PLACEHOLDER.VERTICAL_OBJECT,
            }:
                body_slots += 1
        non_placeholder_count = max(0, len(layout.shapes) - len(placeholders))
        score = 1.0 / (1.0 + non_placeholder_count + max(0, len(placeholders) - 2) * 0.5)
        output.append(
            LayoutEvidence(
                index=index,
                name=layout.name or f"Layout {index}",
                shape_count=len(layout.shapes),
                placeholder_count=len(placeholders),
                title_slots=title_slots,
                body_slots=body_slots,
                score_as_canvas=round(score, 3),
            )
        )
    return output


def _pattern_evidence(prs: Presentation) -> dict[str, Any]:
    counts = Counter()
    density: list[float] = []
    shape_counts: list[int] = []
    picture_counts: list[int] = []
    for slide in prs.slides:
        boxes = 0
        pictures = 0
        tables = 0
        area = 0.0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                boxes += 1
            area += min(1.0, (shape.width * shape.height) / (prs.slide_width * prs.slide_height))
        counts["slides_with_images"] += int(pictures > 0)
        counts["slides_with_tables"] += int(tables > 0)
        counts["slides_with_4plus_text_blocks"] += int(boxes >= 4)
        density.append(min(1.0, area))
        shape_counts.append(len(slide.shapes))
        picture_counts.append(pictures)
    slide_count = max(1, len(prs.slides))
    return {
        "slide_count": len(prs.slides),
        "image_slide_ratio": round(counts["slides_with_images"] / slide_count, 3),
        "table_slide_ratio": round(counts["slides_with_tables"] / slide_count, 3),
        "dense_text_slide_ratio": round(counts["slides_with_4plus_text_blocks"] / slide_count, 3),
        "mean_shape_area_sum": round(sum(density) / slide_count, 3),
        "mean_slide_local_shape_count": round(sum(shape_counts) / slide_count, 2),
        "mean_slide_local_picture_count": round(sum(picture_counts) / slide_count, 2),
        "supported_recipes": [
            "cover",
            "section",
            "statement",
            "title_body",
            "title_body_callout",
            "two_column_equal",
            "three_column",
            "card_grid_2x2",
            "comparison",
            "kpi_grid",
            "native_bar_chart",
            "native_chart",
            "table",
            "process",
            "image_text",
            "generated_visual",
            "closing_cta",
        ],
    }


def build_template_risk_report(system: TemplateSystem) -> dict[str, Any]:
    mean_local_shapes = float(system.pattern_evidence.get("mean_slide_local_shape_count", 0.0))
    mean_local_pictures = float(system.pattern_evidence.get("mean_slide_local_picture_count", 0.0))
    mean_layout_content = (
        sum(max(0, layout.shape_count - layout.placeholder_count) for layout in system.layouts) / max(1, len(system.layouts))
    )
    completed_slide_heavy = mean_local_shapes >= 8 and mean_layout_content <= 2
    risk_level = "high" if completed_slide_heavy and len(system.chrome) < 6 else "medium" if completed_slide_heavy else "low"
    findings: list[dict[str, Any]] = []
    if completed_slide_heavy:
        findings.append(
            {
                "code": "completed_slide_heavy_template",
                "severity": "review_required" if risk_level == "high" else "warning",
                "message": "Most design detail lives on completed slides rather than masters/layouts; token-and-chrome reconstruction cannot preserve every interior ornament automatically.",
                "evidence": {
                    "mean_slide_local_shape_count": mean_local_shapes,
                    "mean_slide_local_picture_count": mean_local_pictures,
                    "mean_layout_non_placeholder_shape_count": round(mean_layout_content, 2),
                    "repeated_chrome_count": len(system.chrome),
                },
            }
        )
    return {
        "schema_version": system.schema_version,
        "risk_level": risk_level,
        "template_mode": "completed_slide_heavy" if completed_slide_heavy else "master_layout_led",
        "composition_mode": "tokens_chrome_reconstruction",
        "extracted": ["theme colors", "observed fonts", "layout canvas candidates", "exact repeated edge/background chrome"],
        "fallback": ["deterministic component styles", "recipe geometry", "human review for non-repeated slide-local ornament"],
        "lower_fidelity_page_types": ["cover", "section", "full_visual", "photo_grid", "organization"] if completed_slide_heavy else [],
        "findings": findings,
    }


def compile_template(path: str | Path) -> TemplateSystem:
    template_path = Path(path).resolve()
    if not template_path.is_file():
        raise FileNotFoundError(template_path)
    prs = Presentation(str(template_path))
    return TemplateSystem(
        schema_version="0.11-mvp.2",
        template_path=str(template_path),
        template_sha256=_hash_file(template_path),
        slide_width_in=round(prs.slide_width / EMU_PER_INCH, 4),
        slide_height_in=round(prs.slide_height / EMU_PER_INCH, 4),
        slide_count=len(prs.slides),
        tokens=_tokens(template_path, prs),
        chrome=_chrome(prs),
        layouts=_layout_evidence(prs),
        pattern_evidence=_pattern_evidence(prs),
    )
