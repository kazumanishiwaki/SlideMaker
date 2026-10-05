from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any
from zipfile import ZipFile
import mimetypes
import re

from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from .models import SourceBundle, SourceFile, SourceUnit
from .ocr import OCRBackend, OCRConfig, ocr_pdf_pages


SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".docx", ".pdf", ".pptx"}
FACT_PATTERNS = [
    re.compile(r"(?:¥|￥)\s?[\d,]+(?:\.\d+)?(?:万|億)?(?:円)?"),
    re.compile(r"[\d,]+(?:\.\d+)?\s?(?:万|億)?円"),
    re.compile(r"\d+(?:\.\d+)?\s?[〜~\-–]\s?\d+(?:\.\d+)?\s?%"),
    re.compile(r"\d+(?:\.\d+)?\s?%"),
    re.compile(r"\d{4}年\s?\d{1,2}月(?:\s?\d{1,2}日)?"),
    re.compile(r"\d{4}[/-]\d{1,2}(?:[/-]\d{1,2})?"),
    re.compile(r"\d+(?:\.\d+)?\s?(?:pt|日|週間|か月|ヶ月|月間|年間|時間|分|秒|名|件|社|枚|ページ|MB|GB)", re.IGNORECASE),
]
REQUIRED_KEYWORDS = ("必須", "条件", "対象", "期限", "期間", "費用", "料金", "注意", "例外", "禁止", "セキュリティ", "法務", "契約")


def _digest(path: Path) -> str:
    result = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def _semantic(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def extract_facts(text: str) -> list[str]:
    matches: list[tuple[int, str]] = []
    for pattern in FACT_PATTERNS:
        matches.extend((match.start(), match.group(0)) for match in pattern.finditer(text))
    output: list[str] = []
    for _, value in sorted(matches):
        if value not in output:
            output.append(value)
    return output


def classify_unit(kind: str, text: str) -> str:
    normalized = _semantic(text)
    if kind == "heading":
        return "section"
    if kind == "table":
        return "data"
    if any(keyword in normalized for keyword in ("注意", "例外", "ただし", "禁止")):
        return "note"
    if any(keyword in normalized for keyword in ("条件", "対象", "必要", "必須", "期限", "期間")):
        return "condition"
    if extract_facts(normalized):
        return "fact"
    if any(keyword in normalized for keyword in ("例えば", "例:", "例：", "事例")):
        return "example"
    if kind in {"bullet", "numbered_item"}:
        return "support"
    return "claim"


def _raw_unit(kind: str, text: str, section_path: list[str], metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    value = text.strip()
    return {"kind": kind, "text": value, "section_path": list(section_path), "metadata": metadata or {}}


def _parse_markdown_text(text: str) -> list[dict[str, Any]]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    units: list[dict[str, Any]] = []
    sections: list[str] = []
    paragraph: list[str] = []
    in_fence = False
    fence_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            units.append(_raw_unit("paragraph", "\n".join(paragraph), sections))
            paragraph.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if in_fence:
                units.append(_raw_unit("code", "\n".join(fence_lines), sections))
                fence_lines.clear()
            in_fence = not in_fence
            index += 1
            continue
        if in_fence:
            fence_lines.append(line)
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            sections[:] = sections[: level - 1]
            sections.append(title)
            units.append(_raw_unit("heading", title, sections, {"level": level}))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
            flush_paragraph()
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            parsed = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in table_lines]
            columns = parsed[0]
            rows = parsed[2:]
            units.append(_raw_unit("table", "\n".join(table_lines), sections, {"columns": columns, "rows": rows}))
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            units.append(_raw_unit("bullet", bullet.group(1), sections))
            index += 1
            continue
        numbered = re.match(r"^\s*(\d+)[.)]\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            units.append(_raw_unit("numbered_item", numbered.group(2), sections, {"number": int(numbered.group(1))}))
            index += 1
            continue
        if not stripped:
            flush_paragraph()
        else:
            paragraph.append(stripped)
        index += 1
    flush_paragraph()
    if fence_lines:
        units.append(_raw_unit("code", "\n".join(fence_lines), sections))
    return units


def _parse_docx(path: Path) -> list[dict[str, Any]]:
    document = Document(str(path))
    units: list[dict[str, Any]] = []
    sections: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        heading = re.match(r"Heading\s+(\d+)", style, flags=re.IGNORECASE)
        if heading:
            level = int(heading.group(1))
            sections[:] = sections[: level - 1]
            sections.append(text)
            units.append(_raw_unit("heading", text, sections, {"level": level, "style": style}))
        elif "List" in style:
            units.append(_raw_unit("bullet", text, sections, {"style": style}))
        else:
            units.append(_raw_unit("paragraph", text, sections, {"style": style}))
    for table_index, table in enumerate(document.tables):
        values = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not values:
            continue
        units.append(
            _raw_unit(
                "table",
                "\n".join(" | ".join(row) for row in values),
                sections,
                {"table_index": table_index, "columns": values[0], "rows": values[1:]},
            )
        )
    return units


def _parse_pdf(
    path: Path,
    ocr_config: OCRConfig | None = None,
    ocr_backend: OCRBackend | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    reader = PdfReader(str(path))
    units: list[dict[str, Any]] = []
    native_texts = [page.extract_text() or "" for page in reader.pages]
    ocr_pages, ocr_report = ocr_pdf_pages(path, native_texts, config=ocr_config, backend=ocr_backend)
    for page_index, text in enumerate(native_texts):
        page_number = page_index + 1
        ocr_page = ocr_pages.get(page_number)
        if ocr_page is not None:
            for block in ocr_page.blocks:
                if not block.text.strip():
                    continue
                units.append(
                    _raw_unit(
                        "paragraph",
                        block.text,
                        [f"Page {page_number}"],
                        {
                            "page": page_number,
                            "extraction": "ocr",
                            "confidence": round(block.confidence, 3),
                            "bbox": [block.left, block.top, block.width, block.height],
                            "ocr_backend": ocr_page.backend,
                            "ocr_languages": ocr_page.languages,
                            "source_sha256": ocr_page.source_sha256,
                            "page_image_sha256": ocr_page.image_sha256,
                            "preprocessing": ocr_page.preprocessing,
                            "cache_hit": ocr_page.cache_hit,
                        },
                    )
                )
            continue
        blocks = [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]
        if len(blocks) == 1:
            blocks = [line.strip() for line in text.splitlines() if line.strip()]
        for block in blocks:
            units.append(_raw_unit("paragraph", block, [f"Page {page_number}"], {"page": page_number, "extraction": "native"}))
    return units, ocr_report


def _parse_pptx(path: Path) -> list[dict[str, Any]]:
    prs = Presentation(str(path))
    units: list[dict[str, Any]] = []
    for slide_index, slide in enumerate(prs.slides):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape and title_shape.text.strip() else f"Slide {slide_index + 1}"
        sections = [title]
        units.append(_raw_unit("heading", title, sections, {"level": 2, "slide": slide_index + 1}))
        for shape in slide.shapes:
            if shape is title_shape or not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                value = paragraph.text.strip()
                if value:
                    kind = "bullet" if paragraph.level > 0 or paragraph.text.startswith(("•", "-")) else "paragraph"
                    units.append(_raw_unit(kind, value.lstrip("•- "), sections, {"slide": slide_index + 1, "level": paragraph.level}))
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            values = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
            if values:
                units.append(_raw_unit("table", "\n".join(" | ".join(row) for row in values), sections, {"slide": slide_index + 1, "columns": values[0], "rows": values[1:]}))
    return units


def _embedded_assets(path: Path, source_id: str) -> list[dict[str, Any]]:
    prefix = "word/media/" if path.suffix.lower() == ".docx" else "ppt/media/" if path.suffix.lower() == ".pptx" else None
    if not prefix:
        return []
    output: list[dict[str, Any]] = []
    with ZipFile(path) as archive:
        for index, name in enumerate(sorted(item for item in archive.namelist() if item.startswith(prefix) and not item.endswith("/"))):
            blob = archive.read(name)
            output.append(
                {
                    "asset_id": f"A-{source_id}-{index + 1:03d}",
                    "source_id": source_id,
                    "package_name": name,
                    "filename": Path(name).name,
                    "size_bytes": len(blob),
                    "sha256": sha256(blob).hexdigest(),
                    "status": "embedded_not_extracted",
                }
            )
    return output


def _parse_file(
    path: Path,
    ocr_config: OCRConfig | None = None,
    ocr_backend: OCRBackend | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return _parse_markdown_text(path.read_text(encoding="utf-8-sig")), None
    if suffix == ".docx":
        return _parse_docx(path), None
    if suffix == ".pdf":
        return _parse_pdf(path, ocr_config, ocr_backend)
    if suffix == ".pptx":
        return _parse_pptx(path), None
    raise ValueError(f"unsupported source format: {suffix}")


def ingest_sources(
    paths: list[str | Path],
    *,
    ocr_config: OCRConfig | None = None,
    ocr_backend: OCRBackend | None = None,
) -> SourceBundle:
    if not paths:
        raise ValueError("at least one source file is required")
    source_files: list[SourceFile] = []
    source_units: list[SourceUnit] = []
    assets: list[dict[str, Any]] = []
    ocr_reports: list[dict[str, Any]] = []
    global_order = 0
    for file_index, value in enumerate(paths):
        path = Path(value).resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError(f"unsupported source format: {path.suffix}")
        file_hash = _digest(path)
        source_id = f"SRC-{file_index + 1:02d}-{file_hash[:8]}"
        media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        source_files.append(SourceFile(source_id, str(path), path.name, media_type, path.stat().st_size, file_hash))
        raw_units, ocr_report = _parse_file(path, ocr_config, ocr_backend)
        if ocr_report is not None:
            ocr_reports.append(ocr_report)
        for local_index, raw in enumerate(raw_units):
            global_order += 1
            source_text = str(raw["text"])
            semantic_text = _semantic(source_text)
            facts = extract_facts(source_text)
            classification = classify_unit(str(raw["kind"]), source_text)
            required = bool(facts or classification in {"condition", "note"} or any(keyword in semantic_text for keyword in REQUIRED_KEYWORDS))
            source_units.append(
                SourceUnit(
                    unit_id=f"U-{file_index + 1:02d}-{local_index + 1:04d}",
                    source_id=source_id,
                    order=global_order,
                    kind=str(raw["kind"]),
                    section_path=list(raw["section_path"]),
                    source_text=source_text,
                    semantic_text=semantic_text,
                    classification=classification,
                    facts=facts,
                    required=required,
                    metadata=dict(raw.get("metadata") or {}),
                )
            )
        assets.extend(_embedded_assets(path, source_id))
    findings = [item for report in ocr_reports for item in report.get("findings", [])]
    blocking = [item for item in findings if item.get("severity") == "blocking"]
    ocr_summary = {
        "schema_version": "0.4.0",
        "status": "FAIL" if blocking else "PASS",
        "file_count": len(ocr_reports),
        "ocr_page_count": sum(int(report.get("ocr_page_count", 0)) for report in ocr_reports),
        "cache_hit_count": sum(int(report.get("cache_hit_count", 0)) for report in ocr_reports),
        "blocking_count": len(blocking),
        "review_count": sum(1 for item in findings if item.get("severity") == "review"),
        "files": ocr_reports,
        "findings": findings,
    }
    return SourceBundle("0.11-mvp.2", source_files, source_units, assets, ocr_summary)
