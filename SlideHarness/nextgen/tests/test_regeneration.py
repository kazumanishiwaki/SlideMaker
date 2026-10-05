from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation

from slideharness_next.composer import compose_deck
from slideharness_next.models import DeckSpec, SlideSpec
from slideharness_next.planner import plan_deck
from slideharness_next.regeneration import isolate_regenerated_slide
from slideharness_next.template_compiler import compile_template


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "nextgen" / "examples" / "test-case" / "inputs" / "source.pptx"
IMAGE = ROOT / "nextgen" / "examples" / "representative" / "assets" / "generated-template-system.png"


class RegenerationIsolationTest(unittest.TestCase):
    def test_only_target_slide_package_changes(self) -> None:
        slides = [
            SlideSpec("S001", "Cover", "Current", "cover", ["U1"], {}),
            SlideSpec("S002", "Target", "Before", "statement", ["U2"], {"statement": "Before"}),
            SlideSpec("S003", "Closing", "Done", "closing_cta", ["U3"], {"actions": ["Done"]}),
        ]
        base_deck = DeckSpec("Isolation", slides)
        candidate_slides = deepcopy(slides)
        candidate_slides[1] = SlideSpec("S002", "Target", "After", "statement", ["U2"], {"statement": "After"})
        candidate_deck = DeckSpec("Isolation", candidate_slides)
        system = compile_template(TEMPLATE)
        with TemporaryDirectory() as directory:
            root = Path(directory)
            base = root / "base.pptx"
            candidate = root / "candidate.pptx"
            output = root / "isolated.pptx"
            report_path = root / "qa-page-isolation.json"
            compose_deck(TEMPLATE, base_deck, plan_deck(base_deck), system, base)
            compose_deck(TEMPLATE, candidate_deck, plan_deck(candidate_deck), system, candidate)
            report = isolate_regenerated_slide(base, candidate, 1, output, report_path)
            self.assertEqual(report["status"], "PASS")
            self.assertTrue(report["non_target_unchanged"])
            self.assertEqual(report["non_target_slide_count"], 2)
            self.assertEqual(
                report["non_target_fingerprints_before"],
                report["non_target_fingerprints_after"],
            )
            prs = Presentation(output)
            self.assertIn("After", "\n".join(shape.text for shape in prs.slides[1].shapes if getattr(shape, "has_text_frame", False)))

    def test_target_only_media_is_removed_when_replacement_no_longer_uses_it(self) -> None:
        slides = [
            SlideSpec("S001", "Cover", "Current", "cover", ["U1"], {}),
            SlideSpec(
                "S002",
                "Target",
                "Illustrated",
                "generated_visual",
                ["U2"],
                {"visual": {"asset_id": "S002-image", "type": "flow", "prompt": "Flow", "required_text": ["Flow"]}},
            ),
            SlideSpec("S003", "Closing", "Done", "closing_cta", ["U3"], {"actions": ["Done"]}),
        ]
        base_deck = DeckSpec("Isolation", slides)
        candidate_slides = deepcopy(slides)
        candidate_slides[1] = SlideSpec(
            "S002",
            "Target",
            "Text only",
            "statement",
            ["U2"],
            {"statement": "Text only"},
        )
        candidate_deck = DeckSpec("Isolation", candidate_slides)
        system = compile_template(TEMPLATE)
        with TemporaryDirectory() as directory:
            root = Path(directory)
            base = root / "base.pptx"
            candidate = root / "candidate.pptx"
            output = root / "isolated.pptx"
            compose_deck(
                TEMPLATE,
                base_deck,
                plan_deck(base_deck),
                system,
                base,
                {"S002-image": IMAGE},
            )
            compose_deck(TEMPLATE, candidate_deck, plan_deck(candidate_deck), system, candidate)
            report = isolate_regenerated_slide(base, candidate, 1, output)
            self.assertEqual(report["status"], "PASS")
            self.assertTrue(
                any(item.startswith("ppt/media/") for item in report["removed_dependencies"]),
                report,
            )


if __name__ == "__main__":
    unittest.main()
