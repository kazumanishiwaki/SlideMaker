from pathlib import Path
from tempfile import TemporaryDirectory
import shutil
import unittest

from docx import Document
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont

try:
    from reportlab.pdfgen import canvas
except ImportError:  # pragma: no cover - optional test helper
    canvas = None

from slideharness_next.ingest import extract_facts, ingest_sources
from slideharness_next.ocr import OCRBlock, OCRConfig, OCRPage


class FakeOCRBackend:
    available = True
    identity = "fake-ocr:1"

    def __init__(self) -> None:
        self.calls = 0

    def recognize(self, image_path: Path, *, page_number: int, languages: str) -> OCRPage:
        self.calls += 1
        text = "対象は25社、期間は3か月です。"
        return OCRPage(
            page_number=page_number,
            text=text,
            confidence=96.0,
            blocks=[OCRBlock(text, 96.0, 10, 20, 300, 40, 1)],
            backend=self.identity,
            languages=languages,
            image_sha256="fake-image-hash",
        )


class IngestTest(unittest.TestCase):
    def test_markdown_preserves_units_sections_tables_and_facts(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "brief.md"
            source.write_text(
                "# Proposal\n\n## Terms\n\n- 料金は月額10万円です。\n- 期間は2026年9月1日からです。\n\n| A | B |\n|---|---|\n| x | y |\n",
                encoding="utf-8",
            )
            bundle = ingest_sources([source])
            self.assertEqual(len(bundle.files), 1)
            self.assertTrue(any(unit.kind == "table" for unit in bundle.units))
            self.assertTrue(any("10万円" in unit.facts for unit in bundle.units))
            self.assertTrue(any("2026年9月1日" in unit.facts for unit in bundle.units))
            self.assertTrue(all(unit.unit_id.startswith("U-01-") for unit in bundle.units))

    def test_docx_and_pptx_are_supported(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "brief.docx"
            document = Document()
            document.add_heading("Word Brief", level=1)
            document.add_paragraph("対象は3社です。")
            document.save(docx_path)
            pptx_path = root / "brief.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "PowerPoint Brief"
            slide.placeholders[1].text = "期間は3か月です。"
            prs.save(pptx_path)
            bundle = ingest_sources([docx_path, pptx_path])
            self.assertEqual(len(bundle.files), 2)
            self.assertTrue(any(unit.source_text == "Word Brief" for unit in bundle.units))
            self.assertTrue(any(unit.source_text == "PowerPoint Brief" for unit in bundle.units))
            self.assertTrue(any("3か月" in unit.facts for unit in bundle.units))

    def test_fact_ranges_remain_exact(self) -> None:
        self.assertIn("55〜85%", extract_facts("占有率は55〜85%、本文は9.5ptです。"))
        self.assertIn("9.5pt", extract_facts("占有率は55〜85%、本文は9.5ptです。"))

    @unittest.skipIf(canvas is None, "reportlab is unavailable")
    def test_text_pdf_is_supported(self) -> None:
        with TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "brief.pdf"
            writer = canvas.Canvas(str(pdf_path))
            writer.drawString(72, 720, "PDF brief covers 25% of the target.")
            writer.save()
            bundle = ingest_sources([pdf_path])
            self.assertTrue(any("PDF brief" in unit.source_text for unit in bundle.units))
            self.assertTrue(any("25%" in unit.facts for unit in bundle.units))

    @unittest.skipIf(canvas is None, "reportlab is unavailable")
    def test_scanned_pdf_uses_ocr_metadata_and_cache(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "scan.png"
            image = Image.new("RGB", (1200, 700), "white")
            ImageDraw.Draw(image).text((80, 120), "scanned source", fill="black")
            image.save(image_path)
            pdf_path = root / "scan.pdf"
            writer = canvas.Canvas(str(pdf_path), pagesize=(600, 350))
            writer.drawImage(str(image_path), 0, 0, width=600, height=350)
            writer.save()
            backend = FakeOCRBackend()
            config = OCRConfig(cache_dir=root / "cache", native_text_min_chars=10)
            first = ingest_sources([pdf_path], ocr_config=config, ocr_backend=backend)
            second = ingest_sources([pdf_path], ocr_config=config, ocr_backend=backend)
            self.assertEqual(first.ocr_report["status"], "PASS")
            self.assertEqual(first.ocr_report["ocr_page_count"], 1)
            self.assertEqual(second.ocr_report["cache_hit_count"], 1)
            self.assertEqual(backend.calls, 1)
            ocr_unit = next(unit for unit in first.units if unit.metadata.get("extraction") == "ocr")
            self.assertEqual(ocr_unit.metadata["page"], 1)
            self.assertEqual(ocr_unit.metadata["confidence"], 96.0)
            self.assertIn("25社", ocr_unit.source_text)
            self.assertEqual(ocr_unit.metadata["preprocessing"], ["rasterize:220dpi", "grayscale", "autocontrast"])
            self.assertEqual(len(ocr_unit.metadata["source_sha256"]), 64)

    @unittest.skipIf(canvas is None or not shutil.which("tesseract") or not (shutil.which("pdftoppm") or shutil.which("pdftocairo")), "OCR tools are unavailable")
    def test_scanned_pdf_runs_real_local_tesseract(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "real-scan.png"
            image = Image.new("RGB", (1600, 900), "white")
            font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 74)
            draw = ImageDraw.Draw(image)
            draw.text((100, 170), "Revenue 2026", fill="black", font=font)
            draw.text((100, 310), "Target 42 percent", fill="black", font=font)
            image.save(image_path)
            pdf_path = root / "real-scan.pdf"
            writer = canvas.Canvas(str(pdf_path), pagesize=(800, 450))
            writer.drawImage(str(image_path), 0, 0, width=800, height=450)
            writer.save()
            config = OCRConfig(cache_dir=root / "cache", native_text_min_chars=10, blocking_confidence=10, review_confidence=30)
            bundle = ingest_sources([pdf_path], ocr_config=config)
            recognized = "\n".join(unit.source_text for unit in bundle.units)
            self.assertEqual(bundle.ocr_report["status"], "PASS", bundle.ocr_report)
            self.assertIn("Revenue", recognized)
            self.assertIn("2026", recognized)


if __name__ == "__main__":
    unittest.main()
