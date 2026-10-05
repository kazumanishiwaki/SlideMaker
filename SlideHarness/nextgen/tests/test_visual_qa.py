from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.util import Inches

from slideharness_next.models import LayoutPlan
from slideharness_next.visual_qa import _normalize_font_name, find_text_overlaps, inspect_visual


class VisualQATest(unittest.TestCase):
    def test_subset_font_names_normalize_to_the_pptx_family_style(self) -> None:
        self.assertEqual(_normalize_font_name("BAAAAA+PretendardJP-Bold"), "pretendardjpbold")
        self.assertEqual(_normalize_font_name("Pretendard JP Bold"), "pretendardjpbold")

    def test_overlapping_generated_text_containers_are_blocking(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "overlap.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            left = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1.5))
            left.name = "SHN_CONTENT:S001:left"
            left.text_frame.text = "Left text"
            right = slide.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(4), Inches(1.5))
            right.name = "SHN_CONTENT:S001:right"
            right.text_frame.text = "Right text"
            prs.save(output)
            plan = LayoutPlan("S001", "title_body", [], [], 1)
            findings = find_text_overlaps(output, [plan])
            self.assertEqual(len(findings), 1)
            self.assertEqual(findings[0].code, "text_container_overlap")
            self.assertEqual(findings[0].severity, "blocking")
            incomplete_render = inspect_visual(output, [plan], {"status": "PASS", "images": []})
            codes = {item["code"] for item in incomplete_render["findings"]}
            self.assertEqual(incomplete_render["status"], "FAIL")
            self.assertIn("rendered_page_count_mismatch", codes)

    def test_visual_qa_skips_without_successful_render(self) -> None:
        plan = LayoutPlan("S001", "cover", [], [], 1)
        report = inspect_visual("not-read.pptx", [plan], {"status": "SKIP", "images": []})
        self.assertEqual(report["status"], "SKIP")
        self.assertIn("successful", report["reason"])


if __name__ == "__main__":
    unittest.main()
