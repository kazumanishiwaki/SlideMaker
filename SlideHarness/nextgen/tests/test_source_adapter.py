from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slideharness_next.composer import compose_deck
from slideharness_next.models import DeckSpec
from slideharness_next.planner import plan_deck
from slideharness_next.source_adapter import analyze_source_slides, select_source_slides, template_fidelity_qa
from slideharness_next.template_compiler import compile_template


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "nextgen" / "examples" / "test-case" / "inputs" / "source.pptx"
DECK_SPEC = ROOT / "nextgen" / "examples" / "source-pipeline" / "output" / "deck_spec.json"


class SourceAdapterTest(unittest.TestCase):
    def setUp(self) -> None:
        self.deck = DeckSpec.load(DECK_SPEC)
        self.plans = plan_deck(self.deck)
        self.evidence = analyze_source_slides(TEMPLATE)
        self.report = select_source_slides(self.evidence, self.plans)

    def test_selects_recipe_appropriate_completed_slides(self) -> None:
        selected = {item["recipe"]: item["source_slide_number"] for item in self.report["selections"]}
        self.assertEqual(selected["cover"], 1)
        self.assertEqual(selected["statement"], 2)
        self.assertEqual(selected["three_column"], 7)
        self.assertEqual(selected["card_grid_2x2"], 15)
        self.assertEqual(selected["process"], 14)
        self.assertEqual(selected["table"], 11)
        self.assertEqual(selected["closing_cta"], 17)

    def test_source_slide_mode_preserves_components_and_native_elements(self) -> None:
        system = compile_template(TEMPLATE)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "deck.pptx"
            compose_deck(
                TEMPLATE,
                self.deck,
                self.plans,
                system,
                output,
                composition_mode="source-slides",
                source_slide_report=self.report,
                source_evidence=self.evidence,
            )
            prs = Presentation(output)
            self.assertEqual(len(prs.slides), len(self.plans))
            self.assertGreater(sum("SOURCE_COMPONENT" in str(shape.name) for shape in prs.slides[2].shapes), 0)
            self.assertGreaterEqual(sum("-step-" in str(shape.name) for shape in prs.slides[4].shapes), 4)
            process_steps = [shape for shape in prs.slides[4].shapes if ":source-step-" in str(shape.name)]
            self.assertLess(min(shape.top / prs.slide_height for shape in process_steps), 0.45)
            self.assertEqual(sum(1 for shape in prs.slides[6].shapes if getattr(shape, "has_table", False)), 1)
            closing = prs.slides[7]
            closing_bottom_slots = [
                shape
                for shape in closing.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and shape.top / prs.slide_height >= 0.79
                and 0.3 <= shape.width / prs.slide_width <= 0.34
                and 0.07 <= shape.height / prs.slide_height <= 0.08
            ]
            self.assertEqual(len(closing_bottom_slots), 2)
            closing_labels = [shape for shape in closing.shapes if str(shape.name).endswith((":source-body", ":source-actions"))]
            self.assertEqual(len(closing_labels), 2)
            self.assertTrue(all(shape.height / prs.slide_height >= 0.05 for shape in closing_labels))
            self.assertEqual(sum(len(slide._element.xpath(".//a:outerShdw")) for slide in prs.slides), 0)
            self.assertEqual(sum(len(slide._element.xpath(".//a:prstShdw")) for slide in prs.slides), 0)
            self.assertEqual(template_fidelity_qa(output, self.report)["status"], "PASS")


if __name__ == "__main__":
    unittest.main()
