from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation

from slideharness_next.composer import compose_deck
from slideharness_next.models import DeckSpec, SlideSpec, SourceBundle, SourceFile, SourceUnit
from slideharness_next.planner import plan_deck
from slideharness_next.provenance import build_source_notes, inspect_provenance
from slideharness_next.template_compiler import compile_template


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "nextgen" / "examples" / "test-case" / "inputs" / "source.pptx"


class ProvenanceTest(unittest.TestCase):
    def test_source_units_are_written_to_speaker_notes_and_verified(self) -> None:
        deck = DeckSpec(
            "Notes test",
            [
                SlideSpec(
                    "S001",
                    "Evidence",
                    "Traceable claim",
                    "statement",
                    ["U-01-0001"],
                    {
                        "statement": "Traceable claim",
                        "asset": {
                            "asset_id": "A-001",
                            "source_url": "https://assets.example/image",
                            "source_note": "licensed source image",
                        },
                    },
                )
            ],
        )
        bundle = SourceBundle(
            "test",
            [SourceFile("SRC-01", "/private/source/brief.md", "brief.md", "text/markdown", 10, "abc")],
            [
                SourceUnit(
                    "U-01-0001",
                    "SRC-01",
                    1,
                    "paragraph",
                    ["Evidence"],
                    "Traceable claim https://example.com/report.",
                    "Traceable claim https://example.com/report.",
                    "claim",
                )
            ],
        )
        notes, ledger = build_source_notes(deck, bundle)
        self.assertEqual(ledger["slides"][0]["records"][0]["filename"], "brief.md")
        self.assertNotIn("/private/source", notes["S001"])
        self.assertIn("https://example.com/report", notes["S001"])
        self.assertIn("A-001 | https://assets.example/image", notes["S001"])
        with TemporaryDirectory() as directory:
            output = Path(directory) / "deck.pptx"
            compose_deck(TEMPLATE, deck, plan_deck(deck), compile_template(TEMPLATE), output, source_notes=notes)
            prs = Presentation(output)
            note_text = prs.slides[0].notes_slide.notes_text_frame.text
            self.assertTrue(note_text.startswith("[Sources]"))
            self.assertIn("U-01-0001", note_text)
            self.assertEqual(inspect_provenance(output, deck, bundle)["status"], "PASS")

            prs.slides[0].notes_slide.notes_text_frame.text = ""
            missing = Path(directory) / "missing-notes.pptx"
            prs.save(missing)
            report = inspect_provenance(missing, deck, bundle)
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("sources_block_missing", {item["code"] for item in report["findings"]})

            prs = Presentation(output)
            prs.slides[0].notes_slide.notes_text_frame.text = "[Sources]\n- U-01-0001 | brief.md | Evidence"
            missing_urls = Path(directory) / "missing-urls.pptx"
            prs.save(missing_urls)
            url_report = inspect_provenance(missing_urls, deck, bundle)
            codes = {item["code"] for item in url_report["findings"]}
            self.assertEqual(url_report["status"], "FAIL")
            self.assertIn("source_note_url_missing", codes)
            self.assertIn("source_note_asset_missing", codes)


if __name__ == "__main__":
    unittest.main()
