from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from slideharness_next.cli import _qa_passes
from slideharness_next.package_qa import inspect_package


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "nextgen" / "examples" / "test-case" / "inputs" / "source.pptx"


class PackageQATest(unittest.TestCase):
    @staticmethod
    def _rewrite(source: Path, target: Path, transform) -> None:
        with zipfile.ZipFile(source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist() if not name.endswith("/")}
        transform(entries)
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, payload in entries.items():
                archive.writestr(name, payload)

    def test_valid_template_passes_with_baseline(self) -> None:
        report = inspect_package(TEMPLATE, TEMPLATE)
        self.assertEqual(report["status"], "PASS", report["findings"])
        self.assertEqual(report["summary"]["blocking_count"], 0)
        self.assertEqual(report["summary"]["baseline_finding_count"], 0)

    def test_missing_relationship_target_is_blocking(self) -> None:
        with TemporaryDirectory() as directory:
            broken = Path(directory) / "broken.pptx"
            with zipfile.ZipFile(TEMPLATE) as source:
                entries = {name: source.read(name) for name in source.namelist() if not name.endswith("/")}
            removable = next(name for name in entries if name.startswith("ppt/theme/") and name.endswith(".xml"))
            entries.pop(removable)
            with zipfile.ZipFile(broken, "w", compression=zipfile.ZIP_DEFLATED) as target:
                for name, payload in entries.items():
                    target.writestr(name, payload)
            report = inspect_package(broken, TEMPLATE)
            codes = {item["code"] for item in report["findings"]}
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("relationship_target_missing", codes)
            self.assertGreater(report["summary"]["regression_count"], 0)

    def test_render_skip_is_fail_closed_by_default(self) -> None:
        skipped = {"status": "SKIP", "reason": "test"}
        self.assertFalse(_qa_passes(skipped))
        self.assertTrue(_qa_passes(skipped, allow_skip=True))
        self.assertTrue(_qa_passes({"status": "PASS"}))

    def test_empty_structural_placeholder_is_blocking(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "empty-placeholder.pptx"
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[0])
            prs.save(path)
            report = inspect_package(path)
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("empty_structural_placeholder", {item["code"] for item in report["findings"]})

    def test_missing_relationship_reference_is_blocking(self) -> None:
        with TemporaryDirectory() as directory:
            broken = Path(directory) / "broken-reference.pptx"

            def transform(entries: dict[str, bytes]) -> None:
                slide = entries["ppt/slides/slide1.xml"]
                relationship_id = next(item.split(b'"', 1)[0] for item in slide.split(b'r:embed="')[1:])
                entries["ppt/slides/slide1.xml"] = slide.replace(relationship_id, b"rIdMissing", 1)

            self._rewrite(TEMPLATE, broken, transform)
            report = inspect_package(broken, TEMPLATE)
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("relationship_reference_missing", {item["code"] for item in report["findings"]})

    def test_unreferenced_media_is_blocking(self) -> None:
        with TemporaryDirectory() as directory:
            broken = Path(directory) / "orphan-media.pptx"

            def transform(entries: dict[str, bytes]) -> None:
                image_name = next(name for name in entries if name.startswith("ppt/media/") and "." in Path(name).name)
                extension = Path(image_name).suffix
                entries[f"ppt/media/orphan{extension}"] = entries[image_name]

            self._rewrite(TEMPLATE, broken, transform)
            report = inspect_package(broken, TEMPLATE)
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("orphan_package_part", {item["code"] for item in report["findings"]})

    def test_template_theme_must_remain_byte_identical(self) -> None:
        with TemporaryDirectory() as directory:
            changed = Path(directory) / "changed-theme.pptx"

            def transform(entries: dict[str, bytes]) -> None:
                theme = entries["ppt/theme/theme1.xml"]
                entries["ppt/theme/theme1.xml"] = theme.replace(b"</a:theme>", b"  </a:theme>")

            self._rewrite(TEMPLATE, changed, transform)
            report = inspect_package(changed, TEMPLATE)
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("template_theme_changed", {item["code"] for item in report["findings"]})

    def test_chart_axis_reference_without_definition_is_blocking(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "chart.pptx"
            broken = Path(directory) / "broken-chart.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            data = CategoryChartData()
            data.categories = ["A", "B"]
            data.add_series("Value", [1, 2])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)
            prs.save(source)

            def transform(entries: dict[str, bytes]) -> None:
                chart_name = next(name for name in entries if name.startswith("ppt/charts/chart") and name.endswith(".xml"))
                payload = entries[chart_name]
                start = payload.index(b"<c:barChart>")
                axis = payload.index(b'<c:axId val="', start) + len(b'<c:axId val="')
                end = payload.index(b'"', axis)
                entries[chart_name] = payload[:axis] + b"999999999" + payload[end:]

            self._rewrite(source, broken, transform)
            report = inspect_package(broken)
            self.assertEqual(report["status"], "FAIL")
            self.assertIn("chart_axis_definition_missing", {item["code"] for item in report["findings"]})


if __name__ == "__main__":
    unittest.main()
