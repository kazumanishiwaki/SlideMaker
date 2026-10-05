from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from slideharness_next.charts import choose_chart_type, validate_chart_content
from slideharness_next.composer import compose_deck
from slideharness_next.models import DeckSpec, SlideSpec
from slideharness_next.package_qa import inspect_package
from slideharness_next.planner import plan_deck
from slideharness_next.qa import inspect_pptx
from slideharness_next.template_compiler import compile_template


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "nextgen" / "examples" / "test-case" / "inputs" / "source.pptx"


class NativeChartTest(unittest.TestCase):
    def test_deterministic_chart_selection(self) -> None:
        self.assertEqual(choose_chart_type({"categories": ["2025/1", "2025/2"], "series": [{"values": [1, 2]}]}), "line")
        self.assertEqual(choose_chart_type({"intent": "composition", "categories": ["A", "B"], "series": [{"values": [40, 60]}]}), "doughnut")
        self.assertEqual(choose_chart_type({"series": [{"x_values": [1, 2], "y_values": [3, 4]}]}), "scatter")
        self.assertEqual(choose_chart_type({"categories": ["A", "B"], "series": [{"values": [1, 2]}]}), "column")

    def test_chart_contract_rejects_invalid_shapes(self) -> None:
        self.assertTrue(validate_chart_content({"chart_type": "pie", "categories": ["A"], "series": []}))
        self.assertTrue(validate_chart_content({"chart_type": "scatter", "series": [{"x_values": [1], "y_values": [2, 3]}]}))
        self.assertTrue(validate_chart_content({"chart_type": "line", "categories": ["Q1"], "series": [{"values": ["not-a-number"]}]}))
        self.assertTrue(validate_chart_content({"chart_type": "pie", "categories": ["A", "B"], "series": [{"values": [0, 0]}]}))
        self.assertTrue(validate_chart_content({"chart_type": "unknown", "categories": ["A"], "series": [{"values": [1]}]}))

    def test_all_supported_types_remain_editable_powerpoint_charts(self) -> None:
        category_types = ["bar", "column", "stacked-bar", "stacked-column", "line", "area", "pie", "doughnut"]
        slides = []
        for index, chart_type in enumerate(category_types, start=1):
            series = [{"name": "Value", "values": [25, 35, 40]}]
            if chart_type in {"stacked-bar", "stacked-column"}:
                series.append({"name": "Other", "values": [15, 25, 20]})
            slides.append(
                SlideSpec(
                    f"S{index:03d}",
                    f"{chart_type} chart",
                    "Editable chart",
                    "native_chart",
                    [f"U{index}"],
                    {
                        "chart": {"chart_type": chart_type, "categories": ["A", "B", "C"], "series": series},
                        "insight": {"heading": "Insight", "body": "PowerPoint native"},
                    },
                )
            )
        slides.append(
            SlideSpec(
                "S009",
                "scatter chart",
                "Editable scatter",
                "native_chart",
                ["U9"],
                {
                    "chart": {"chart_type": "scatter", "series": [{"name": "XY", "x_values": [1, 2, 3], "y_values": [3, 5, 8]}]},
                    "insight": {"heading": "Insight", "body": "PowerPoint native"},
                },
            )
        )
        deck = DeckSpec("Native chart matrix", slides)
        plans = plan_deck(deck)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "charts.pptx"
            compose_deck(TEMPLATE, deck, plans, compile_template(TEMPLATE), output)
            prs = Presentation(output)
            actual = [next(shape.chart for shape in slide.shapes if getattr(shape, "has_chart", False)).chart_type for slide in prs.slides]
            descriptions = [next(shape for shape in slide.shapes if getattr(shape, "has_chart", False))._element.xpath(".//p:cNvPr")[0].get("descr") for slide in prs.slides]
            expected = [
                XL_CHART_TYPE.BAR_CLUSTERED,
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                XL_CHART_TYPE.BAR_STACKED,
                XL_CHART_TYPE.COLUMN_STACKED,
                XL_CHART_TYPE.LINE_MARKERS,
                XL_CHART_TYPE.AREA,
                XL_CHART_TYPE.PIE,
                XL_CHART_TYPE.DOUGHNUT,
                # python-pptx reports its marker-only XY scatter as the
                # line-marker enum even when the connecting line is hidden.
                XL_CHART_TYPE.XY_SCATTER_LINES,
            ]
            self.assertEqual(actual, expected)
            self.assertTrue(all(description and "Editable" in description for description in descriptions))
            self.assertEqual(inspect_pptx(output, plans)["status"], "PASS")
            self.assertEqual(inspect_package(output, TEMPLATE)["status"], "PASS")


if __name__ == "__main__":
    unittest.main()
