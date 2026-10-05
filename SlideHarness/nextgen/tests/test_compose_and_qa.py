from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from slideharness_next.composer import compose_deck
from slideharness_next.models import DeckSpec, SlideSpec
from slideharness_next.planner import plan_deck
from slideharness_next.qa import inspect_pptx
from slideharness_next.template_compiler import compile_template


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "nextgen" / "examples" / "test-case" / "inputs" / "source.pptx"


class ComposeAndQATest(unittest.TestCase):
    def test_editable_chart_table_and_process_survive_save(self) -> None:
        deck = DeckSpec(
            "Editable native visual test",
            [
                SlideSpec(
                    "S001",
                    "Native chart",
                    "Editable after generation",
                    "native_bar_chart",
                    ["SRC-1"],
                    {
                        "chart": {"categories": ["A", "B"], "series": [{"name": "Value", "values": [1, 2]}]},
                        "insight": {"heading": "Insight", "body": "The chart remains native."},
                    },
                ),
                SlideSpec(
                    "S002",
                    "Native table",
                    "Rows and cells remain editable",
                    "table",
                    ["SRC-2"],
                    {"columns": ["Item", "Value"], "rows": [["A", "10"], ["B", "20"], ["C", "30"]]},
                ),
                SlideSpec(
                    "S003",
                    "Editable process",
                    "Steps are native PowerPoint shapes",
                    "process",
                    ["SRC-3"],
                    {"steps": [{"heading": "Plan", "body": "Define"}, {"heading": "Build", "body": "Compose"}, {"heading": "QA", "body": "Render"}]},
                ),
            ],
        )
        system = compile_template(TEMPLATE)
        plans = plan_deck(deck)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "deck.pptx"
            compose_deck(TEMPLATE, deck, plans, system, output)
            prs = Presentation(output)
            self.assertEqual(len(prs.slides), 3)
            self.assertEqual(sum(1 for shape in prs.slides[0].shapes if getattr(shape, "has_chart", False)), 1)
            chart = next(shape.chart for shape in prs.slides[0].shapes if getattr(shape, "has_chart", False))
            self.assertTrue(chart.plots[0].has_data_labels)
            self.assertTrue(chart.plots[0].data_labels.show_value)
            self.assertEqual(sum(1 for shape in prs.slides[1].shapes if getattr(shape, "has_table", False)), 1)
            table_shape = next(shape for shape in prs.slides[1].shapes if getattr(shape, "has_table", False))
            header_cell = table_shape.table.cell(0, 0)
            body_cell = table_shape.table.cell(1, 0)
            self.assertGreaterEqual(header_cell.text_frame.margin_left, Inches(0.16))
            self.assertGreaterEqual(header_cell.text_frame.margin_right, Inches(0.16))
            self.assertGreaterEqual(header_cell.text_frame.margin_top, Inches(0.08))
            self.assertEqual(header_cell.text_frame.paragraphs[0].alignment, PP_ALIGN.CENTER)
            self.assertEqual(body_cell.text_frame.paragraphs[0].alignment, PP_ALIGN.LEFT)
            self.assertGreaterEqual(sum(1 for shape in prs.slides[2].shapes if "-step-" in shape.name), 3)
            process_names = [str(shape.name) for shape in prs.slides[2].shapes]
            arrow_positions = [index for index, name in enumerate(process_names) if "-arrow-" in name]
            step_positions = [index for index, name in enumerate(process_names) if "-step-" in name]
            self.assertTrue(arrow_positions)
            self.assertLess(max(arrow_positions), min(step_positions))
            report = inspect_pptx(output, plans)
            self.assertEqual(report["status"], "PASS", report["findings"])
            title_shape = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False) and "Native chart" in shape.text)
            title_shape.text_frame.text = "TODO"
            placeholder_output = Path(directory) / "placeholder.pptx"
            prs.save(placeholder_output)
            placeholder_report = inspect_pptx(placeholder_output, plans)
            placeholder_codes = {item["code"] for item in placeholder_report["findings"]}
            self.assertEqual(placeholder_report["status"], "FAIL")
            self.assertIn("leftover_placeholder", placeholder_codes)
            title_shape.text_frame.text = "タイトルを入力してください。"
            japanese_placeholder = Path(directory) / "japanese-placeholder.pptx"
            prs.save(japanese_placeholder)
            japanese_report = inspect_pptx(japanese_placeholder, plans)
            self.assertEqual(japanese_report["status"], "FAIL")
            self.assertIn("leftover_placeholder", {item["code"] for item in japanese_report["findings"]})

            overflow_prs = Presentation(output)
            overflow = overflow_prs.slides[0].shapes.add_textbox(Inches(1), Inches(5), Inches(1.2), Inches(0.25))
            overflow.name = "SHN_CONTENT:S001:overflow-test"
            overflow.text_frame.auto_size = MSO_AUTO_SIZE.NONE
            overflow.text_frame.word_wrap = True
            overflow.text_frame.text = "固定枠には収まらない長い本文テキストです。"
            overflow.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
            overflow_path = Path(directory) / "overflow.pptx"
            overflow_prs.save(overflow_path)
            overflow_report = inspect_pptx(overflow_path, plans)
            self.assertEqual(overflow_report["status"], "FAIL")
            self.assertIn("likely_text_overflow", {item["code"] for item in overflow_report["findings"]})


if __name__ == "__main__":
    unittest.main()
