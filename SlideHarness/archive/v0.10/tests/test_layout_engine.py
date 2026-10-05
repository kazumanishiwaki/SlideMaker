import json, shutil, tempfile, unittest
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
from analyze_ppt import analyze as analyze_ppt
from compose_deck import compose
from layout_engine import Box, _metrics, plan_layouts
from PIL import Image


class LayoutEngineTests(unittest.TestCase):
    def _project(self, td):
        p=Path(td)
        for d in ['inputs','analysis','planning/slide_plans','render/layout_plans','output','state']:
            (p/d).mkdir(parents=True,exist_ok=True)
        shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
        shutil.copy2(ROOT/'examples/test-case/design.md',p/'inputs/design.md')
        analyze_ppt(p/'inputs/source.pptx',p/'analysis')
        (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','order':1,'primary_message':'主要設定を4領域で管理'}]},ensure_ascii=False),encoding='utf-8')
        (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S015','page_number':15,'recommended_mode':'RECOMPOSE'}}]}),encoding='utf-8')
        plan={
          'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S015','adaptation_mode':'RECOMPOSE',
          'composer_support':'layout_engine','content_blocks':[],'source_object_mutations':[],
          'layout_request':{
            'family':'grid','density':'medium_low','title':'管理画面から主要設定を変更','subtitle':'公開後も担当者が更新できます',
            'items':[
              {'number':'01','title':'キャンペーン','body':'商品・店舗・応募方法\n期間・規約・問い合わせ先'},
              {'number':'02','title':'フォーム','body':'追加質問・回答形式\n必須設定・レシート枚数'},
              {'number':'03','title':'ビジュアル','body':'テンプレート・フォント・配色\n画像・OGP・favicon'},
              {'number':'04','title':'メール・計測','body':'自動返信メール・GA4\nGoogle Tag Manager'},
            ],
            'clear_region_object_ids':[f'SRC-S015-O{x:03d}' for x in list(range(3,16))+list(range(19,34))],
            'style_source_object_ids':{
              'title':'SRC-S015-O014','subtitle':'SRC-S015-O015','item_title':'SRC-S015-O022','item_body':'SRC-S015-O020','number':'SRC-S015-O023'
            }
          },
          'status':'APPROVED'
        }
        (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
        return p

    def _statement_plan(self, p: Path, *, fill_mode: str | None = None):
        (p/'planning/deck_plan.json').write_text(json.dumps({
            'slides':[{
                'slide_id':'OUT-S001','order':1,'role':'statement',
                'primary_message':'素材が揃えば、最短即日で公開可能。'
            }]
        },ensure_ascii=False),encoding='utf-8')
        (p/'planning/slide_matching.json').write_text(json.dumps({
            'slides':[{
                'slide_id':'OUT-S001',
                'selected':{'source_slide_id':'SRC-S002','page_number':2,'recommended_mode':'RECOMPOSE'}
            }]
        },ensure_ascii=False),encoding='utf-8')
        plan={
          'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S002','adaptation_mode':'RECOMPOSE',
          'composer_support':'layout_engine','content_blocks':[],'source_object_mutations':[],
          'layout_request':{
            'family':'statement','density':'medium_low',
            'title':'素材が揃えば、最短即日で公開可能。',
            'clear_region_object_ids':['SRC-S002-O004','SRC-S002-O005','SRC-S002-O006'],
            'style_source_object_ids':{'title':'SRC-S002-O005','subtitle':'SRC-S002-O006'}
          },
          'status':'APPROVED'
        }
        if fill_mode:
            plan['fill_policy']={
                'mode':fill_mode,
                'source':'user',
                'reason':['test override'],
                'preserve_empty_zones':fill_mode in {'intentional_whitespace','preserve_source'},
            }
        (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
        return plan

    def test_recompose_grid_fills_canvas_and_preserves_chrome(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            src=Presentation(p/'inputs/source.pptx'); src_slide=src.slides[14]
            chrome=next(s for s in src_slide.shapes if s.shape_id==2)
            chrome_bbox=(chrome.left,chrome.top,chrome.width,chrome.height)
            result=compose(p,slide_ids=['OUT-S001'],output_name='layout.pptx')
            self.assertEqual(result['errors'],0)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['family'],'grid')
            self.assertEqual(lp['quality']['status'],'pass')
            self.assertGreaterEqual(lp['quality']['content_extent_width_ratio'],0.88)
            self.assertGreaterEqual(lp['quality']['content_extent_height_ratio'],0.62)
            self.assertGreaterEqual(len(lp['generated_objects']),10)
            self.assertEqual(len(lp['geometry']['item_boxes']),4)
            ys=sorted({round(b['y'],2) for b in lp['geometry']['item_boxes']})
            self.assertEqual(len(ys),2, 'four items should form a balanced 2x2 grid, not 3+1')

            out=Presentation(result['pptx']); slide=out.slides[0]
            chrome2=next(s for s in slide.shapes if s.shape_id==2)
            self.assertEqual((chrome2.left,chrome2.top,chrome2.width,chrome2.height),chrome_bbox)
            self.assertFalse(any(s.shape_id==3 for s in slide.shapes), 'old content visual should be cleared')
            all_text='\n'.join(s.text for s in slide.shapes if getattr(s,'has_text_frame',False))
            self.assertIn('管理画面から主要設定を変更',all_text)
            self.assertIn('Google Tag Manager',all_text)

    def test_intentional_whitespace_passes_without_content_fill(self):
        canvas=Box(0,0,10,6)
        focus=[Box(0,2,6,2)]
        intentional=_metrics(canvas,focus,'statement','medium_low',{
            'mode':'intentional_whitespace','source':'user','reason':['single strong message'],'preserve_empty_zones':True,
        })
        filled=_metrics(canvas,focus,'statement','medium_low',{
            'mode':'content_filled','source':'user','reason':[],'preserve_empty_zones':False,
        })
        self.assertEqual(intentional['status'],'pass')
        self.assertTrue(intentional['intentional_whitespace'])
        self.assertFalse(intentional['accidental_dead_space'])
        self.assertEqual(filled['status'],'fail')

    def test_tiny_content_island_fails_even_when_whitespace_is_intentional(self):
        canvas=Box(0,0,10,6)
        q=_metrics(canvas,[Box(0.2,0.2,2.0,0.5)],'statement','airy',{
            'mode':'intentional_whitespace','source':'user','reason':[],'preserve_empty_zones':True,
        })
        self.assertEqual(q['status'],'fail')
        self.assertTrue(q['accidental_dead_space'])
        self.assertIn('CONTENT_ISLAND_TOO_SMALL',q['reason_codes'])

    def test_statement_slide_records_fill_policy(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','order':1,'role':'statement','primary_message':'素材が揃えば、最短即日で公開可能。'}]},ensure_ascii=False),encoding='utf-8')
            (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S002','page_number':2,'recommended_mode':'RECOMPOSE'}}]}),encoding='utf-8')
            plan={
              'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S002','adaptation_mode':'RECOMPOSE',
              'composer_support':'layout_engine','content_blocks':[],'source_object_mutations':[],
              'fill_policy':{'mode':'intentional_whitespace','source':'user','reason':['一つの訴求を強調'],'preserve_empty_zones':True},
              'layout_request':{
                'family':'statement','density':'airy','title':'素材が揃えば、最短即日で公開可能。',
                'subtitle':'キャンペーン素材・規約・実施条件が確定している場合。',
                'title_font_pt':52,'subtitle_font_pt':24,
                'clear_region_object_ids':['SRC-S002-O004','SRC-S002-O005','SRC-S002-O006'],
                'style_source_object_ids':{'title':'SRC-S002-O005','subtitle':'SRC-S002-O006'}
              },
              'status':'APPROVED'
            }
            (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
            result=compose(p,slide_ids=['OUT-S001'],output_name='statement.pptx')
            self.assertEqual(result['errors'],0)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['fill_policy']['mode'],'intentional_whitespace')
            self.assertEqual(lp['quality']['check'],'LAYOUT_FIT_TO_INTENT')
            self.assertEqual(lp['quality']['status'],'pass')
            self.assertTrue(lp['quality']['intentional_whitespace'])

    def test_requires_layout_engine_without_request_still_refuses(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            plan_path=p/'planning/slide_plans/OUT-S001.json'
            plan=json.loads(plan_path.read_text())
            plan['composer_support']='requires_layout_engine'; plan.pop('layout_request',None)
            plan_path.write_text(json.dumps(plan),encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'no executable layout_request|cannot be composed'):
                compose(p,slide_ids=['OUT-S001'])

    def test_statement_auto_policy_preserves_intentional_whitespace(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._statement_plan(p)
            result=compose(p,slide_ids=['OUT-S001'],output_name='statement.pptx')
            self.assertEqual(result['errors'],0)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['fill_policy']['mode'],'intentional_whitespace')
            self.assertEqual(lp['fill_policy']['source'],'design_role')
            self.assertEqual(lp['quality']['status'],'pass')
            self.assertTrue(lp['quality']['intentional_whitespace'])
            self.assertFalse(lp['quality']['accidental_dead_space'])
            self.assertLess(lp['quality']['occupied_area_ratio'],0.40)
            # Deliberate empty zones are protected by Pixel Diff; the mutable region is
            # the generated message cluster rather than the whole content canvas.
            self.assertLess(lp['allowed_regions'][0]['w'],lp['content_canvas']['w'])

    def test_sparse_statement_is_rejected_when_forced_content_filled(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._statement_plan(p,fill_mode='content_filled')
            result=compose(p,slide_ids=['OUT-S001'],output_name='statement-filled.pptx')
            self.assertEqual(result['errors'],1)
            report=json.loads((p/'output/statement-filled_compose_report.json').read_text())
            error=report['slides'][0]['errors'][0]['error']
            self.assertIn('fit-to-intent failed',error)
            self.assertIn("'mode': 'content_filled'",error)
            failed_plan=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(failed_plan['quality']['status'],'fail')
            self.assertIn('POLICY_SEMANTIC_MISMATCH',failed_plan['quality']['reason_codes'])


    def test_effective_fill_policy_is_not_logged_in_compact_mode(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._statement_plan(p)
            first=compose(p,slide_ids=['OUT-S001'],output_name='statement-v1.pptx')
            second=compose(p,slide_ids=['OUT-S001'],output_name='statement-v2.pptx')
            self.assertEqual(first['errors'],0)
            self.assertEqual(second['errors'],0)
            trace_path=p/'decisions/decision_trace.jsonl'
            if trace_path.exists():
                rows=[json.loads(line) for line in trace_path.read_text(encoding='utf-8').splitlines() if line.strip()]
                decisions=[r for r in rows if r.get('decision_type')=='effective_fill_policy' and r.get('target',{}).get('id')=='OUT-S001']
                self.assertEqual(decisions,[])

    def test_full_bleed_policy_is_rejected_for_non_visual_family(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._statement_plan(p,fill_mode='full_bleed')
            result=compose(p,slide_ids=['OUT-S001'],output_name='statement-full-bleed.pptx')
            self.assertEqual(result['errors'],1)
            report=json.loads((p/'output/statement-full-bleed_compose_report.json').read_text())
            error=report['slides'][0]['errors'][0]['error']
            self.assertIn('fit-to-intent failed',error)
            self.assertIn('POLICY_SEMANTIC_MISMATCH',error)
            failed_plan=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(failed_plan['quality']['status'],'fail')
            self.assertEqual(failed_plan['fill_policy']['mode'],'full_bleed')


    def _write_single_plan(self, p: Path, *, slide_id: str, source_page: int, role: str, request: dict):
        source_id=f'SRC-S{source_page:03d}'
        (p/'planning/deck_plan.json').write_text(json.dumps({
            'slides':[{'slide_id':slide_id,'order':1,'role':role,'primary_message':request.get('title','')}]
        },ensure_ascii=False),encoding='utf-8')
        (p/'planning/slide_matching.json').write_text(json.dumps({
            'slides':[{'slide_id':slide_id,'selected':{'source_slide_id':source_id,'page_number':source_page,'recommended_mode':'RECOMPOSE'}}]
        },ensure_ascii=False),encoding='utf-8')
        plan={
            'schema_version':'0.1','slide_id':slide_id,'source_slide_id':source_id,
            'adaptation_mode':'RECOMPOSE','composer_support':'layout_engine',
            'content_blocks':[],'source_object_mutations':[],
            'layout_request':request,'status':'APPROVED'
        }
        (p/f'planning/slide_plans/{slide_id}.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
        return plan

    def _seed_layout_image(self, p: Path, *, slide_id: str, recipe_id: str, negative_space: str = 'right'):
        (p/'inputs/assets').mkdir(parents=True,exist_ok=True)
        (p/'render/generated_assets').mkdir(parents=True,exist_ok=True)
        image_path=p/'inputs/assets/layout-photo.png'
        Image.new('RGB',(1600,1000),(218,146,92)).save(image_path)
        planned=plan_layouts(p,slide_ids=[slide_id])
        if planned['errors']:
            raise AssertionError(planned['errors'])
        recipes=json.loads((p/'planning/image_recipes.json').read_text())
        recipe=next(row for row in recipes['recipes'] if row['recipe_id']==recipe_id)
        recipe['composition']['negative_space']=negative_space
        recipe['composition_origin']='user'
        recipe['selected_asset_id']=f'GAST-{recipe_id}-V001'
        recipe['status']='READY'
        recipe['provider']='manual'
        recipes['generated_at']='test'
        manifest={
            'schema_version':'0.1','assets':[{
                'asset_id':f'GAST-{recipe_id}-V001','recipe_id':recipe_id,'slide_id':slide_id,
                'provider':'manual','source_path':'inputs/assets/layout-photo.png',
                'processed_path':'inputs/assets/layout-photo.png',
                'process':{'crop':{'focal_point':[0.5,0.5]}},'status':'READY'
            }],
            'selected_by_recipe':{recipe_id:f'GAST-{recipe_id}-V001'}
        }
        (p/'planning/image_recipes.json').write_text(json.dumps(recipes,ensure_ascii=False),encoding='utf-8')
        (p/'render/generated_assets/manifest.json').write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')

    def test_plan_layouts_seeds_photo_recipe_before_asset_selection(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._write_single_plan(p,slide_id='OUT-S001',source_page=4,role='photo_text',request={
                'family':'photo_text','density':'medium_low','title':'元PPTの写真枠を活かす',
                'body':'写真と本文を内容に合わせて再構成します。','image_recipe_id':'IMG-OUT-S001-01',
                'image_source_object_id':'SRC-S004-O004','image_position':'auto','text_side':'right',
                'reuse_source_image_frame':True,
                'clear_region_object_ids':['SRC-S004-O003','SRC-S004-O005','SRC-S004-O006','SRC-S004-O007','SRC-S004-O011','SRC-S004-O012'],
                'style_source_object_ids':{'title':'SRC-S004-O006','subtitle':'SRC-S004-O007','item_body':'SRC-S004-O011','image':'SRC-S004-O004'}
            })
            result=plan_layouts(p,slide_ids=['OUT-S001'])
            self.assertEqual(result['planned'],1)
            self.assertFalse(result['errors'])
            recipe=json.loads((p/'planning/image_recipes.json').read_text())['recipes'][0]
            self.assertEqual(recipe['status'],'NEEDS_ASSET')
            self.assertEqual(recipe['target']['source_object_id'],'SRC-S004-O004')
            self.assertEqual(recipe['target']['layout_slot_id'],'OUT-S001:primary_visual')
            self.assertEqual(recipe['composition']['negative_space'],'right')

    def test_layout_generated_photo_slot_clears_stale_source_lineage_and_avoids_image_mutation(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            plan=self._write_single_plan(p,slide_id='OUT-S001',source_page=4,role='photo_text',request={
                'family':'photo_text','density':'medium_low','title':'内容に合わせて写真領域を再構成',
                'body':'元ページのタイポグラフィは継承し、画像枠はLayout Engineが新しく設計します。',
                'image_recipe_id':'IMG-OUT-S001-01','image_source_object_id':'SRC-S004-O004',
                'image_position':'left','reuse_source_image_frame':True,
                'clear_region_object_ids':['SRC-S004-O003','SRC-S004-O004','SRC-S004-O005','SRC-S004-O006','SRC-S004-O007','SRC-S004-O011','SRC-S004-O012'],
                'style_source_object_ids':{'title':'SRC-S004-O006','subtitle':'SRC-S004-O007','item_body':'SRC-S004-O011','image':'SRC-S004-O004'}
            })
            first=plan_layouts(p,slide_ids=['OUT-S001'])
            self.assertFalse(first['errors'])
            recipes=json.loads((p/'planning/image_recipes.json').read_text())
            recipe=recipes['recipes'][0]
            self.assertEqual(recipe['target']['source_object_id'],'SRC-S004-O004')
            # Simulate an approved image from the earlier frame-based design.
            recipe['selected_asset_id']='GAST-OLD-V001'; recipe['status']='READY'
            (p/'planning/image_recipes.json').write_text(json.dumps(recipes,ensure_ascii=False),encoding='utf-8')

            plan_path=p/'planning/slide_plans/OUT-S001.json'
            plan=json.loads(plan_path.read_text())
            plan['layout_request']['reuse_source_image_frame']=False
            plan['layout_request'].pop('image_source_object_id',None)
            plan['layout_request']['style_source_object_ids'].pop('image',None)
            plan_path.write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
            second=plan_layouts(p,slide_ids=['OUT-S001'])
            self.assertFalse(second['errors'])
            recipe=json.loads((p/'planning/image_recipes.json').read_text())['recipes'][0]
            self.assertIsNone(recipe['target']['source_object_id'])
            self.assertIsNone(recipe['selected_asset_id'])
            self.assertEqual(recipe['status'],'NEEDS_ASSET')
            self.assertEqual(recipe['treatment']['frame'],'layout_generated_frame')

            self._seed_layout_image(p,slide_id='OUT-S001',recipe_id='IMG-OUT-S001-01',negative_space='right')
            result=compose(p,slide_ids=['OUT-S001'],output_name='layout-generated-photo.pptx')
            self.assertEqual(result['errors'],0)
            final_plan=json.loads(plan_path.read_text())
            self.assertFalse(any(m.get('action')=='replace_image' for m in final_plan['source_object_mutations']))
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            visual=next(x for x in lp['generated_objects'] if x.get('role')=='primary_visual')
            self.assertFalse(visual['source_picture_reused'])
            out=Presentation(result['pptx'])
            self.assertGreaterEqual(sum(1 for shape in out.slides[0].shapes if shape.shape_type==13),1)

    def test_photo_text_uses_recipe_negative_space_and_reuses_picture_shape(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._write_single_plan(p,slide_id='OUT-S001',source_page=4,role='photo_text',request={
                'family':'photo_text','density':'medium_low','title':'スマートフォンで使いやすい画面',
                'body':'必要な情報を読みやすく配置し、操作の流れを迷わせません。',
                'image_recipe_id':'IMG-OUT-S001-01','image_source_object_id':'SRC-S004-O004',
                'image_position':'auto','reuse_source_image_frame':True,
                'clear_region_object_ids':['SRC-S004-O003','SRC-S004-O005','SRC-S004-O006','SRC-S004-O007','SRC-S004-O011','SRC-S004-O012'],
                'style_source_object_ids':{'title':'SRC-S004-O006','subtitle':'SRC-S004-O007','item_body':'SRC-S004-O011','image':'SRC-S004-O004'}
            })
            self._seed_layout_image(p,slide_id='OUT-S001',recipe_id='IMG-OUT-S001-01',negative_space='right')
            result=compose(p,slide_ids=['OUT-S001'],output_name='photo-text.pptx')
            self.assertEqual(result['errors'],0)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['family'],'photo_text')
            self.assertEqual(lp['geometry']['image_position'],'left')
            self.assertEqual(lp['layout_intent']['resolved_text_side'],'right')
            visual=next(x for x in lp['generated_objects'] if x.get('role')=='primary_visual')
            self.assertEqual(visual['output_shape_id'],4)
            self.assertTrue(visual['source_picture_reused'])
            out=Presentation(result['pptx']); slide=out.slides[0]
            picture=next(x for x in slide.shapes if x.shape_id==4)
            self.assertAlmostEqual(picture.left/914400,lp['geometry']['image_box']['x'],places=3)
            self.assertGreater(lp['geometry']['text_box']['x'],lp['geometry']['image_box']['x'])

    def test_full_bleed_reuses_picture_and_keeps_chrome_above(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._write_single_plan(p,slide_id='OUT-S001',source_page=4,role='full_bleed',request={
                'family':'full_bleed','density':'medium_low','title':'ブランドの印象を、全画面で伝える',
                'body':'写真を主役にしながら、元PPTの共通装飾とタイポグラフィを維持します。',
                'image_recipe_id':'IMG-OUT-S001-01','image_source_object_id':'SRC-S004-O004',
                'reuse_source_image_frame':True,'bleed_to_slide':True,'overlay_position':'left',
                'overlay_panel':'solid','overlay_color':'FFFFFF','overlay_opacity':0.86,
                'clear_region_object_ids':['SRC-S004-O003','SRC-S004-O005','SRC-S004-O006','SRC-S004-O007','SRC-S004-O011','SRC-S004-O012'],
                'style_source_object_ids':{'title':'SRC-S004-O006','item_body':'SRC-S004-O011','image':'SRC-S004-O004'}
            })
            self._seed_layout_image(p,slide_id='OUT-S001',recipe_id='IMG-OUT-S001-01',negative_space='left')
            result=compose(p,slide_ids=['OUT-S001'],output_name='full-bleed.pptx')
            self.assertEqual(result['errors'],0)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['fill_policy']['mode'],'full_bleed')
            self.assertEqual(lp['quality']['status'],'pass')
            out=Presentation(result['pptx']); slide=out.slides[0]
            order=[shape.shape_id for shape in slide.shapes]
            self.assertLess(order.index(4),order.index(2),'full-bleed picture must sit behind template chrome')
            visual=next(x for x in lp['generated_objects'] if x.get('role')=='background_visual')
            self.assertEqual(visual['output_shape_id'],4)
            self.assertAlmostEqual(next(x for x in slide.shapes if x.shape_id==4).width/914400,26.6667,places=2)

    def test_two_column_comparison_is_editable_and_balanced(self):
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._write_single_plan(p,slide_id='OUT-S001',source_page=13,role='two_column',request={
                'family':'two_column','density':'medium_low','title':'Liteとフル版の使い分け',
                'subtitle':'必要な範囲に合わせて選べます。','items':[
                    {'title':'Aha! レシート Lite','body':'標準機能ですぐに開始\nシンプルな運用に対応','style':'card'},
                    {'title':'Aha! レシート','body':'OCR・高度なカスタマイズ\n複雑な条件や外部連携に対応','style':'card'}
                ],
                'clear_region_object_ids':[f'SRC-S013-O{x:03d}' for x in range(3,23)],
                'style_source_object_ids':{'title':'SRC-S013-O005','subtitle':'SRC-S013-O006','item_title':'SRC-S013-O018','item_body':'SRC-S013-O019'}
            })
            result=compose(p,slide_ids=['OUT-S001'],output_name='two-column.pptx')
            self.assertEqual(result['errors'],0)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(len(lp['geometry']['item_boxes']),2)
            self.assertAlmostEqual(lp['geometry']['item_boxes'][0]['w'],lp['geometry']['item_boxes'][1]['w'],places=2)
            self.assertTrue(any(x.get('role')=='column_divider' for x in lp['generated_objects']))
            from pptx import Presentation
            out=Presentation(result['pptx']); slide=out.slides[0]
            tinted=[]
            for shape in slide.shapes:
                try:
                    rgb=str(shape.fill.fore_color.rgb) if shape.fill.type else ''
                except Exception:
                    rgb=''
                if rgb=='FFF7F0':
                    tinted.append(shape)
            self.assertEqual(len(tinted),2,'card-style columns should create two editable brand-tinted panels')

    def test_comparison_matrix_is_native_powerpoint_table(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._write_single_plan(p,slide_id='OUT-S001',source_page=11,role='comparison_matrix',request={
                'family':'comparison_matrix','density':'medium_low','title':'主要機能の比較',
                'subtitle':'○：標準対応　△：別作業が必要　×：非対応',
                'matrix':{
                    'row_labels':['フォーム項目','レシート応募','抽選','OCR'],
                    'columns':[
                        {'title':'Googleフォーム','values':['○','△','△','×']},
                        {'title':'Lite','values':['○','○','○','×']},
                        {'title':'フル版','values':['○','○','○','○']}
                    ],'highlight_column':1
                },
                'clear_region_object_ids':[f'SRC-S011-O{x:03d}' for x in range(3,15)],
                'style_source_object_ids':{'title':'SRC-S011-O005','subtitle':'SRC-S011-O006','item_title':'SRC-S011-O012','item_body':'SRC-S011-O011','table':'SRC-S011-O010'}
            })
            result=compose(p,slide_ids=['OUT-S001'],output_name='matrix.pptx')
            self.assertEqual(result['errors'],0)
            out=Presentation(result['pptx']); slide=out.slides[0]
            tables=[x for x in slide.shapes if getattr(x,'has_table',False)]
            self.assertEqual(len(tables),1)
            self.assertEqual(len(tables[0].table.rows),5)
            self.assertEqual(len(tables[0].table.columns),4)
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['quality']['status'],'pass')
            self.assertTrue(next(x for x in lp['generated_objects'] if x.get('role')=='comparison_matrix')['editable'])

    def test_relationship_uses_native_nodes_and_connectors(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=self._project(td)
            self._write_single_plan(p,slide_id='OUT-S001',source_page=13,role='relationship',request={
                'family':'relationship','density':'medium_low','title':'運用を支える機能の関係',
                'nodes':[
                    {'id':'CENTER','title':'管理画面','body':'一元管理'},
                    {'id':'LP','title':'公開','body':'LP・期間・規約'},
                    {'id':'ENTRY','title':'受付','body':'フォーム・画像'},
                    {'id':'CHECK','title':'確認','body':'応募・重複候補'},
                    {'id':'DRAW','title':'抽選','body':'CSV・簡易抽選'}
                ],
                'edges':[{'from':'CENTER','to':x} for x in ['LP','ENTRY','CHECK','DRAW']],
                'center_node_id':'CENTER',
                'clear_region_object_ids':[f'SRC-S013-O{x:03d}' for x in range(3,23)],
                'style_source_object_ids':{'title':'SRC-S013-O005','subtitle':'SRC-S013-O006','item_title':'SRC-S013-O018','item_body':'SRC-S013-O019'}
            })
            result=compose(p,slide_ids=['OUT-S001'],output_name='relationship.pptx')
            self.assertEqual(result['errors'],0)
            out=Presentation(result['pptx']); slide=out.slides[0]
            lp=json.loads((p/'render/layout_plans/OUT-S001.json').read_text())
            self.assertEqual(lp['geometry']['hub_id'],'CENTER')
            edges=[x for x in lp['generated_objects'] if x.get('role')=='relationship_edge']
            nodes=[x for x in lp['generated_objects'] if x.get('role')=='relationship_node']
            self.assertEqual(len(edges),4)
            self.assertEqual(len(nodes),5)
            self.assertGreaterEqual(sum(1 for x in slide.shapes if x.shape_type==9),4)


if __name__=='__main__': unittest.main()
