import json, shutil, tempfile, unittest
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
from analyze_ppt import analyze
from compose_deck import compose
from validate_output import validate
from render_source import render_pptx




class PixelDiffMaskUnitTests(unittest.TestCase):
    def test_removed_text_mask_includes_render_overflow_halo(self):
        from PIL import Image, ImageDraw
        from validate_output import _pixel_diff_outside_mutations

        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            source=root/'source.png'; output=root/'output.png'
            # 10in × 5.625in at 96px/in. The declared text box ends at
            # y=2.20in, while the source glyph pixels extend another 0.16in.
            # Clearing the declared text must not be reported as an unrelated
            # mutation; a real change farther away still would be protected.
            src=Image.new('RGB',(960,540),'white')
            out=Image.new('RGB',(960,540),'white')
            draw=ImageDraw.Draw(src)
            draw.rectangle((96,192,390,227),fill='black')
            src.save(source); out.save(output)
            srcslide={'objects':[{'source_object_id':'SRC-S001-O001','text':'overflowing source text','bbox':{'x':1.0,'y':2.0,'w':3.0,'h':0.2}}]}
            layout={'removed_source_object_ids':['SRC-S001-O001'],'allowed_regions':[]}
            result=_pixel_diff_outside_mutations(
                source,output,srcslide,{},10.0,5.625,layout_plan=layout
            )
            self.assertTrue(result['available'])
            self.assertEqual(result['changed_pixels'],0)

@unittest.skipUnless(shutil.which('libreoffice') and shutil.which('pdftoppm') and shutil.which('pdftotext'), 'render QA dependencies unavailable')
class ValidateOutputTests(unittest.TestCase):
    def test_intentional_whitespace_uses_fit_to_intent_qa(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)
            for d in ['inputs','analysis','planning/slide_plans','output','render/source_renders','render/layout_plans','qa','state']:
                (p/d).mkdir(parents=True,exist_ok=True)
            shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
            shutil.copy2(ROOT/'examples/test-case/design.md',p/'inputs/design.md')
            analyze(p/'inputs/source.pptx',p/'analysis')
            rendered = render_pptx(p/'inputs/source.pptx', p/'render/source_renders')
            self.assertGreaterEqual(len(rendered), 2)
            (p/'planning/deck_plan.json').write_text(json.dumps({
                'slides':[{'slide_id':'OUT-S001','order':1,'role':'statement','primary_message':'素材が揃えば、最短即日で公開可能。'}]
            },ensure_ascii=False),encoding='utf-8')
            (p/'planning/slide_matching.json').write_text(json.dumps({
                'slides':[{'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S002','page_number':2,'recommended_mode':'RECOMPOSE'}}]
            },ensure_ascii=False),encoding='utf-8')
            plan={
                'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S002',
                'adaptation_mode':'RECOMPOSE','composer_support':'layout_engine',
                'content_blocks':[],'source_object_mutations':[],
                'fill_policy':{
                    'mode':'intentional_whitespace','source':'user',
                    'reason':['一つの訴求を強調'],'preserve_empty_zones':True,
                },
                'layout_request':{
                    'family':'statement','density':'airy',
                    'title':'素材が揃えば、最短即日で公開可能。',
                    'subtitle':'キャンペーン素材・規約・実施条件が確定している場合。',
                    'title_font_pt':52,'subtitle_font_pt':24,
                    'clear_region_object_ids':['SRC-S002-O004','SRC-S002-O005','SRC-S002-O006'],
                    'style_source_object_ids':{'title':'SRC-S002-O005','subtitle':'SRC-S002-O006'},
                },
                'status':'APPROVED',
            }
            (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
            result=compose(p,slide_ids=['OUT-S001'],output_name='statement.pptx')
            self.assertEqual(result['errors'],0)
            qa=validate(p,'output/statement.pptx',kind='test')
            fit=[c for c in qa['checks'] if c['check_id']=='LAYOUT_FIT_TO_INTENT']
            self.assertEqual(len(fit),1)
            self.assertEqual(fit[0]['result'],'pass')
            self.assertEqual(fit[0]['fill_policy'],'intentional_whitespace')
            self.assertTrue(fit[0]['intentional_whitespace'])
            self.assertFalse(fit[0]['accidental_dead_space'])
            self.assertFalse(any(c['check_id']=='LAYOUT_CANVAS_FILL' for c in qa['checks']))
            self.assertEqual(qa['status'],'pass')

    def test_pixel_diff_rejects_unplanned_mutation(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)
            for d in ['inputs','analysis','planning/slide_plans','output','render/source_renders','qa','state']:
                (p/d).mkdir(parents=True,exist_ok=True)
            shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
            analyze(p/'inputs/source.pptx',p/'analysis')
            rendered = render_pptx(p/'inputs/source.pptx', p/'render/source_renders')
            self.assertGreaterEqual(len(rendered), 2)

            src=Presentation(p/'inputs/source.pptx'); slide=src.slides[1]
            text_shapes=[s for s in slide.shapes if getattr(s,'has_text_frame',False) and s.text.strip()]
            title=max(text_shapes,key=lambda s:len(s.text))
            other=next(s for s in text_shapes if s.shape_id != title.shape_id)
            (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','order':1,'primary_message':'新しいメッセージ'}]}),encoding='utf-8')
            (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S002','page_number':2,'recommended_mode':'DIRECT'}}]}),encoding='utf-8')
            plan={
                'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S002',
                'adaptation_mode':'DIRECT','composer_support':'supported',
                'content_blocks':[{'id':'CB-01','role':'title','source_text':title.text,'semantic_text':'新しいメッセージ','render_text':'新しいメッセージ'}],
                'source_object_mutations':[
                    {'source_object_id':f'SRC-S002-O{title.shape_id:03d}','action':'replace_text','content_block_id':'CB-01','render_text':'新しいメッセージ','reason':'planned'},
                    {'source_object_id':f'SRC-S002-O{other.shape_id:03d}','action':'replace_text','content_block_id':None,'render_text':'UNAUTHORIZED CHANGE','reason':'temporarily composed for tamper test'}
                ],
                'status':'APPROVED'
            }
            (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
            result=compose(p,slide_ids=['OUT-S001'],output_name='one.pptx')
            # Simulate an out-of-contract mutation: output contains the second change,
            # but the effective Slide Plan no longer authorizes that object region.
            effective=json.loads((p/'planning/slide_plans/OUT-S001.json').read_text())
            effective['source_object_mutations']=effective['source_object_mutations'][:1]
            (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(effective,ensure_ascii=False),encoding='utf-8')

            qa=validate(p,'output/one.pptx',kind='test')
            pixel=[c for c in qa['checks'] if c['check_id']=='UNEXPECTED_PIXEL_DIFF']
            self.assertEqual(len(pixel),1)
            self.assertEqual(pixel[0]['result'],'fail')
            self.assertEqual(pixel[0]['severity'],'blocking')
            self.assertGreater(pixel[0]['changed_pixels'],50)
            self.assertEqual(qa['status'],'fail')


if __name__=='__main__': unittest.main()
