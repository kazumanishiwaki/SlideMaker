import json
import tempfile
import unittest
from pathlib import Path
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'engine'))
from analyze_ppt import analyze


def rename_picture(shape, name: str) -> None:
    shape._element.nvPicPr.cNvPr.set('name', name)


class AnalyzePptClassificationTests(unittest.TestCase):
    def test_context_zorder_and_role_hints_are_recorded(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            logo = root / 'logo.png'
            photo = root / 'photo.png'
            icon = root / 'icon.png'
            Image.new('RGBA', (120, 40), (242, 125, 48, 255)).save(logo)
            Image.new('RGB', (900, 600), (115, 155, 190)).save(photo)
            im = Image.new('RGBA', (160, 160), (0, 0, 0, 0))
            draw = ImageDraw.Draw(im); draw.ellipse((20, 20, 140, 140), fill=(242, 125, 48, 255))
            im.save(icon)

            prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
            for page in range(3):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                logo_shape = slide.shapes.add_picture(str(logo), Inches(11.7), Inches(0.25), Inches(1.0), Inches(0.33))
                rename_picture(logo_shape, 'Brand Logo')
                if page == 0:
                    slide.shapes.add_textbox(Inches(0.9), Inches(1.0), Inches(4.4), Inches(0.8)).text = '導入担当者が画面を確認'
                    photo_shape = slide.shapes.add_picture(str(photo), Inches(0.9), Inches(2.0), Inches(5.5), Inches(3.6))
                    rename_picture(photo_shape, 'Campaign Photo')
                    icon_shape = slide.shapes.add_picture(str(icon), Inches(7.2), Inches(2.2), Inches(0.8), Inches(0.8))
                    rename_picture(icon_shape, 'Feature Icon')
                    slide.shapes.add_textbox(Inches(8.1), Inches(2.25), Inches(3.2), Inches(0.6)).text = '管理画面で変更可能'
            pptx = root / 'fixture.pptx'; prs.save(pptx)
            result = analyze(pptx, root / 'analysis')
            assets = json.loads((root / 'analysis/source_assets.json').read_text())['assets']
            logos = [row for row in assets if row.get('shape_name') == 'Brand Logo']
            self.assertEqual(len(logos), 3)
            self.assertTrue(all(row['classification'] == 'template_chrome' for row in logos))
            photo_row = next(row for row in assets if row.get('shape_name') == 'Campaign Photo')
            self.assertEqual(photo_row['classification'], 'content_visual')
            self.assertEqual(photo_row['content_role'], 'photo')
            self.assertIsInstance(photo_row.get('z_index'), int)
            self.assertIn('nearest_text_distance', photo_row.get('context', {}))
            icon_row = next(row for row in assets if row.get('shape_name') == 'Feature Icon')
            self.assertEqual(icon_row['content_role'], 'icon')
            self.assertIn(icon_row['classification'], {'content_visual', 'unknown'})
            self.assertEqual(result['source_assets']['summary']['template_chrome'], 3)

    def test_geometry_repetition_and_default_picture_names_are_classified(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
            for page in range(3):
                corner = root / f'corner-{page}.png'
                Image.new('RGBA', (100, 60), (220 - page * 10, 80 + page * 10, 90, 255)).save(corner)
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                # Different hashes, same geometry: recurring corner decoration.
                slide.shapes.add_picture(str(corner), Inches(12.0), Inches(0.25), Inches(0.8), Inches(0.48))
                if page == 0:
                    icon = root / 'default-icon.png'
                    im = Image.new('RGBA', (120, 120), (0, 0, 0, 0))
                    draw = ImageDraw.Draw(im); draw.rectangle((24, 24, 96, 96), fill=(230, 120, 55, 255))
                    im.save(icon)
                    slide.shapes.add_picture(str(icon), Inches(1.0), Inches(2.0), Inches(0.72), Inches(0.72))
                    slide.shapes.add_textbox(Inches(1.78), Inches(2.02), Inches(3.0), Inches(0.6)).text = '本文に対応する小さなアイコン'
                    rule = root / 'rule.png'
                    Image.new('RGBA', (800, 8), (230, 120, 55, 255)).save(rule)
                    slide.shapes.add_picture(str(rule), Inches(1.0), Inches(4.0), Inches(5.0), Inches(0.05))
            pptx = root / 'geometry-fixture.pptx'; prs.save(pptx)
            analyze(pptx, root / 'analysis')
            assets = json.loads((root / 'analysis/source_assets.json').read_text())['assets']
            corner_rows = [row for row in assets if row.get('geometry_repeat_count') == 3]
            self.assertEqual(len(corner_rows), 3)
            self.assertTrue(all(row['classification'] == 'template_chrome' for row in corner_rows))
            icon_row = next(row for row in assets if row.get('content_role') == 'icon')
            self.assertEqual(icon_row['classification'], 'content_visual')
            thin_row = next(row for row in assets if 'THIN_DECORATIVE_RASTER' in row.get('reason_codes', []))
            self.assertEqual(thin_row['classification'], 'template_chrome')


if __name__ == '__main__':
    unittest.main()
