import json
import shutil
import tempfile
import unittest
from pathlib import Path
import sys

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'engine'))
from analyze_ppt import analyze
from compose_deck import compose


class LayoutFamiliesV090Tests(unittest.TestCase):
    def _project(self, td: str) -> Path:
        p = Path(td)
        for directory in [
            'inputs', 'analysis', 'planning/slide_plans', 'render/layout_plans',
            'output', 'state', 'decisions', 'versions', 'qa'
        ]:
            (p / directory).mkdir(parents=True, exist_ok=True)
        shutil.copy2(ROOT / 'examples/test-case/inputs/source.pptx', p / 'inputs/source.pptx')
        shutil.copy2(ROOT / 'examples/test-case/design.md', p / 'inputs/design.md')
        analyze(p / 'inputs/source.pptx', p / 'analysis')
        return p

    def _clear_ids(self, p: Path, source_id: str) -> list[str]:
        slides = json.loads((p / 'analysis/source_slides.json').read_text())['slides']
        assets = json.loads((p / 'analysis/source_assets.json').read_text())['assets']
        protected = {
            row['source_object_id'] for row in assets
            if row['source_slide_id'] == source_id and row['classification'] == 'template_chrome'
        }
        slide = next(row for row in slides if row['id'] == source_id)
        return [row['source_object_id'] for row in slide['objects'] if row['source_object_id'] not in protected]

    def _compose_family(self, p: Path, family: str, request: dict, *, source_page: int = 15, role: str | None = None):
        sid = 'OUT-S001'; source_id = f'SRC-S{source_page:03d}'
        request = {
            'family': family,
            'density': 'medium_low',
            'clear_region_object_ids': self._clear_ids(p, source_id),
            'style_source_object_ids': {
                'title': f'{source_id}-O014',
                'subtitle': f'{source_id}-O015',
                'item_title': f'{source_id}-O022',
                'item_body': f'{source_id}-O020',
                'number': f'{source_id}-O023',
            },
            **request,
        }
        (p / 'planning/deck_plan.json').write_text(json.dumps({
            'slides': [{'slide_id': sid, 'order': 1, 'role': role or family, 'primary_message': request.get('title', '')}]
        }, ensure_ascii=False), encoding='utf-8')
        (p / 'planning/slide_matching.json').write_text(json.dumps({
            'slides': [{'slide_id': sid, 'selected': {'source_slide_id': source_id, 'page_number': source_page, 'recommended_mode': 'RECOMPOSE'}}]
        }, ensure_ascii=False), encoding='utf-8')
        plan = {
            'schema_version': '0.1', 'slide_id': sid, 'source_slide_id': source_id,
            'adaptation_mode': 'RECOMPOSE', 'composer_support': 'layout_engine',
            'content_blocks': [], 'source_object_mutations': [],
            'layout_request': request, 'status': 'APPROVED',
        }
        (p / f'planning/slide_plans/{sid}.json').write_text(json.dumps(plan, ensure_ascii=False), encoding='utf-8')
        result = compose(p, slide_ids=[sid], output_name=f'{family}.pptx')
        self.assertEqual(result['errors'], 0, result)
        layout = json.loads((p / f'render/layout_plans/{sid}.json').read_text())
        self.assertEqual(layout['quality']['status'], 'pass')
        return Presentation(p / f'output/{family}.pptx'), layout

    def test_timeline_kpi_and_chart_are_editable(self):
        with tempfile.TemporaryDirectory() as td:
            p = self._project(td)
            timeline, plan = self._compose_family(p, 'timeline', {
                'title': '導入までの流れ',
                'items': [
                    {'date': 'STEP 1', 'title': '素材確認', 'body': '規約と素材を確認'},
                    {'date': 'STEP 2', 'title': '初期設定', 'body': '管理画面を設定'},
                    {'date': 'STEP 3', 'title': 'プレビュー', 'body': '画面とメールを確認'},
                    {'date': 'STEP 4', 'title': '公開', 'body': 'キャンペーン開始'},
                ],
            }, role='timeline')
            self.assertGreaterEqual(sum(1 for shape in timeline.slides[0].shapes if shape.shape_type is not None), 10)
            self.assertEqual(plan['geometry']['orientation'], 'horizontal')

        with tempfile.TemporaryDirectory() as td:
            p = self._project(td)
            kpi, _plan = self._compose_family(p, 'kpi', {
                'title': '運用状況を一目で把握',
                'items': [
                    {'label': '総応募数', 'value': '12,480件', 'note': 'キャンペーン累計'},
                    {'label': '今日の応募', 'value': '386件', 'note': '前日比 +12%'},
                    {'label': '重複候補', 'value': '24件', 'note': '確認対象'},
                ],
            }, role='kpi')
            texts = '\n'.join(shape.text for shape in kpi.slides[0].shapes if getattr(shape, 'has_text_frame', False))
            self.assertIn('12,480件', texts)
            self.assertIn('重複候補', texts)

        with tempfile.TemporaryDirectory() as td:
            p = self._project(td)
            chart, _plan = self._compose_family(p, 'chart', {
                'title': '応募数は公開後に伸長',
                'categories': ['1週目', '2週目', '3週目', '4週目'],
                'series': [{'name': '応募数', 'values': [120, 260, 430, 690]}],
                'chart_type': 'line',
                'note': 'テストデータ',
                'has_legend': False,
            }, role='chart')
            charts = [shape.chart for shape in chart.slides[0].shapes if getattr(shape, 'has_chart', False)]
            self.assertEqual(len(charts), 1)
            self.assertEqual(list(charts[0].series[0].values), [120.0, 260.0, 430.0, 690.0])

    def test_cover_section_and_closing_preserve_deliberate_whitespace(self):
        cases = [
            ('cover', {'title': 'レシートキャンペーンを、ひとつに。', 'subtitle': 'Aha! レシート Lite', 'body': '公開準備から抽選まで'}, 'cover'),
            ('section', {'kicker': '02', 'title': '一つの管理画面で運用', 'subtitle': '準備から終了後まで流れを見失いません'}, 'section'),
            ('closing', {'title': '導入までのご相談', 'subtitle': 'キャンペーン内容と開始希望日をお知らせください。', 'items': [
                {'title': '30分デモ', 'body': '管理画面をご案内'},
                {'title': '見積依頼', 'body': '実施条件を整理'},
                {'title': 'プラン相談', 'body': 'Lite／フル版を比較'},
            ]}, 'closing'),
        ]
        for family, request, role in cases:
            with self.subTest(family=family), tempfile.TemporaryDirectory() as td:
                p = self._project(td)
                prs, layout = self._compose_family(p, family, request, source_page=2 if family != 'closing' else 17, role=role)
                self.assertEqual(len(prs.slides), 1)
                self.assertFalse(layout['quality']['accidental_dead_space'])
                self.assertIn(layout['fill_policy']['mode'], {'intentional_whitespace', 'balanced'})


if __name__ == '__main__':
    unittest.main()
