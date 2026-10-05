import shutil
import tempfile
import unittest
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'engine'))
from render_source import render_pptx


@unittest.skipUnless(
    shutil.which('libreoffice') and (shutil.which('pdftocairo') or shutil.which('pdftoppm')),
    'render dependencies unavailable',
)
class RenderSourceTests(unittest.TestCase):
    def test_repeated_render_in_same_process_is_ordered_and_does_not_hang(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            pptx = root / 'twelve-pages.pptx'
            prs = Presentation()
            blank = prs.slide_layouts[6]
            # More than nine pages verifies that page 10 is not lexically sorted
            # between page 1 and page 2.
            for page in range(1, 13):
                slide = prs.slides.add_slide(blank)
                box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
                run = box.text_frame.paragraphs[0].add_run()
                run.text = f'PAGE {page}'
                run.font.size = Pt(28)
            prs.save(pptx)

            for attempt in range(2):
                files = render_pptx(pptx, root / f'render-{attempt}')
                self.assertEqual(len(files), 12)
                self.assertEqual([Path(f).name for f in files], [f'SRC-S{i:03d}.png' for i in range(1, 13)])
                self.assertTrue(all(Path(f).stat().st_size > 0 for f in files))


if __name__ == '__main__':
    unittest.main()
