import json, shutil, tempfile, unittest
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
from build_skeleton import build

class SkeletonTests(unittest.TestCase):
    def test_builds_editable_clone_deck(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=Path(td); (p/'inputs').mkdir(); (p/'planning').mkdir(); (p/'output').mkdir()
            shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
            (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[
                {'slide_id':'OUT-S001','order':1,'primary_message':'A'},
                {'slide_id':'OUT-S002','order':2,'primary_message':'B'},
                {'slide_id':'OUT-S003','order':3,'primary_message':'C'},
            ]}),encoding='utf-8')
            (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[
                {'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S002','page_number':2,'recommended_mode':'DIRECT'}},
                {'slide_id':'OUT-S002','selected':{'source_slide_id':'SRC-S006','page_number':6,'recommended_mode':'ADAPT'}},
                {'slide_id':'OUT-S003','selected':{'source_slide_id':'SRC-S014','page_number':14,'recommended_mode':'RECOMPOSE'}},
            ]}),encoding='utf-8')
            result=build(p)
            self.assertEqual(result['slide_count'],3)
            out=Presentation(result['pptx'])
            self.assertEqual(len(out.slides),3)
            self.assertGreater(len(out.slides[0].shapes),0)
            source=Presentation(p/'inputs/source.pptx')
            self.assertIsNotNone(out.slides[0].element.cSld.bg)
            self.assertEqual(
                out.slides[0].element.cSld.bg.xml,
                source.slides[1].element.cSld.bg.xml,
                'slide-specific background must survive source slide cloning',
            )
            with open(result['map'],encoding='utf-8') as fh:
                mapping=json.load(fh)
            self.assertEqual(mapping['slides'][1]['source_page_number'],6)

if __name__=='__main__': unittest.main()
