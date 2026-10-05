import hashlib
import json
import tempfile
import shutil
import unittest
from unittest.mock import patch
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'engine'))

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

from analyze_ppt import analyze
from compose_deck import compose
from design_continuity import build_design_system, build_continuity_plans
from image_pipeline import _palette_treatment, generate_images, update_image_recipe, validate_image_harmony, validate_deck_image_coherence
from state import default_project_state, ensure_slide, load_state, save_state
from render_source import render_pptx
from validate_output import validate


class ImagePipelineTests(unittest.TestCase):
    def test_low_key_brand_render_uses_duotone_palette_treatment(self):
        image = Image.new('RGB', (800, 500), '#D8E2DC')
        draw = ImageDraw.Draw(image)
        draw.rectangle((420, 40, 760, 460), fill='#424B4C')
        treated, info = _palette_treatment(
            image,
            {'brightness': .18, 'saturation': .62},
            {'color_roles': {'primary_accent': '500000', 'background_dark': '341212'}},
        )
        self.assertEqual(info['mode'], 'deck_duotone')
        sample = treated.resize((1, 1)).getpixel((0, 0))
        self.assertGreater(sample[0], sample[1])
        self.assertGreater(sample[0], sample[2])

    def _make_project(self, td: str, *, include_asset: bool = True) -> Path:
        root = Path(td)
        for rel in [
            'inputs/assets', 'analysis', 'planning/slide_plans', 'render/source_renders',
            'render/generated_assets', 'output', 'qa', 'state', 'decisions'
        ]:
            (root / rel).mkdir(parents=True, exist_ok=True)

        source_visual = root / 'source-visual.png'
        im = Image.new('RGB', (1600, 900), '#E8E0D2')
        draw = ImageDraw.Draw(im)
        draw.rectangle((80, 80, 720, 820), fill='#E96531')
        draw.ellipse((900, 160, 1480, 740), fill='#242724')
        im.save(source_visual)

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.8), Inches(.45), Inches(7.4), Inches(.7))
        title.text = '元PPTの写真レイアウト'
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        picture = slide.shapes.add_picture(str(source_visual), Inches(6.9), Inches(1.4), width=Inches(5.5), height=Inches(4.8))
        picture.rotation = 7.0
        source_ppt = root / 'inputs/source.pptx'
        prs.save(source_ppt)

        if include_asset:
            uploaded = Image.new('RGB', (1800, 1200), '#D5E2DB')
            draw = ImageDraw.Draw(uploaded)
            draw.rounded_rectangle((180, 180, 930, 1040), radius=90, fill='#E96531')
            draw.ellipse((1050, 230, 1680, 970), fill='#303431')
            uploaded_path = root / 'inputs/assets/キャンペーン_スマートフォン_運用.png'
            uploaded.save(uploaded_path)

        analyze(source_ppt, root / 'analysis')
        (root / 'project.json').write_text(json.dumps({
            'schema_version': '0.1',
            'project_id': 'image-test',
            'preferences': {'source_continuity': 'balanced'},
        }, ensure_ascii=False), encoding='utf-8')
        build_design_system(root)

        (root / 'planning/deck_plan.json').write_text(json.dumps({
            'slides': [{
                'slide_id': 'OUT-S001', 'order': 1, 'role': 'overview',
                'primary_message': 'スマートフォンで運用しやすいキャンペーン画面',
            }]
        }, ensure_ascii=False), encoding='utf-8')
        (root / 'planning/visual_plan.json').write_text(json.dumps({
            'slides': [{
                'slide_id': 'OUT-S001',
                'visual_intent': {'selected_type': 'photo', 'variant': 'editorial'},
            }]
        }, ensure_ascii=False), encoding='utf-8')
        (root / 'planning/slide_matching.json').write_text(json.dumps({
            'slides': [{
                'slide_id': 'OUT-S001',
                'selected': {'source_slide_id': 'SRC-S001', 'page_number': 1, 'recommended_mode': 'ADAPT'},
            }]
        }, ensure_ascii=False), encoding='utf-8')
        plan = {
            'schema_version': '0.1', 'slide_id': 'OUT-S001', 'source_slide_id': 'SRC-S001',
            'adaptation_mode': 'ADAPT', 'composer_support': 'supported',
            'content_blocks': [], 'source_object_mutations': [], 'status': 'APPROVED',
        }
        (root / 'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan), encoding='utf-8')
        build_continuity_plans(root)

        state = default_project_state()
        ss = ensure_slide(state, 'OUT-S001')
        for phase in ('deck_plan', 'visual_plan', 'source_match', 'slide_plan'):
            ss[phase] = 'APPROVED'
        save_state(root, state)
        return root

    def test_auto_selects_local_asset_grades_and_preserves_picture_shape(self):
        with tempfile.TemporaryDirectory() as td:
            root = self._make_project(td)
            recipes = json.loads((root / 'planning/image_recipes.json').read_text())['recipes']
            self.assertEqual(len(recipes), 1)
            recipe = recipes[0]
            self.assertTrue(recipe['target']['source_object_id'])

            result = generate_images(root, provider='auto', recipe_ids=[recipe['recipe_id']])
            self.assertEqual(result['results'][0]['status'], 'READY')
            asset_id = result['results'][0]['asset_id']
            asset = next(x for x in result['manifest']['assets'] if x['asset_id'] == asset_id)
            self.assertEqual(asset['provider'], 'local')
            self.assertTrue((root / asset['processed_path']).exists())
            self.assertGreater(asset['provider_meta']['candidate_ranking'][0]['score']['total'], 0.45)

            updated_recipe = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            self.assertEqual(updated_recipe['selected_asset_id'], asset_id)
            plan = json.loads((root / 'planning/slide_plans/OUT-S001.json').read_text())
            mutation = next(m for m in plan['source_object_mutations'] if m['action'] == 'replace_image')
            self.assertEqual(mutation['fit_mode'], 'preserve_format')
            self.assertEqual(mutation['source_object_id'], recipe['target']['source_object_id'])

            qa = validate_image_harmony(root)
            self.assertEqual(qa['blocking_count'], 0)
            self.assertTrue(any(c['check_id'] == 'IMAGE_TEMPLATE_HARMONY' for c in qa['checks']))
            coherence = validate_deck_image_coherence(root)
            self.assertEqual(coherence['status'], 'pass')

            source_prs = Presentation(root / 'inputs/source.pptx')
            source_picture = next(s for s in source_prs.slides[0].shapes if s.shape_type == 13)
            source_geometry = (source_picture.shape_id, source_picture.left, source_picture.top, source_picture.width, source_picture.height, source_picture.rotation, source_picture._pic.spPr.xml)

            output = compose(root, slide_ids=['OUT-S001'], output_name='image-output.pptx')
            self.assertEqual(output['errors'], 0)
            out = Presentation(output['pptx'])
            pictures = [s for s in out.slides[0].shapes if s.shape_type == 13]
            self.assertEqual(len(pictures), 1)
            expected_hash = hashlib.sha256((root / asset['processed_path']).read_bytes()).hexdigest()
            actual_hash = hashlib.sha256(pictures[0].image.blob).hexdigest()
            self.assertEqual(actual_hash, expected_hash)
            self.assertEqual(pictures[0].shape_id, int(recipe['target']['source_object_id'].split('-O')[-1]))
            output_geometry = (pictures[0].shape_id, pictures[0].left, pictures[0].top, pictures[0].width, pictures[0].height, pictures[0].rotation, pictures[0]._pic.spPr.xml)
            self.assertEqual(output_geometry, source_geometry, 'picture shape geometry, rotation, mask and z-order identity must remain from the source PPT')

            # Refreshing higher-level continuity must not erase a still-compatible user/AI selection.
            build_continuity_plans(root)
            refreshed = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            self.assertEqual(refreshed['selected_asset_id'], asset_id)


    @unittest.skipUnless((shutil.which('libreoffice') or shutil.which('soffice')) and (shutil.which('pdftocairo') or shutil.which('pdftoppm')), 'render toolchain unavailable')
    def test_rendered_output_runs_image_harmony_and_mutation_scope_qa(self):
        with tempfile.TemporaryDirectory() as td:
            root = self._make_project(td)
            recipe = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            generate_images(root, provider='auto', recipe_ids=[recipe['recipe_id']])
            source_renders = render_pptx(root / 'inputs/source.pptx', root / 'render/source_renders')
            self.assertEqual(len(source_renders), 1)
            output = compose(root, slide_ids=['OUT-S001'], output_name='image-render-qa.pptx')
            qa = validate(root, 'output/image-render-qa.pptx', kind='image-render')
            rendered = next(c for c in qa['checks'] if c['check_id'] == 'IMAGE_RENDERED_IN_TARGET_SLOT')
            self.assertEqual(rendered['result'], 'pass')
            self.assertEqual(qa['blocking_count'], 0)

    def test_auto_without_candidates_or_api_requires_review(self):
        with tempfile.TemporaryDirectory() as td, patch.dict('os.environ', {'OPENAI_API_KEY': ''}, clear=False):
            root = self._make_project(td, include_asset=False)
            recipe = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            result = generate_images(root, provider='auto', recipe_ids=[recipe['recipe_id']], allow_mock=False)
            self.assertEqual(result['results'][0]['status'], 'REVIEW_REQUIRED')
            refreshed = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            self.assertIsNone(refreshed['selected_asset_id'])
            self.assertEqual(refreshed['status'], 'REVIEW_REQUIRED')

    def test_mock_fallback_is_explicit_and_warned(self):
        with tempfile.TemporaryDirectory() as td, patch.dict('os.environ', {'OPENAI_API_KEY': '', 'IMAGE_OFFLINE_MOCK': ''}, clear=False):
            root = self._make_project(td, include_asset=False)
            recipe = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            result = generate_images(root, provider='auto', recipe_ids=[recipe['recipe_id']], allow_mock=True)
            asset = result['manifest']['assets'][-1]
            self.assertEqual(asset['provider'], 'mock')
            qa = validate_image_harmony(root)
            placeholder = next(c for c in qa['checks'] if c['check_id'] == 'IMAGE_PROVIDER_PLACEHOLDER')
            self.assertEqual(placeholder['result'], 'fail')
            self.assertEqual(placeholder['severity'], 'warning')

    def test_recipe_override_clears_selection_and_only_invalidates_image_descendants(self):
        with tempfile.TemporaryDirectory() as td:
            root = self._make_project(td)
            recipe = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            generated = generate_images(root, provider='auto', recipe_ids=[recipe['recipe_id']])
            asset_id = generated['results'][0]['asset_id']
            state_before = load_state(root)
            self.assertEqual(state_before['slides']['OUT-S001']['image'], 'APPROVED')
            result = update_image_recipe(root, recipe['recipe_id'], prompt_seed='担当者がスマートフォンで応募状況を確認する場面', negative_space='left')
            self.assertIn('prompt_seed', result['changed'])
            self.assertEqual(result['previous_asset_id'], asset_id)
            updated = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            self.assertIsNone(updated['selected_asset_id'])
            self.assertEqual(updated['status'], 'NEEDS_ASSET')
            plan = json.loads((root / 'planning/slide_plans/OUT-S001.json').read_text())
            self.assertFalse(any(m.get('image_recipe_id') == recipe['recipe_id'] for m in plan['source_object_mutations']))
            state_after = load_state(root)['slides']['OUT-S001']
            self.assertEqual(state_after['deck_plan'], 'APPROVED')
            self.assertEqual(state_after['slide_plan'], 'APPROVED')
            self.assertEqual(state_after['image'], 'DRAFT')
            self.assertEqual(state_after['render'], 'STALE')
            self.assertEqual(state_after['qa'], 'STALE')

    def test_negative_space_change_invalidates_photo_layout_but_prompt_change_does_not(self):
        with tempfile.TemporaryDirectory() as td:
            root = self._make_project(td)
            plan_path = root / 'planning/slide_plans/OUT-S001.json'
            plan = json.loads(plan_path.read_text())
            plan['adaptation_mode'] = 'RECOMPOSE'
            plan['composer_support'] = 'layout_engine'
            plan['layout_request'] = {
                'family': 'photo_text',
                'title': 'スマートフォンで運用しやすいキャンペーン画面',
                'image_recipe_id': 'IMG-OUT-S001-01',
            }
            plan_path.write_text(json.dumps(plan), encoding='utf-8')

            recipe = json.loads((root / 'planning/image_recipes.json').read_text())['recipes'][0]
            state = load_state(root)
            state['slides']['OUT-S001']['layout'] = 'APPROVED'
            save_state(root, state)
            first = update_image_recipe(root, recipe['recipe_id'], prompt_seed='別の運用シーンを使う')
            self.assertFalse(first['layout_invalidated'])
            self.assertEqual(load_state(root)['slides']['OUT-S001']['layout'], 'APPROVED')

            # Only composition-side changes propagate back into photo-aware geometry.
            current_space = (recipe.get('composition') or {}).get('negative_space')
            next_space = 'right' if current_space == 'left' else 'left'
            second = update_image_recipe(root, recipe['recipe_id'], negative_space=next_space)
            self.assertTrue(second['layout_invalidated'])
            state_after = load_state(root)['slides']['OUT-S001']
            self.assertEqual(state_after['layout'], 'STALE')
            trace = [json.loads(line) for line in (root / 'decisions/decision_trace.jsonl').read_text().splitlines() if line.strip()]
            self.assertTrue(trace[-1]['layout_invalidated'])
            self.assertIn('OUT-S001:layout', trace[-1]['rerun_scope'])

    def _write_coherence_fixture(self, root: Path, rows: list[tuple[str, str, str]]) -> None:
        (root / 'planning').mkdir(parents=True, exist_ok=True)
        (root / 'render/generated_assets').mkdir(parents=True, exist_ok=True)
        recipes = []
        assets = []
        selected = {}
        for index, (recipe_id, slide_id, average_hash) in enumerate(rows, start=1):
            asset_id = f'GAST-{index:03d}'
            recipes.append({'recipe_id': recipe_id, 'slide_id': slide_id, 'selected_asset_id': asset_id})
            assets.append({
                'asset_id': asset_id, 'recipe_id': recipe_id, 'slide_id': slide_id,
                'metrics': {
                    'average_hash': average_hash, 'brightness': .70, 'saturation': .12,
                    'contrast': .55, 'edge_density': .08, 'warmth': .05,
                },
            })
            selected[recipe_id] = asset_id
        (root / 'planning/image_recipes.json').write_text(json.dumps({'recipes': recipes}), encoding='utf-8')
        (root / 'render/generated_assets/manifest.json').write_text(json.dumps({
            'schema_version': '0.1', 'assets': assets, 'selected_by_recipe': selected,
        }), encoding='utf-8')

    def test_same_slide_image_series_is_not_reported_as_cross_slide_duplicate(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            self._write_coherence_fixture(root, [
                ('IMG-A', 'OUT-S001', 'f1e4e6e6e6ece4f0'),
                ('IMG-B', 'OUT-S001', 'f1e4eee6e6eee4f0'),
            ])
            report = validate_deck_image_coherence(root)
            self.assertEqual(report['warning_count'], 0)
            self.assertFalse(any(c['check_id'] == 'DECK_IMAGE_NEAR_DUPLICATE' for c in report['checks']))
            self.assertTrue(any(c['check_id'] == 'DECK_IMAGE_SHARED_SERIES_STYLE' for c in report['checks']))

    def test_near_duplicate_across_slides_is_reported(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            self._write_coherence_fixture(root, [
                ('IMG-A', 'OUT-S001', 'f1e4e6e6e6ece4f0'),
                ('IMG-B', 'OUT-S002', 'f1e4eee6e6eee4f0'),
            ])
            report = validate_deck_image_coherence(root)
            duplicate = next(c for c in report['checks'] if c['check_id'] == 'DECK_IMAGE_NEAR_DUPLICATE')
            self.assertEqual(duplicate['result'], 'fail')
            self.assertEqual(duplicate['severity'], 'warning')
            self.assertEqual(duplicate['pairs'][0]['slide_ids'], ['OUT-S001', 'OUT-S002'])


if __name__ == '__main__':
    unittest.main()
