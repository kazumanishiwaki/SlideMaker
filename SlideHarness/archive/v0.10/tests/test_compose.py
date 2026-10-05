import json, shutil, tempfile, unittest
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
from compose_deck import compose


class ComposeTests(unittest.TestCase):
    def _project(self, td):
        from pptx import Presentation
        p=Path(td)
        for d in ['inputs','planning/slide_plans','output']:
            (p/d).mkdir(parents=True,exist_ok=True)
        shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
        source=Presentation(p/'inputs/source.pptx')
        slide=source.slides[1]  # source page 2, orange statement page
        text_shapes=[s for s in slide.shapes if getattr(s,'has_text_frame',False) and s.text.strip()]
        self.assertGreaterEqual(len(text_shapes),2)
        title=max(text_shapes,key=lambda s: len(s.text))
        (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[{
            'slide_id':'OUT-S001','order':1,'primary_message':'新しいメッセージ'
        }]}),encoding='utf-8')
        (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[{
            'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S002','page_number':2,'recommended_mode':'DIRECT'}
        }]}),encoding='utf-8')
        plan={
            'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S002',
            'adaptation_mode':'DIRECT','composer_support':'supported',
            'content_blocks':[{'id':'CB-01','role':'title','source_text':title.text,'semantic_text':'新しいメッセージ','render_text':'新しいメッセージ'}],
            'source_object_mutations':[{'source_object_id':f'SRC-S002-O{title.shape_id:03d}','action':'replace_text','content_block_id':'CB-01','render_text':'新しいメッセージ','reason':'test'}],
            'status':'APPROVED'
        }
        (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
        return p, slide

    def test_composes_text_and_preserves_background(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p, source_slide=self._project(td)
            result=compose(p,slide_ids=['OUT-S001'],output_name='one.pptx')
            self.assertEqual(result['errors'],0)
            out=Presentation(result['pptx'])
            self.assertEqual(len(out.slides),1)
            self.assertIn('新しいメッセージ','\n'.join(s.text for s in out.slides[0].shapes if getattr(s,'has_text_frame',False)))
            self.assertEqual(out.slides[0].element.cSld.bg.xml,source_slide.element.cSld.bg.xml)

    def test_refuses_plan_requiring_layout_engine(self):
        with tempfile.TemporaryDirectory() as td:
            p,_=self._project(td)
            plan_path=p/'planning/slide_plans/OUT-S001.json'
            plan=json.loads(plan_path.read_text())
            plan['composer_support']='requires_layout_engine'
            plan['layout_engine_reason']='needs new diagram geometry'
            plan_path.write_text(json.dumps(plan),encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'requires layout engine|cannot be composed'):
                compose(p,slide_ids=['OUT-S001'])

    def test_replaces_source_table_with_editable_native_table(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)
            for d in ['inputs','planning/slide_plans','output']:
                (p/d).mkdir(parents=True,exist_ok=True)
            shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
            source=Presentation(p/'inputs/source.pptx'); slide=source.slides[10]
            table_shape=next(s for s in slide.shapes if getattr(s,'has_table',False))
            (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','order':1,'primary_message':'比較'}]}),encoding='utf-8')
            (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S011','page_number':11,'recommended_mode':'ADAPT'}}]}),encoding='utf-8')
            data=[['機能','Lite','フル版'],['抽選','○','○'],['OCR','×','○']]
            plan={
                'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S011','adaptation_mode':'ADAPT','composer_support':'supported',
                'content_blocks':[],
                'source_object_mutations':[{'source_object_id':f'SRC-S011-O{table_shape.shape_id:03d}','action':'replace_table','table_data':data,'column_width_weights':[2.4,1,1],'header_rows':1,'reason':'test editable comparison'}],
                'status':'APPROVED'
            }
            (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
            result=compose(p,slide_ids=['OUT-S001'],output_name='table.pptx')
            out=Presentation(result['pptx']); tbl=next(s.table for s in out.slides[0].shapes if getattr(s,'has_table',False))
            self.assertEqual(len(tbl.rows),3); self.assertEqual(len(tbl.columns),3)
            self.assertEqual(tbl.cell(0,0).text,'機能'); self.assertEqual(tbl.cell(2,2).text,'○')


    def test_replaces_picture_region_with_editable_chart(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)
            for d in ['inputs','planning/slide_plans','output']:
                (p/d).mkdir(parents=True,exist_ok=True)
            shutil.copy2(ROOT/'examples/test-case/inputs/source.pptx',p/'inputs/source.pptx')
            source=Presentation(p/'inputs/source.pptx'); slide=source.slides[4]
            target=next(s for s in slide.shapes if s.shape_id == 9)
            (p/'planning/deck_plan.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','order':1,'primary_message':'推移'}]}),encoding='utf-8')
            (p/'planning/slide_matching.json').write_text(json.dumps({'slides':[{'slide_id':'OUT-S001','selected':{'source_slide_id':'SRC-S005','page_number':5,'recommended_mode':'ADAPT'}}]}),encoding='utf-8')
            plan={
                'schema_version':'0.1','slide_id':'OUT-S001','source_slide_id':'SRC-S005','adaptation_mode':'ADAPT','composer_support':'supported',
                'content_blocks':[],
                'source_object_mutations':[{'source_object_id':'SRC-S005-O009','action':'replace_chart','chart_type':'line','categories':['Q1','Q2','Q3'],'series':[{'name':'応募数','values':[100,140,210]}],'has_legend':False,'reason':'test editable chart'}],
                'status':'APPROVED'
            }
            (p/'planning/slide_plans/OUT-S001.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
            result=compose(p,slide_ids=['OUT-S001'],output_name='chart.pptx')
            out=Presentation(result['pptx']); charts=[s.chart for s in out.slides[0].shapes if getattr(s,'has_chart',False)]
            self.assertEqual(len(charts),1)
            self.assertEqual(len(charts[0].series),1)
            self.assertEqual(list(charts[0].series[0].values),[100.0,140.0,210.0])


if __name__=='__main__': unittest.main()
