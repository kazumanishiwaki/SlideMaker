from __future__ import annotations
import shutil
import subprocess
import tempfile
import re
import os
import signal
from pathlib import Path


LIBREOFFICE_TIMEOUT_SECONDS = 120
PDF_RENDER_TIMEOUT_SECONDS = 120

def _run_bounded(args: list[str], *, timeout: int) -> subprocess.CompletedProcess:
    """Run an external renderer and terminate its entire process group on timeout.

    LibreOffice launches oosplash/soffice.bin descendants. subprocess.run() only
    guarantees termination of the direct child, which can leave an orphan holding
    files/profile locks in a long-running App Server.
    """
    proc = subprocess.Popen(
        args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True,
    )
    try:
        stdout, stderr = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        try:
            os.killpg(proc.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        stdout, stderr = proc.communicate()
        raise subprocess.TimeoutExpired(args, timeout, output=stdout, stderr=stderr) from exc
    if proc.returncode != 0:
        raise subprocess.CalledProcessError(proc.returncode, args, output=stdout, stderr=stderr)
    return subprocess.CompletedProcess(args, proc.returncode, stdout, stderr)


def _libreoffice_profile_arg(profile_dir: Path) -> str:
    """Return an isolated LibreOffice profile argument.

    LibreOffice may otherwise reuse or wait on a process-wide user profile lock.
    App Server requests can render concurrently, so every conversion gets its own
    temporary profile to keep rendering deterministic and bounded.
    """
    profile_dir.mkdir(parents=True, exist_ok=True)
    return f"-env:UserInstallation={profile_dir.resolve().as_uri()}"


def render_pptx(pptx: str | Path, output_dir: str | Path) -> list[str]:
    pptx = Path(pptx).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    # pdftoppm 25.06 has been observed to hang on a second conversion in the
    # same long-running Python/App Server process. pdftocairo uses the same
    # Poppler rendering stack but is stable across repeated requests here.
    pdf_renderers = [x for x in (shutil.which("pdftocairo"), shutil.which("pdftoppm")) if x]
    # Keep order while removing duplicate resolved paths.
    pdf_renderers = list(dict.fromkeys(pdf_renderers))
    if not libreoffice or not pdf_renderers:
        return []

    # Remove only files owned by this renderer. Keeping unrelated diagnostics in
    # the same directory is useful, but stale page images must never be mistaken
    # for the current render.
    for pattern in ("source-*.png", "SRC-S*.png"):
        for old in output_dir.glob(pattern):
            old.unlink(missing_ok=True)

    pdf_dir = output_dir / "_pdf"
    shutil.rmtree(pdf_dir, ignore_errors=True)
    pdf_dir.mkdir(exist_ok=True)

    try:
        conversion_errors = []
        pdf = pdf_dir / (pptx.stem + ".pdf")
        for attempt in range(2):
            shutil.rmtree(pdf_dir, ignore_errors=True)
            pdf_dir.mkdir(exist_ok=True)
            try:
                with tempfile.TemporaryDirectory(prefix="ppt-lo-profile-") as profile:
                    _run_bounded([
                        libreoffice,
                        _libreoffice_profile_arg(Path(profile)),
                        "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
                        "--convert-to", "pdf",
                        "--outdir", str(pdf_dir),
                        str(pptx),
                    ], timeout=LIBREOFFICE_TIMEOUT_SECONDS)
                if pdf.exists() or list(pdf_dir.glob("*.pdf")):
                    break
                conversion_errors.append(f"attempt {attempt + 1}: no PDF produced")
            except (subprocess.TimeoutExpired, subprocess.CalledProcessError) as exc:
                stderr = getattr(exc, "stderr", b"") or b""
                if isinstance(stderr, bytes):
                    stderr = stderr.decode("utf-8", errors="replace")
                conversion_errors.append(f"attempt {attempt + 1}: {type(exc).__name__}: {str(stderr)[-400:]}")
        else:
            raise RuntimeError("LibreOffice conversion failed; " + " | ".join(conversion_errors))

        pdf = pdf_dir / (pptx.stem + ".pdf")
        if not pdf.exists():
            candidates = list(pdf_dir.glob("*.pdf"))
            if not candidates:
                return []
            pdf = candidates[0]

        prefix = output_dir / "source"
        render_errors = []
        rendered_ok = False
        for pdf_renderer in pdf_renderers:
            for stale in output_dir.glob("source-*.png"):
                stale.unlink(missing_ok=True)
            try:
                _run_bounded([
                    pdf_renderer, "-png", "-r", "96", str(pdf), str(prefix)
                ], timeout=PDF_RENDER_TIMEOUT_SECONDS)
                if list(output_dir.glob("source-*.png")):
                    rendered_ok = True
                    break
                render_errors.append(f"{Path(pdf_renderer).name}: no PNG pages produced")
            except subprocess.TimeoutExpired as exc:
                render_errors.append(f"{Path(pdf_renderer).name}: timed out after {exc.timeout}s")
            except subprocess.CalledProcessError as exc:
                stderr = (exc.stderr or b"").decode("utf-8", errors="replace").strip()
                render_errors.append(f"{Path(pdf_renderer).name}: exit {exc.returncode}: {stderr[-500:]}")
        if not rendered_ok:
            raise RuntimeError("PDF page rendering failed; " + " | ".join(render_errors))

        def page_number(file: Path) -> int:
            match = re.search(r"-(\d+)\.png$", file.name)
            return int(match.group(1)) if match else 10**9

        files = sorted(output_dir.glob("source-*.png"), key=page_number)
        renamed = []
        for i, file in enumerate(files, start=1):
            dst = output_dir / f"SRC-S{i:03d}.png"
            dst.unlink(missing_ok=True)
            file.replace(dst)
            renamed.append(str(dst))
        return renamed
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"render command timed out after {exc.timeout}s: {pptx.name}") from exc
    finally:
        shutil.rmtree(pdf_dir, ignore_errors=True)


def render_selected_source_slides(
    pptx: str | Path,
    slide_numbers: list[int] | tuple[int, ...] | set[int],
    output_dir: str | Path,
) -> list[str]:
    """Render only the source slides referenced by the current output deck.

    Large real-world templates may contain dozens or hundreds of pages while one
    generation uses only a handful.  Pixel-diff QA needs the selected source pages,
    not an eager render of the entire template.  The original source IDs are retained
    in the output file names (SRC-S042.png etc.).
    """
    from pptx import Presentation

    pptx = Path(pptx).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    wanted = sorted({int(value) for value in slide_numbers if int(value) >= 1})
    if not wanted:
        return []
    prs = Presentation(pptx)
    total = len(prs.slides)
    invalid = [value for value in wanted if value > total]
    if invalid:
        raise IndexError(f'source slide numbers outside deck: {invalid}; total={total}')

    # Reuse a complete cache when it already exists.
    existing = [output_dir / f'SRC-S{number:03d}.png' for number in wanted]
    if all(path.exists() for path in existing):
        return [str(path) for path in existing]

    with tempfile.TemporaryDirectory(prefix='ppt-selected-source-') as td:
        td_path = Path(td)
        subset_path = td_path / 'selected-source.pptx'
        keep = set(wanted)
        # Removing sldId entries preserves each selected slide part, master, layout,
        # theme, and relationships without cross-package cloning.
        for index in range(len(prs.slides), 0, -1):
            if index in keep:
                continue
            sld_id = prs.slides._sldIdLst[index - 1]
            rel_id = sld_id.rId
            prs.part.drop_rel(rel_id)
            prs.slides._sldIdLst.remove(sld_id)
        prs.save(subset_path)
        rendered_dir = td_path / 'rendered'
        rendered = render_pptx(subset_path, rendered_dir)
        if len(rendered) != len(wanted):
            raise RuntimeError(f'selected source render mismatch: expected {len(wanted)}, got {len(rendered)}')
        outputs: list[str] = []
        for original_number, rendered_path in zip(wanted, rendered):
            target = output_dir / f'SRC-S{original_number:03d}.png'
            shutil.copy2(rendered_path, target)
            outputs.append(str(target))
        return outputs
