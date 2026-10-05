from __future__ import annotations
from pathlib import Path
from typing import Any

from pptx import Presentation
from ppt_clone import clone_slide_into

from helpers import now_iso, read_json, write_json


def build(project: str | Path, output_name: str = 'template_skeleton.pptx') -> dict[str, Any]:
    root = Path(project).resolve()
    source_path = root / 'inputs/source.pptx'
    matches = read_json(root / 'planning/slide_matching.json', {})
    deck = read_json(root / 'planning/deck_plan.json', {})
    if not source_path.exists():
        raise FileNotFoundError(source_path)
    match_by_id = {x['slide_id']: x for x in matches.get('slides', [])}
    ordered = sorted(deck.get('slides', []), key=lambda x: x.get('order', 9999))
    if not ordered:
        raise ValueError('deck_plan has no slides')

    prs = Presentation(source_path)
    originals = list(prs.slides)
    keep_ids = []
    mapping = []
    for out_slide in ordered:
        slide_id = out_slide['slide_id']
        match = match_by_id.get(slide_id, {})
        selected = match.get('selected') or (match.get('candidates') or [None])[0]
        if not selected:
            raise ValueError(f'No source match for {slide_id}')
        page_number = int(selected.get('page_number') or str(selected['source_slide_id']).split('S')[-1])
        if page_number < 1 or page_number > len(originals):
            raise ValueError(f'Invalid source page {page_number} for {slide_id}')
        clone_slide_into(prs, originals[page_number - 1])
        keep_ids.append(prs.slides._sldIdLst[-1])
        mapping.append({
            'slide_id': slide_id,
            'order': out_slide.get('order'),
            'source_slide_id': selected['source_slide_id'],
            'source_page_number': page_number,
            'adaptation_mode': selected.get('recommended_mode', 'ADAPT'),
            'primary_message': out_slide.get('primary_message'),
        })

    # Remove originals and any slide not explicitly cloned for this skeleton.
    for sld_id in list(prs.slides._sldIdLst):
        if sld_id not in keep_ids:
            prs.slides._sldIdLst.remove(sld_id)

    output_dir = root / 'output'
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / output_name
    prs.save(out_path)
    map_path = output_dir / 'template_skeleton_map.json'
    write_json(map_path, {
        'schema_version': '0.1',
        'generated_at': now_iso(),
        'kind': 'pre_content_skeleton',
        'warning': 'Layout/source-slide clone only. Not a final presentation; content adaptation has not run.',
        'slides': mapping,
    })
    return {'pptx': str(out_path), 'map': str(map_path), 'slide_count': len(mapping)}
