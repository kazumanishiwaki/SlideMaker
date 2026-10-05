from __future__ import annotations
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean
from typing import Any
import math
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from helpers import now_iso, sha256_bytes, write_json
from image_metrics import analyze_image

EMU_PER_INCH = 914400


def inch(value: int) -> float:
    return round(value / EMU_PER_INCH, 4)


def rgb_to_hex(rgb) -> str | None:
    try:
        if rgb is None:
            return None
        return str(rgb)
    except Exception:
        return None


def shape_text_style(shape) -> dict[str, Any]:
    fonts = Counter()
    sizes = Counter()
    colors = Counter()
    bold_count = 0
    run_count = 0
    if not getattr(shape, "has_text_frame", False):
        return {}
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if not r.text:
                continue
            run_count += 1
            if r.font.name:
                fonts[r.font.name] += len(r.text)
            if r.font.size:
                sizes[round(r.font.size.pt, 1)] += len(r.text)
            if r.font.bold:
                bold_count += 1
            try:
                c = rgb_to_hex(r.font.color.rgb)
                if c:
                    colors[c] += len(r.text)
            except Exception:
                pass
    return {
        "fonts": fonts.most_common(5),
        "sizes_pt": sizes.most_common(5),
        "colors": colors.most_common(5),
        "bold_run_ratio": round(bold_count / run_count, 3) if run_count else 0,
    }


def infer_role(page_number: int, texts: list[str], shape_count: int, picture_count: int) -> tuple[str, list[str], float]:
    joined = "\n".join(texts)
    if page_number == 1:
        return "cover", [], 0.98
    rules = [
        (("時間の経過", "タイムライン"), "timeline", ["process"]),
        (("3項目", "3段レイアウト"), "three_items", []),
        (("4段レイアウト",), "four_items", []),
        (("写真とテキストによる4段",), "photo_grid", ["four_items"]),
        (("組織図",), "organization", ["hierarchy"]),
        (("業績ハイライト",), "kpi", ["chart"]),
        (("ベン図",), "comparison", ["relationship"]),
        (("会社概要",), "table", ["profile"]),
        (("長文テキスト",), "long_text", ["statement"]),
        (("QnA", "Q&A"), "qna", []),
        (("Thank you", "THANK YOU"), "closing", []),
        (("目次", "CONTENT"), "overview", ["agenda"]),
    ]
    for terms, role, secondary in rules:
        if any(term in joined for term in terms):
            return role, secondary, 0.94
    if picture_count >= 4:
        return "photo_grid", [], 0.62
    if shape_count <= 8:
        return "statement", [], 0.55
    return "overview", [], 0.50


def positional_variance(occurrences: list[dict[str, Any]]) -> float:
    if len(occurrences) <= 1:
        return 0.0
    xs = [o["bbox"]["x"] for o in occurrences]
    ys = [o["bbox"]["y"] for o in occurrences]
    ws = [o["bbox"]["w"] for o in occurrences]
    hs = [o["bbox"]["h"] for o in occurrences]
    vals = []
    for arr in (xs, ys, ws, hs):
        m = mean(arr)
        vals.extend((v - m) ** 2 for v in arr)
    return math.sqrt(sum(vals) / len(vals))



def bbox_distance(a: dict[str, float], b: dict[str, float]) -> float:
    ax0, ay0, ax1, ay1 = a['x'], a['y'], a['x'] + a['w'], a['y'] + a['h']
    bx0, by0, bx1, by1 = b['x'], b['y'], b['x'] + b['w'], b['y'] + b['h']
    dx = max(bx0 - ax1, ax0 - bx1, 0.0)
    dy = max(by0 - ay1, ay0 - by1, 0.0)
    return math.hypot(dx, dy)


def picture_role_hint(obj: dict[str, Any], *, area_ratio: float, near_edge: bool, slide_role: str) -> tuple[str, list[str]]:
    name = str(obj.get('name') or '').lower()
    default_picture_name = bool(re.fullmatch(r'picture\s*\d*', name))
    metrics = ((obj.get('picture') or {}).get('visual_metrics') or {})
    transparent = float(metrics.get('transparent_ratio') or 0.0)
    edge_density = float(metrics.get('edge_density') or 0.0)
    saturation = float(metrics.get('saturation') or 0.0)
    aspect = float(((obj.get('picture') or {}).get('display') or {}).get('slot_aspect_ratio') or 0.0)
    reasons: list[str] = []
    if re.search(r'logo|brand|watermark|header|footer|corner|ornament|decor|background|\bbg\b', name):
        reasons.append('SHAPE_NAME_DECORATIVE')
        return ('logo' if 'logo' in name or 'brand' in name else 'background_visual' if 'background' in name or 'bg' in name else 'decorative_visual', reasons)
    if re.search(r'screenshot|screen|capture|mockup|ui|browser|device', name):
        reasons.append('SHAPE_NAME_SCREENSHOT')
        return 'screenshot', reasons
    if re.search(r'icon|symbol|pictogram', name):
        reasons.append('SHAPE_NAME_ICON')
        return 'icon', reasons
    if not default_picture_name and re.search(r'photo|portrait|hero|product|team|people', name):
        reasons.append('SHAPE_NAME_PHOTO')
        return 'photo', reasons
    if area_ratio >= 0.58 and int(obj.get('z_index') or 0) <= 2:
        reasons.append('LARGE_BACK_LAYER')
        return 'background_visual', reasons
    if transparent >= 0.18 and area_ratio <= 0.08:
        reasons.append('TRANSPARENT_SMALL_GRAPHIC')
        return 'icon' if edge_density >= 0.06 else 'decorative_visual', reasons
    if edge_density >= 0.14 and saturation <= 0.32 and area_ratio >= 0.06:
        reasons.append('HIGH_EDGE_LOW_SATURATION')
        return 'screenshot', reasons
    if aspect >= 4.5 and near_edge:
        reasons.append('WIDE_EDGE_GRAPHIC')
        return 'decorative_visual', reasons
    if slide_role in {'photo_grid', 'profile', 'case_study'} and area_ratio >= 0.03:
        reasons.append('SLIDE_ROLE_PHOTO_LED')
        return 'photo', reasons
    return 'photo' if area_ratio >= 0.035 else 'unknown', reasons


def classify_picture_seed(
    obj: dict[str, Any], *, repeat_count: int, geometry_repeat_count: int, variance: float, area_ratio: float,
    near_edge: bool, slide_role: str,
) -> tuple[str, str, float, list[str], str]:
    name = str(obj.get('name') or '').lower()
    context = obj.get('context') or {}
    nearest_text_distance = context.get('nearest_text_distance')
    role_hint, role_reasons = picture_role_hint(obj, area_ratio=area_ratio, near_edge=near_edge, slide_role=slide_role)
    reasons = list(role_reasons)
    bbox = obj.get('bbox') or {}
    bw = max(0.0, float(bbox.get('w') or 0.0)); bh = max(0.0, float(bbox.get('h') or 0.0))
    aspect = max(bw, bh) / max(0.001, min(bw, bh)) if bw and bh else 0.0
    nearest_text_distance = context.get('nearest_text_distance')
    if re.search(r'logo|brand|watermark|header|footer|corner|ornament|decor', name):
        return 'template_chrome', 'keep', 0.97, reasons + ['NAME_OVERRIDES_AMBIGUITY'], role_hint
    if repeat_count >= 3 and variance < 0.08:
        return 'template_chrome', 'keep', 0.96, reasons + ['REPEATED_ACROSS_SLIDES', 'POSITION_STABLE'], role_hint
    if repeat_count >= 2 and near_edge and variance < 0.18:
        return 'template_chrome', 'keep', 0.88, reasons + ['REPEATED', 'EDGE_PLACEMENT'], role_hint
    if geometry_repeat_count >= 3 and near_edge and area_ratio <= 0.035:
        return 'template_chrome', 'keep', 0.86, reasons + ['GEOMETRY_REPEATED_ACROSS_SLIDES', 'EDGE_PLACEMENT'], 'decorative_visual'
    if aspect >= 5.0 and min(bw, bh) <= 0.12:
        confidence = 0.84 if near_edge or geometry_repeat_count >= 2 else 0.72
        return 'template_chrome', 'keep', confidence, reasons + ['THIN_DECORATIVE_RASTER'], 'decorative_visual'
    if role_hint == 'background_visual' and near_edge and area_ratio >= 0.45:
        return 'template_chrome', 'keep', 0.80, reasons + ['BACKGROUND_LAYER_CANDIDATE'], role_hint
    if repeat_count >= 2 and 0.04 <= area_ratio <= 0.28 and variance < 0.25:
        return 'container_or_frame', 'replace_inner', 0.66, reasons + ['REPEATED_CONTAINER_LIKE', 'CENTRAL_OR_CONTENT_AREA'], 'frame'
    if role_hint in {'screenshot', 'photo'} and area_ratio >= 0.025:
        confidence = 0.90 if nearest_text_distance is not None and nearest_text_distance <= 0.30 else 0.84
        if nearest_text_distance is not None and nearest_text_distance <= 0.30:
            reasons.append('ADJACENT_TO_TEXT')
        return 'content_visual', 'replace', confidence, reasons + ['PAGE_CONTENT_VISUAL'], role_hint
    if role_hint == 'icon' and nearest_text_distance is not None and nearest_text_distance <= 0.30:
        return 'content_visual', 'replace', 0.78, reasons + ['SEMANTIC_ICON_ADJACENT_TO_TEXT'], role_hint
    if repeat_count == 1 and area_ratio >= 0.035:
        return 'content_visual', 'replace', 0.82, reasons + ['PAGE_SPECIFIC', 'VISUALLY_SIGNIFICANT'], role_hint
    if repeat_count == 1 and area_ratio < 0.01 and near_edge:
        return 'template_chrome', 'keep', 0.74, reasons + ['PAGE_SPECIFIC', 'TINY_EDGE_DECORATION'], role_hint
    if area_ratio < 0.02 and 0.55 <= aspect <= 1.8 and nearest_text_distance is not None and nearest_text_distance <= 0.24:
        return 'content_visual', 'replace', 0.76, reasons + ['SMALL_SEMANTIC_ICON', 'ADJACENT_TO_TEXT'], 'icon'
    if repeat_count == 1 and area_ratio < 0.02:
        return 'unknown', 'review_required', 0.54, reasons + ['PAGE_SPECIFIC', 'SMALL_AMBIGUOUS_VISUAL'], role_hint
    if repeat_count == 1:
        return 'content_visual', 'replace', 0.68, reasons + ['PAGE_SPECIFIC', 'CONTENT_AREA_CANDIDATE'], role_hint
    return 'unknown', 'review_required', 0.45, reasons + ['AMBIGUOUS'], role_hint

def analyze(pptx_path: str | Path, output_dir: str | Path) -> dict[str, Any]:
    pptx_path = Path(pptx_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    source_media_dir = output_dir / "source_media"
    source_media_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(pptx_path)
    slide_w = inch(prs.slide_width)
    slide_h = inch(prs.slide_height)
    slide_area = slide_w * slide_h

    slides: list[dict[str, Any]] = []
    picture_occurrences: dict[str, list[dict[str, Any]]] = defaultdict(list)
    picture_metadata: dict[str, dict[str, Any]] = {}
    global_fonts = Counter()
    global_sizes = Counter()
    global_colors = Counter()

    for s_idx, slide in enumerate(prs.slides, start=1):
        objects = []
        texts = []
        picture_count = 0
        for z_index, shape in enumerate(slide.shapes):
            obj = {
                "source_object_id": f"SRC-S{s_idx:03d}-O{shape.shape_id:03d}",
                "shape_id": shape.shape_id,
                "name": shape.name,
                "shape_type": str(shape.shape_type),
                "z_index": z_index,
                "bbox": {"x": inch(shape.left), "y": inch(shape.top), "w": inch(shape.width), "h": inch(shape.height)},
            }
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                obj["text"] = text
                obj["text_chars"] = len(text)
                style = shape_text_style(shape)
                obj["text_style"] = style
                texts.append(text)
                for font, weight in style.get("fonts", []):
                    global_fonts[font] += weight
                for size, weight in style.get("sizes_pt", []):
                    global_sizes[size] += weight
                for color, weight in style.get("colors", []):
                    global_colors[color] += weight
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
                try:
                    blob = shape.image.blob
                    h = sha256_bytes(blob)
                    ext = str(shape.image.ext or "png").lower().lstrip(".")
                    media_name = f"{h}.{ext}"
                    media_path = source_media_dir / media_name
                    if not media_path.exists():
                        media_path.write_bytes(blob)
                    visual_metrics = None
                    try:
                        visual_metrics = analyze_image(blob)
                    except Exception:
                        # Vector or malformed picture payloads remain usable as PPT assets,
                        # but cannot participate in raster harmony metrics.
                        visual_metrics = None
                    obj["picture"] = {
                        "sha256": h,
                        "ext": ext,
                        "size_bytes": len(blob),
                        "pixel_size": [shape.image.size[0], shape.image.size[1]],
                        "media_path": f"analysis/source_media/{media_name}",
                        "visual_metrics": visual_metrics,
                        "display": {
                            "rotation": round(float(shape.rotation or 0), 3),
                            "crop_left": round(float(shape.crop_left or 0), 5),
                            "crop_top": round(float(shape.crop_top or 0), 5),
                            "crop_right": round(float(shape.crop_right or 0), 5),
                            "crop_bottom": round(float(shape.crop_bottom or 0), 5),
                            "slot_aspect_ratio": round(float(shape.width) / float(shape.height), 5) if shape.height else None,
                        },
                    }
                    picture_metadata[h] = obj["picture"]
                    occurrence = {
                        "slide_id": f"SRC-S{s_idx:03d}",
                        "object_id": obj["source_object_id"],
                        "bbox": obj["bbox"],
                    }
                    picture_occurrences[h].append(occurrence)
                except Exception as e:
                    obj["picture_error"] = str(e)
            objects.append(obj)

        text_objects = [row for row in objects if row.get("text")]
        for picture in [row for row in objects if row.get("picture")]:
            nearby = sorted(
                ((bbox_distance(picture["bbox"], text_row["bbox"]), text_row) for text_row in text_objects),
                key=lambda pair: pair[0],
            )
            if nearby:
                distance, text_row = nearby[0]
                picture["context"] = {
                    "nearest_text_distance": round(distance, 4),
                    "nearest_text_object_id": text_row.get("source_object_id"),
                    "nearest_text": str(text_row.get("text") or "")[:240],
                }
            else:
                picture["context"] = {"nearest_text_distance": None, "nearest_text_object_id": None, "nearest_text": ""}

        primary_role, secondary_roles, role_conf = infer_role(s_idx, texts, len(objects), picture_count)
        text_chars = [o.get("text_chars", 0) for o in objects if o.get("text_chars")]
        slides.append({
            "id": f"SRC-S{s_idx:03d}",
            "page_number": s_idx,
            "roles": {"primary": primary_role, "secondary": secondary_roles, "confidence": role_conf},
            "layout_family": primary_role,
            "stats": {
                "shape_count": len(objects),
                "picture_count": picture_count,
                "text_shape_count": sum(1 for o in objects if o.get("text")),
                "text_char_total": sum(text_chars),
                "text_char_max": max(text_chars) if text_chars else 0,
            },
            "objects": objects,
            "status": "DRAFT"
        })

    asset_records = []
    repeated_chrome = []
    object_lookup = {o["source_object_id"]: o for s in slides for o in s["objects"]}
    picture_geometry_rows = [
        {"slide_id": slide["id"], "object_id": obj["source_object_id"], "bbox": obj["bbox"]}
        for slide in slides for obj in slide["objects"] if obj.get("picture")
    ]

    def geometry_repeat_count(target: dict[str, Any]) -> int:
        box = target["bbox"]
        x_tol=max(0.10,slide_w*0.012); y_tol=max(0.10,slide_h*0.012)
        w_tol=max(0.10,slide_w*0.012); h_tol=max(0.10,slide_h*0.012)
        slides_seen=set()
        for row in picture_geometry_rows:
            other=row["bbox"]
            if (abs(float(box["x"])-float(other["x"])) <= x_tol and
                abs(float(box["y"])-float(other["y"])) <= y_tol and
                abs(float(box["w"])-float(other["w"])) <= w_tol and
                abs(float(box["h"])-float(other["h"])) <= h_tol):
                slides_seen.add(row["slide_id"])
        return len(slides_seen)

    for h, occs in picture_occurrences.items():
        repeat_count = len(occs)
        variance = positional_variance(occs)
        for occ in occs:
            obj = object_lookup[occ["object_id"]]
            bbox = occ["bbox"]
            area_ratio = (bbox["w"] * bbox["h"]) / slide_area if slide_area else 0
            near_edge = bbox["x"] < 0.8 or bbox["y"] < 0.8 or (bbox["x"] + bbox["w"]) > slide_w - 0.8 or (bbox["y"] + bbox["h"]) > slide_h - 0.8
            slide_row = next((row for row in slides if row.get("id") == occ["slide_id"]), None)
            slide_role = ((slide_row or {}).get("roles") or {}).get("primary") or "overview"
            classification, action, confidence, reasons, content_role = classify_picture_seed(
                obj, repeat_count=repeat_count, geometry_repeat_count=geometry_repeat_count(occ),
                variance=variance, area_ratio=area_ratio, near_edge=near_edge, slide_role=str(slide_role),
            )


            record = {
                "id": f"AST-{len(asset_records)+1:05d}",
                "source_slide_id": occ["slide_id"],
                "source_object_id": occ["object_id"],
                "sha256": h,
                "media_path": (picture_metadata.get(h) or {}).get("media_path"),
                "picture": picture_metadata.get(h),
                "visual_metrics": (picture_metadata.get(h) or {}).get("visual_metrics"),
                "source_display": (obj.get("picture") or {}).get("display"),
                "z_index": obj.get("z_index"),
                "shape_name": obj.get("name"),
                "context": obj.get("context") or {},
                "content_role": content_role,
                "classification": classification,
                "action_policy": action,
                "confidence": round(confidence, 2),
                "reason_codes": reasons,
                "repeat_count": repeat_count,
                "geometry_repeat_count": geometry_repeat_count(occ),
                "position_variance": round(variance, 4),
                "bbox": bbox,
                "area_ratio": round(area_ratio, 4),
                "status": "DRAFT" if confidence >= 0.7 else "REVIEW_REQUIRED",
            }
            asset_records.append(record)
            if classification == "template_chrome" and h not in repeated_chrome:
                repeated_chrome.append(h)

    source_slides = {
        "schema_version": "0.1",
        "generated_at": now_iso(),
        "source": str(pptx_path),
        "presentation": {"width_in": slide_w, "height_in": slide_h, "slide_count": len(slides)},
        "slides": slides,
    }
    source_assets = {
        "schema_version": "0.1",
        "generated_at": now_iso(),
        "assets": asset_records,
        "summary": {
            k: sum(1 for a in asset_records if a["classification"] == k)
            for k in ["template_chrome", "content_visual", "container_or_frame", "unknown"]
        }
    }
    template_profile = {
        "schema_version": "0.1",
        "generated_at": now_iso(),
        "presentation": {"width_in": slide_w, "height_in": slide_h},
        "visual_grammar": {
            "fonts": global_fonts.most_common(12),
            "font_sizes_pt": global_sizes.most_common(12),
            "text_colors": global_colors.most_common(12),
            "repeated_chrome_hashes": repeated_chrome,
        },
        "layout_families": sorted({s["layout_family"] for s in slides}),
        "notes": [
            "Observed template facts only. User preferences belong in design.md.",
            "Image classification is a deterministic seed and must be user-reviewable when confidence is low."
        ]
    }
    write_json(output_dir / "source_slides.json", source_slides)
    write_json(output_dir / "source_assets.json", source_assets)
    write_json(output_dir / "template_profile.json", template_profile)
    return {"source_slides": source_slides, "source_assets": source_assets, "template_profile": template_profile}
