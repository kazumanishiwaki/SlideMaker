from __future__ import annotations
from pathlib import Path
from copy import deepcopy
from typing import Any
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE

from helpers import now_iso, read_json, write_json
from ppt_clone import clone_slide_into
from schema_validation import validate_slide_plan
from layout_engine import apply_layout
from picture_ops import replace_picture

OBJ_RE = re.compile(r'-O(\d+)$')


def _shape_id(object_id: str) -> int:
    m = OBJ_RE.search(object_id or '')
    if not m:
        raise ValueError(f'Invalid source_object_id: {object_id}')
    return int(m.group(1))


def _font_snapshot(shape):
    if not getattr(shape, 'has_text_frame', False):
        return {}
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text or True:
                f = r.font
                rgb = None
                try: rgb = f.color.rgb
                except Exception: pass
                return {
                    'name': f.name,
                    'size': f.size,
                    'bold': f.bold,
                    'italic': f.italic,
                    'underline': f.underline,
                    'rgb': rgb,
                    'alignment': p.alignment,
                    'level': p.level,
                }
    p = shape.text_frame.paragraphs[0] if shape.text_frame.paragraphs else None
    return {'alignment': p.alignment if p else None, 'level': p.level if p else 0}


def _replace_text(shape, text: str) -> None:
    if not getattr(shape, 'has_text_frame', False):
        raise ValueError(f'Shape {shape.shape_id} is not a text shape')
    snap = _font_snapshot(shape)
    tf = shape.text_frame
    tf.clear()
    lines = str(text).split('\n')
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if snap.get('alignment') is not None: p.alignment = snap['alignment']
        try: p.level = snap.get('level', 0)
        except Exception: pass
        run = p.add_run(); run.text = line
        f = run.font
        if snap.get('name') is not None: f.name = snap['name']
        if snap.get('size') is not None: f.size = snap['size']
        if snap.get('bold') is not None: f.bold = snap['bold']
        if snap.get('italic') is not None: f.italic = snap['italic']
        if snap.get('underline') is not None: f.underline = snap['underline']
        if snap.get('rgb') is not None:
            try: f.color.rgb = snap['rgb']
            except Exception: pass


def _replace_picture(slide, old_shape, image_path: Path, *, fit_mode: str = "contain", focal_point=None):
    # Compatibility wrapper retained for the existing mutation call sites.
    return replace_picture(slide, old_shape, image_path, fit_mode=fit_mode, focal_point=focal_point)



def _text_style_from_cell(cell):
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            f=r.font
            rgb=None
            try: rgb=f.color.rgb
            except Exception: pass
            return {
                'name':f.name,'size':f.size,'bold':f.bold,'italic':f.italic,
                'rgb':rgb,'alignment':p.alignment,'vertical_anchor':cell.vertical_anchor,
            }
    p=cell.text_frame.paragraphs[0] if cell.text_frame.paragraphs else None
    return {'alignment':p.alignment if p else None,'vertical_anchor':cell.vertical_anchor}


def _apply_cell_text(cell, text, style):
    cell.text=''
    p=cell.text_frame.paragraphs[0]
    if style.get('alignment') is not None: p.alignment=style['alignment']
    run=p.add_run(); run.text=str(text)
    f=run.font
    if style.get('name') is not None: f.name=style['name']
    if style.get('size') is not None: f.size=style['size']
    if style.get('bold') is not None: f.bold=style['bold']
    if style.get('italic') is not None: f.italic=style['italic']
    if style.get('rgb') is not None:
        try: f.color.rgb=style['rgb']
        except Exception: pass
    if style.get('vertical_anchor') is not None:
        try: cell.vertical_anchor=style['vertical_anchor']
        except Exception: pass


def _copy_cell_properties(source_cell, target_cell):
    src=source_cell._tc.tcPr
    dst=target_cell._tc.tcPr
    if src is None: return
    if dst is not None: target_cell._tc.remove(dst)
    target_cell._tc.insert(0,deepcopy(src))


def _replace_table(slide, old_shape, data, column_weights=None, header_rows=1, region_shapes=None):
    if not getattr(old_shape,'has_table',False):
        raise ValueError(f'Shape {old_shape.shape_id} is not a table')
    if not isinstance(data,list) or not data or not all(isinstance(r,list) and r for r in data):
        raise ValueError('table_data must be a non-empty rectangular 2D array')
    cols=max(len(r) for r in data)
    if any(len(r)!=cols for r in data): raise ValueError('table_data rows must have equal column count')
    rows=len(data)
    old_table=old_shape.table
    region_shapes=list(region_shapes or [old_shape])
    if old_shape not in region_shapes: region_shapes.append(old_shape)
    left=min(s.left for s in region_shapes); top=min(s.top for s in region_shapes)
    right=max(s.left+s.width for s in region_shapes); bottom=max(s.top+s.height for s in region_shapes)
    tw=right-left; th=bottom-top
    new_shape=slide.shapes.add_table(rows,cols,left,top,tw,th)
    table=new_shape.table

    weights=column_weights if isinstance(column_weights,list) and len(column_weights)==cols and all(float(x)>0 for x in column_weights) else [1]*cols
    total=sum(float(x) for x in weights)
    assigned=0
    for c,w in enumerate(weights):
        width=tw-assigned if c==cols-1 else int(tw*float(w)/total)
        table.columns[c].width=width; assigned += width
    # Distribute rows uniformly. Source table cell styling, not geometry, is inherited.
    assigned=0
    for r in range(rows):
        height=th-assigned if r==rows-1 else int(th/rows)
        table.rows[r].height=height; assigned += height

    def sample_cell(r,c):
        # Header derives from source first row. Body first column derives from source first
        # column; remaining body derives from the closest available source body cell.
        sr=0 if r < max(0,int(header_rows)) else min(max(1,r),len(old_table.rows)-1)
        sc=0 if c==0 else min(c,len(old_table.columns)-1)
        return old_table.cell(sr,sc)

    for r,row in enumerate(data):
        for c,value in enumerate(row):
            src=sample_cell(r,c); dst=table.cell(r,c)
            _copy_cell_properties(src,dst)
            _apply_cell_text(dst,value,_text_style_from_cell(src))

    # Replace the selected semantic content region while preserving the earliest
    # source z-order slot. This lets a table expand into old photo/text regions without
    # asking the model to invent coordinates.
    new_el=new_shape.element; parent=old_shape.element.getparent(); parent.remove(new_el)
    region_elements=[s.element for s in region_shapes if s.element.getparent() is parent]
    idx=min(parent.index(el) for el in region_elements) if region_elements else len(parent)-1
    for el in region_elements:
        if el.getparent() is parent: parent.remove(el)
    parent.insert(idx,new_el)
    return new_shape


def _replace_chart(slide, old_shape, categories, series, chart_type='column', region_shapes=None, has_legend=None):
    if not isinstance(categories,list) or not categories:
        raise ValueError('chart categories must be a non-empty array')
    if not isinstance(series,list) or not series:
        raise ValueError('chart series must be a non-empty array')
    n=len(categories)
    for s in series:
        if not isinstance(s,dict) or len(s.get('values') or []) != n:
            raise ValueError('every chart series must have exactly one value per category')
    chart_types={
        'column':XL_CHART_TYPE.COLUMN_CLUSTERED,
        'bar':XL_CHART_TYPE.BAR_CLUSTERED,
        'line':XL_CHART_TYPE.LINE_MARKERS,
        'pie':XL_CHART_TYPE.PIE,
        'area':XL_CHART_TYPE.AREA,
    }
    if chart_type not in chart_types: raise ValueError(f'unsupported chart_type: {chart_type}')
    region_shapes=list(region_shapes or [old_shape])
    if old_shape not in region_shapes: region_shapes.append(old_shape)
    left=min(s.left for s in region_shapes); top=min(s.top for s in region_shapes)
    right=max(s.left+s.width for s in region_shapes); bottom=max(s.top+s.height for s in region_shapes)
    data=ChartData(); data.categories=[str(x) for x in categories]
    for ser in series: data.add_series(str(ser.get('name','')), [float(v) for v in ser['values']])
    frame=slide.shapes.add_chart(chart_types[chart_type],left,top,right-left,bottom-top,data)
    chart=frame.chart
    chart.has_title=False
    chart.has_legend = bool(has_legend) if has_legend is not None else len(series)>1
    if chart.has_legend:
        try: chart.legend.include_in_layout=False
        except Exception: pass
    # Use the source deck theme by leaving series colors automatic. The chart remains
    # a native editable PowerPoint chart instead of a raster image.
    new_el=frame.element; parent=old_shape.element.getparent(); parent.remove(new_el)
    region_elements=[s.element for s in region_shapes if s.element.getparent() is parent]
    idx=min(parent.index(el) for el in region_elements) if region_elements else len(parent)-1
    for el in region_elements:
        if el.getparent() is parent: parent.remove(el)
    parent.insert(idx,new_el)
    return frame

def _safe_asset(root: Path, rel: str) -> Path:
    p = (root / rel).resolve()
    if not str(p).startswith(str(root) + '/'): raise ValueError('asset path outside project')
    if not p.exists(): raise FileNotFoundError(p)
    return p


def compose(project: str | Path, *, slide_ids: list[str] | None = None, output_name='presentation_draft.pptx') -> dict[str, Any]:
    root = Path(project).resolve()
    source_path = root / 'inputs/source.pptx'
    deck = read_json(root / 'planning/deck_plan.json', {})
    matches = read_json(root / 'planning/slide_matching.json', {})
    match_by = {m['slide_id']:m for m in matches.get('slides',[])}
    ordered = sorted(deck.get('slides',[]), key=lambda s:s.get('order',9999))
    if slide_ids is not None:
        wanted=set(slide_ids); ordered=[s for s in ordered if s['slide_id'] in wanted]
    if not ordered: raise ValueError('No slides selected')

    state=read_json(root/'state/project_state.json',{})
    image_recipes=read_json(root/'planning/image_recipes.json',{})
    image_by_slide={}
    for recipe in image_recipes.get('recipes',[]):
        sid=recipe.get('slide_id')
        if sid:
            image_by_slide.setdefault(sid,[]).append(recipe)
    stale_values={'STALE','STALE_LOCKED','ERROR'}
    for out in ordered:
        sid=out['slide_id']; ss=(state.get('slides') or {}).get(sid,{})
        phases=['deck_plan','slide_plan']
        slide_recipes=image_by_slide.get(sid,[])
        if slide_recipes:
            phases.append('image')
        bad={phase:ss.get(phase) for phase in phases if ss.get(phase) in stale_values}
        if bad:
            raise ValueError(f'{sid} has stale/error dependencies and cannot be composed: {bad}')
        plan=read_json(root / f'planning/slide_plans/{sid}.json', {})
        family=str((plan.get('layout_request') or {}).get('family') or '')
        requires_layout_visual=family in {'photo_text','full_bleed'}
        missing=[]
        invalid_targets=[]
        for recipe in slide_recipes:
            target=recipe.get('target') or {}
            has_target=bool((target.get('source_object_id') or target.get('layout_slot_id')) and target.get('bbox'))
            if not has_target:
                invalid_targets.append(recipe.get('recipe_id'))
            elif not recipe.get('selected_asset_id'):
                missing.append(recipe.get('recipe_id'))
        if requires_layout_visual and not slide_recipes:
            missing.append(str((plan.get('layout_request') or {}).get('image_recipe_id') or f'IMG-{sid}-01'))
        if invalid_targets:
            raise ValueError(f'{sid} has unresolved image recipe targets: {invalid_targets}')
        if missing:
            raise ValueError(f'{sid} requires selected image assets for: {missing}')

    prs=Presentation(source_path); originals=list(prs.slides); keep_ids=[]; report=[]
    for out in ordered:
        sid=out['slide_id']; match=match_by.get(sid,{}); selected=match.get('selected') or (match.get('candidates') or [None])[0]
        if not selected: raise ValueError(f'No source match for {sid}')
        page=int(selected.get('page_number') or selected['source_slide_id'].split('S')[-1])
        new=clone_slide_into(prs,originals[page-1]); keep_ids.append(prs.slides._sldIdLst[-1])
        by_shape={shape.shape_id:shape for shape in new.shapes}
        plan=read_json(root / f'planning/slide_plans/{sid}.json', {})
        if plan:
            validate_slide_plan(plan)
        support=plan.get('composer_support','supported')
        wants_layout = support == 'layout_engine' or (support == 'requires_layout_engine' and bool(plan.get('layout_request')))
        if support == 'requires_layout_engine' and not wants_layout:
            reason = plan.get('layout_engine_reason') or 'Slide Plan requires layout engine support but has no executable layout_request'
            raise ValueError(f'{sid} cannot be composed: {reason}')
        mutations=plan.get('source_object_mutations',[])
        applied=[]; errors=[]
        for mut in mutations:
            try:
                shape=by_shape.get(_shape_id(mut.get('source_object_id')))
                if shape is None: raise ValueError('shape not found in cloned slide')
                action=mut.get('action')
                if action=='replace_text':
                    _replace_text(shape,mut.get('render_text','')); applied.append(mut)
                elif action=='clear_text':
                    _replace_text(shape,''); applied.append(mut)
                elif action=='replace_image':
                    asset=_safe_asset(root,mut['asset_path'])
                    _replace_picture(new, shape, asset, fit_mode=mut.get('fit_mode','contain'), focal_point=mut.get('focal_point')); applied.append(mut)
                elif action=='replace_table':
                    region=[]
                    for oid in mut.get('replace_region_object_ids') or []:
                        rs=by_shape.get(_shape_id(oid))
                        if rs is not None: region.append(rs)
                    _replace_table(new,shape,mut.get('table_data'),mut.get('column_width_weights'),mut.get('header_rows',1),region_shapes=region); applied.append(mut)
                elif action=='replace_chart':
                    region=[]
                    for oid in mut.get('replace_region_object_ids') or []:
                        rs=by_shape.get(_shape_id(oid))
                        if rs is not None: region.append(rs)
                    _replace_chart(new,shape,mut.get('categories'),mut.get('series'),mut.get('chart_type','column'),region_shapes=region,has_legend=mut.get('has_legend')); applied.append(mut)
                elif action=='remove':
                    el=shape.element; el.getparent().remove(el); applied.append(mut)
                elif action=='keep':
                    applied.append(mut)
                else:
                    raise ValueError(f'unsupported mutation action: {action}')
            except Exception as e:
                errors.append({'mutation':mut,'error':str(e)})
        layout_result=None
        if wants_layout and not errors:
            try:
                layout_result=apply_layout(root,new,plan)
            except Exception as e:
                errors.append({'layout_request':plan.get('layout_request'),'error':str(e)})
        report.append({'slide_id':sid,'source_page':page,'applied':len(applied),'errors':errors,'plan_present':bool(plan),
                       'layout_engine': {'used':bool(layout_result),'family':(layout_result or {}).get('family'),
                                         'quality':(layout_result or {}).get('quality')}})

    for sld_id in list(prs.slides._sldIdLst):
        if sld_id not in keep_ids: prs.slides._sldIdLst.remove(sld_id)
    outdir=root/'output'; outdir.mkdir(parents=True,exist_ok=True); outpath=outdir/output_name; prs.save(outpath)
    report_path=outdir/(Path(output_name).stem+'_compose_report.json')
    write_json(report_path,{'schema_version':'0.1','generated_at':now_iso(),'output':str(outpath),'slides':report})
    return {'pptx':str(outpath),'report':str(report_path),'slide_count':len(report),'errors':sum(len(x['errors']) for x in report)}
