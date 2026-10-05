from __future__ import annotations
import re
from pathlib import Path
from typing import Any
from helpers import now_iso, write_json

HEADING_RE = re.compile(r"^(?:\d{1,2}[｜|\.．]|#+\s+|第\d+章)")
NUMERIC_RE = re.compile(r"\d")
CONDITION_WORDS = ("場合", "条件", "異なり", "異なります", "限り", "のみ", "※", "確定している", "状況により")
EDITORIAL_WORDS = ("省略可能", "省略可", "追加：", "追加:", "メモ", "制作指示", "要検討")
VISUAL_WORDS = ("ビジュアル", "写真", "画像", "グラフ", "図解", "イラスト", "スクリーンショット")
COMPARISON_WORDS = ("比較", "より", "違い", "○", "△", "×")
PROCESS_WORDS = ("公開する", "受け付ける", "見届ける", "抽選する", "ステップ", "流れ", "手順")


def extract_docx(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    from docx import Document
    doc = Document(path)
    paragraphs: list[dict[str, Any]] = []
    for idx, p in enumerate(doc.paragraphs, start=1):
        text = p.text.strip()
        if text:
            paragraphs.append({"index": idx, "text": text, "style": p.style.name if p.style else None})
    tables = []
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": t_idx, "rows": rows})
    return paragraphs, tables


def extract_text(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    paragraphs = [{"index": i, "text": line.strip(), "style": None}
                  for i, line in enumerate(lines, start=1) if line.strip()]
    return paragraphs, []


def classify(text: str) -> tuple[str, list[str], float, list[str]]:
    reasons: list[str] = []
    secondary: list[str] = []
    if any(w in text for w in EDITORIAL_WORDS):
        reasons.append("editorial-marker")
        return "editorial_instruction", secondary, 0.96, reasons
    if HEADING_RE.search(text):
        reasons.append("heading-pattern")
        return "heading", secondary, 0.93, reasons
    if any(w in text for w in CONDITION_WORDS):
        reasons.append("condition-marker")
        if any(w in text for w in VISUAL_WORDS):
            secondary.append("visual_instruction")
        return "conditional_claim", secondary, 0.88, reasons
    if any(w in text for w in PROCESS_WORDS):
        reasons.append("process-marker")
        return "process", secondary, 0.82, reasons
    if any(w in text for w in VISUAL_WORDS) and ("追加" in text or "〜" in text or "する" in text):
        reasons.append("visual-marker")
        return "visual_instruction", secondary, 0.78, reasons
    if any(w in text for w in COMPARISON_WORDS):
        reasons.append("comparison-marker")
        secondary.append("customer_copy")
        return "comparison", secondary, 0.72, reasons
    if NUMERIC_RE.search(text):
        reasons.append("contains-number")
        secondary.append("customer_copy")
        return "numerical_data", secondary, 0.66, reasons
    return "customer_copy", secondary, 0.60, ["default-customer-copy"]


def build_sections(paragraphs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sections = []
    current = {"heading": None, "items": []}
    for p in paragraphs:
        primary, secondary, confidence, reasons = classify(p["text"])
        item = {
            "paragraph_index": p["index"],
            "source_text": p["text"],
            "classification": {"primary": primary, "secondary": secondary},
            "confidence": confidence,
            "reasons": reasons,
        }
        if primary == "heading":
            if current["heading"] or current["items"]:
                sections.append(current)
            current = {"heading": p["text"], "items": []}
        else:
            current["items"].append(item)
    if current["heading"] or current["items"]:
        sections.append(current)
    return sections


def analyze(manuscript: str | Path, output: str | Path) -> dict[str, Any]:
    path = Path(manuscript)
    if path.suffix.lower() == ".docx":
        paragraphs, tables = extract_docx(path)
    else:
        paragraphs, tables = extract_text(path)

    sections = build_sections(paragraphs)
    content_units = []
    unit_no = 1
    for section in sections:
        unit_items = section["items"]
        source_text = "\n".join(i["source_text"] for i in unit_items)
        classification_counts: dict[str, int] = {}
        for i in unit_items:
            k = i["classification"]["primary"]
            classification_counts[k] = classification_counts.get(k, 0) + 1
        primary = max(classification_counts, key=classification_counts.get) if classification_counts else "heading"
        content_units.append({
            "id": f"CU-{unit_no:03d}",
            "heading": section["heading"],
            "classification": {"primary": primary},
            "source_text": source_text,
            "items": unit_items,
            "importance": "primary" if unit_no <= 7 else "secondary",
            "status": "DRAFT",
        })
        unit_no += 1

    result = {
        "schema_version": "0.1",
        "generated_at": now_iso(),
        "source": str(path),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "content_units": content_units,
        "tables": tables,
        "notes": [
            "This is deterministic seed analysis. Codex should refine semantic classification before approval.",
            "source_text must remain immutable; derived semantic/render text belongs downstream."
        ]
    }
    write_json(output, result)
    return result
