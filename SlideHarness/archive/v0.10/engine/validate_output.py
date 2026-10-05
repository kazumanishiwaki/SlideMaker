from __future__ import annotations
from pathlib import Path
from typing import Any
import json
import re
import math
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import unicodedata

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageChops, ImageDraw

from helpers import now_iso, read_json, write_json
from render_source import render_pptx, _run_bounded
from image_pipeline import validate_image_harmony
from japanese_typography import diagnose_rendered_lines

EMU = 914400
OBJ_RE = re.compile(r'-O(\d+)$')


TECHNICAL_CHECK_PREFIXES = (
    'SLIDE_', 'TEMPLATE_', 'MUTATED_', 'LAYOUT_GENERATED_', 'LIBREOFFICE_',
    'UNEXPECTED_PIXEL_', 'RENDER_TEXT_OVERFLOW', 'RENDER_LAYOUT_TEXT_OVERFLOW',
    'RENDER_TEXT_OVERLAP', 'PDF_TEXT_', 'PHOTO_LAYOUT_', 'FULL_BLEED_', 'IMAGE_',
    'EDITABLE_',
)
CONTENT_CHECK_PREFIXES = (
    'RENDER_TEXT_PRESENT', 'RENDER_LAYOUT_TEXT_PRESENT', 'RENDER_TABLE_TEXT_PRESENT',
    'NUMBER_UNIT_', 'LATIN_TERM_', 'SEMANTIC_LAYOUT_FORBIDDEN',
)
DESIGN_CHECK_PREFIXES = (
    'LAYOUT_FIT_', 'SEMANTIC_LAYOUT_MISMATCH', 'JAPANESE_', 'TITLE_',
    'MINIMUM_PRESENTATION_FONT',
)

ROLE_MIN_FONT_PT = {
    'title': 28.0, 'subtitle': 16.0, 'item_title': 18.0, 'item_body': 14.0,
    'body': 15.0, 'number': 12.0, 'kpi_value': 28.0, 'caption': 10.0,
    'note': 10.0, 'kicker': 12.0,
}

def _qa_category(code: str) -> str:
    code = str(code or '')
    if code.startswith(CONTENT_CHECK_PREFIXES):
        return 'content'
    if code.startswith(DESIGN_CHECK_PREFIXES):
        return 'design'
    if code.startswith(TECHNICAL_CHECK_PREFIXES):
        return 'technical'
    return 'technical'

def _shape_font_sizes(shape) -> list[float]:
    values: list[float] = []
    if not getattr(shape, 'has_text_frame', False):
        return values
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            try:
                if run.font.size is not None:
                    values.append(float(run.font.size.pt))
            except Exception:
                continue
    return values


def _shape_id(object_id: str) -> int | None:
    m=OBJ_RE.search(object_id or '')
    return int(m.group(1)) if m else None


def _norm(s: str) -> str:
    # PDF exporters may substitute compatibility glyphs (for example ｜ -> |).
    # NFKC keeps presence checks stable without changing the semantic source text.
    normalized = unicodedata.normalize('NFKC', str(s or '')).replace('\u00a0', ' ')
    return re.sub(r'\s+', '', normalized)


def _extract_pdf_text(pptx: Path, workdir: Path) -> tuple[list[str], str | None]:
    lo=shutil.which('libreoffice') or shutil.which('soffice')
    pdftotext=shutil.which('pdftotext')
    if not lo or not pdftotext: return [], None
    pdfdir=workdir/'_qa_pdf'
    shutil.rmtree(pdfdir,ignore_errors=True)
    pdfdir.mkdir(parents=True,exist_ok=True)
    try:
        # Isolate the LibreOffice profile for every conversion. Without this,
        # concurrent App Server requests can block on the shared profile lock and
        # leave the QA turn hanging indefinitely.
        with tempfile.TemporaryDirectory(prefix='ppt-qa-lo-profile-') as profile:
            profile_arg=f'-env:UserInstallation={Path(profile).resolve().as_uri()}'
            _run_bounded([
                lo,profile_arg,'--headless','--nologo','--nodefault','--nofirststartwizard',
                '--convert-to','pdf','--outdir',str(pdfdir),str(pptx)
            ],timeout=120)
        pdf=pdfdir/(pptx.stem+'.pdf')
        if not pdf.exists():
            candidates=list(pdfdir.glob('*.pdf'))
            if not candidates: return [], None
            pdf=candidates[0]
        proc=_run_bounded([
            pdftotext,'-layout',str(pdf),'-'
        ],timeout=60)
        text=proc.stdout.decode('utf-8','replace')
        pages=text.split('\f')
        if pages and not pages[-1].strip(): pages=pages[:-1]
        return pages, str(pdf)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f'PDF text extraction timed out after {exc.timeout}s: {pptx.name}') from exc



def _extract_pdf_words(pdf: Path) -> list[dict[str, Any]]:
    """Return rendered word boxes from Poppler in PDF point coordinates."""
    pdftotext=shutil.which('pdftotext')
    if not pdftotext: return []
    try:
        proc=_run_bounded(
            [pdftotext,'-bbox-layout',str(pdf),'-'],timeout=60,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f'PDF word-box extraction timed out after {exc.timeout}s: {pdf.name}') from exc
    root=ET.fromstring(proc.stdout)
    pages=[]
    for page in [x for x in root.iter() if x.tag.endswith('page')]:
        words=[]
        for w in [x for x in page.iter() if x.tag.endswith('word')]:
            text=''.join(w.itertext())
            try:
                words.append({
                    'text':text,
                    'xMin':float(w.attrib['xMin']),'yMin':float(w.attrib['yMin']),
                    'xMax':float(w.attrib['xMax']),'yMax':float(w.attrib['yMax']),
                })
            except Exception:
                continue
        pages.append({'width':float(page.attrib.get('width',0)),'height':float(page.attrib.get('height',0)),'words':words})
    return pages


def _box_distance(word: dict[str,Any], box: tuple[float,float,float,float]) -> float:
    # Distance to the *rectangle*, not its center. Text in a wide textbox is often
    # left-aligned; center distance would incorrectly prefer duplicate words in a
    # table below the box over the correct word near the box's left edge.
    cx=(word['xMin']+word['xMax'])/2; cy=(word['yMin']+word['yMax'])/2
    dx=box[0]-cx if cx < box[0] else (cx-box[2] if cx > box[2] else 0.0)
    dy=box[1]-cy if cy < box[1] else (cy-box[3] if cy > box[3] else 0.0)
    return dx*dx+dy*dy


def _words_near_box(words: list[dict[str,Any]], box: tuple[float,float,float,float], pad: float = 24.0) -> list[dict[str,Any]]:
    """Return only rendered words close enough to plausibly belong to a shape.

    PDF text extraction may return hundreds of words per page.  Matching every
    textbox against the whole page is both slow and vulnerable to duplicates in
    tables or repeated labels.  Restricting candidates to an expanded shape
    rectangle keeps validation deterministic and makes runtime roughly linear in
    the amount of text actually rendered near each object.
    """
    x0,y0,x1,y1=box
    return [w for w in words if not (w.get('xMax',0) < x0-pad or w.get('xMin',0) > x1+pad or w.get('yMax',0) < y0-pad or w.get('yMin',0) > y1+pad)]


def _match_rendered_words(words: list[dict[str,Any]], expected: str, shape_box: tuple[float,float,float,float]) -> list[dict[str,Any]] | None:
    """Match expected text to rendered PDF word boxes independent of flow order."""
    target=_norm(expected)
    if not target: return None
    chosen=[]; pos=0; used=set()
    while pos < len(target):
        candidates=[]
        for i,w in enumerate(words):
            if i in used: continue
            t=_norm(w.get('text',''))
            if not t: continue
            if target.startswith(t,pos):
                candidates.append((i,w,t))
        if not candidates:
            remain=target[pos:]
            for i,w in enumerate(words):
                if i in used: continue
                t=_norm(w.get('text',''))
                if remain and remain in t:
                    candidates.append((i,w,remain))
        if not candidates: return None
        candidates.sort(key=lambda x: (_box_distance(x[1],shape_box), -len(x[2])))
        i,w,t=candidates[0]; used.add(i); chosen.append(w); pos += len(t)
    return chosen


def _match_rendered_bbox(words: list[dict[str,Any]], expected: str, shape_box: tuple[float,float,float,float]) -> tuple[float,float,float,float] | None:
    chosen=_match_rendered_words(words,expected,shape_box)
    if not chosen: return None
    return (min(w['xMin'] for w in chosen), min(w['yMin'] for w in chosen), max(w['xMax'] for w in chosen), max(w['yMax'] for w in chosen))


def _group_rendered_lines(words: list[dict[str,Any]]) -> list[list[dict[str,Any]]]:
    if not words: return []
    ordered=sorted(words,key=lambda w:((w['yMin']+w['yMax'])/2,w['xMin']))
    lines=[]
    for w in ordered:
        cy=(w['yMin']+w['yMax'])/2; h=max(1.0,w['yMax']-w['yMin'])
        best=None; best_delta=None
        for line in lines:
            lcy=sum((x['yMin']+x['yMax'])/2 for x in line)/len(line)
            delta=abs(cy-lcy)
            threshold=max(3.0,min(h, sum(max(1.0,x['yMax']-x['yMin']) for x in line)/len(line))*0.35)
            if delta <= threshold and (best_delta is None or delta<best_delta): best=line; best_delta=delta
        if best is None: lines.append([w])
        else: best.append(w)
    lines.sort(key=lambda line:min(x['yMin'] for x in line))
    for line in lines: line.sort(key=lambda w:w['xMin'])
    return lines


def _line_text(line: list[dict[str,Any]]) -> str:
    # Spaces help preserve Latin multi-word terms for split diagnostics.
    return ' '.join(str(w.get('text','')) for w in line).strip()


def _wrap_issues(lines: list[list[dict[str,Any]]], render_text: str, role: str | None) -> list[tuple[str,str,dict[str,Any]]]:
    return diagnose_rendered_lines(lines, render_text, role)




def _pixel_diff_outside_mutations(source_png: Path, output_png: Path, srcslide: dict[str,Any], plan: dict[str,Any], slide_width_in: float, slide_height_in: float, debug_path: Path | None = None, layout_plan: dict[str,Any] | None = None) -> dict[str,Any]:
    """Compare rendered source/output pixels outside semantically allowed mutation regions.

    Region-based masking intentionally derives from source object geometry rather than
    model-generated coordinates. For replace_table/chart region replacement, the union
    rectangle of the declared source region is allowed because the new editable object
    may legitimately occupy gaps between the old source objects.
    """
    if not source_png.exists() or not output_png.exists():
        return {'available':False}
    src=Image.open(source_png).convert('RGB')
    out=Image.open(output_png).convert('RGB')
    if src.size != out.size:
        return {'available':False,'reason':'size_mismatch','source_size':src.size,'output_size':out.size}
    objects={o.get('source_object_id'):o for o in srcslide.get('objects',[])}
    mask=Image.new('1',src.size,0)
    draw=ImageDraw.Draw(mask)
    sx=src.width/max(1e-9,slide_width_in); sy=src.height/max(1e-9,slide_height_in)
    # About 1/8 inch at a 96dpi render. Allows antialias/shadow spill without
    # weakening the invariant over the rest of the slide.
    pad=max(6,int(round(min(sx,sy)*0.125)))

    def draw_bbox(bb: dict[str,Any], extra_pad: int = 0):
        local_pad=pad+max(0,int(extra_pad))
        x0=float(bb.get('x',0))*sx-local_pad; y0=float(bb.get('y',0))*sy-local_pad
        x1=(float(bb.get('x',0))+float(bb.get('w',0)))*sx+local_pad
        y1=(float(bb.get('y',0))+float(bb.get('h',0)))*sy+local_pad
        draw.rectangle([x0,y0,x1,y1],fill=1)

    def mutation_bbox(obj: dict[str,Any]) -> dict[str,float] | None:
        bb=obj.get('bbox')
        if not bb: return None
        # PowerPoint reports an unrotated shape rectangle, while rendered pixels can
        # extend beyond it. Expand to the axis-aligned bounds of the rotated rectangle
        # so an in-place image swap does not look like an unauthorized mutation.
        display=(obj.get('picture') or {}).get('display') or {}
        try: rotation=float(display.get('rotation') or obj.get('rotation') or 0.0)
        except Exception: rotation=0.0
        if abs(rotation) < .01:
            return bb
        x=float(bb.get('x',0)); y=float(bb.get('y',0)); w=float(bb.get('w',0)); h=float(bb.get('h',0))
        cx=x+w/2; cy=y+h/2; rad=math.radians(rotation)
        points=[]
        for px,py in ((x,y),(x+w,y),(x+w,y+h),(x,y+h)):
            dx,dy=px-cx,py-cy
            points.append((cx+dx*math.cos(rad)-dy*math.sin(rad), cy+dx*math.sin(rad)+dy*math.cos(rad)))
        minx=min(v[0] for v in points); maxx=max(v[0] for v in points)
        miny=min(v[1] for v in points); maxy=max(v[1] for v in points)
        return {'x':minx,'y':miny,'w':maxx-minx,'h':maxy-miny}

    for mut in plan.get('source_object_mutations',[]):
        region=mut.get('replace_region_object_ids') or []
        region_bboxes=[objects[x].get('bbox') for x in region if x in objects and objects[x].get('bbox')]
        if region_bboxes:
            minx=min(float(b['x']) for b in region_bboxes); miny=min(float(b['y']) for b in region_bboxes)
            maxx=max(float(b['x'])+float(b['w']) for b in region_bboxes)
            maxy=max(float(b['y'])+float(b['h']) for b in region_bboxes)
            draw_bbox({'x':minx,'y':miny,'w':maxx-minx,'h':maxy-miny})
            continue
        obj=objects.get(mut.get('source_object_id'))
        mutation_region=mutation_bbox(obj) if obj else None
        if obj and mutation_region:
            # Some source templates visually overflow their declared PowerPoint text
            # box before adaptation. Permit a restrained extra halo for declared
            # text mutations so removal of that pre-existing overflow is not treated
            # as an unrelated slide mutation. Non-text and undeclared regions remain
            # protected by the normal tight mask.
            declared_margin=max(0.0,min(0.75,float(mut.get('pixel_diff_margin_in') or 0.0)))
            default_text_margin=0.22 if mut.get('action') in {'replace_text','clear_text'} else 0.0
            text_halo=int(round(min(sx,sy)*(default_text_margin+declared_margin)))
            image_halo=int(round(min(sx,sy)*0.08)) if mut.get('action') == 'replace_image' else 0
            draw_bbox(mutation_region, text_halo + image_halo)

    for oid in (layout_plan or {}).get('removed_source_object_ids',[]):
        obj=objects.get(oid)
        if not obj or not obj.get('bbox'):
            continue
        # PowerPoint/LibreOffice can render glyphs slightly outside the declared
        # text-frame rectangle (descenders, font substitution, line-height, or a
        # source box whose text already overflowed). Removing that declared text is
        # an authorized mutation, so give removed text frames the same restrained
        # halo used by explicit replace_text/clear_text mutations. Pictures and
        # other shapes keep the tighter default padding.
        text_halo=int(round(min(sx,sy)*0.22)) if str(obj.get('text') or '').strip() else 0
        draw_bbox(obj['bbox'], text_halo)
    for region in (layout_plan or {}).get('allowed_regions',[]):
        if isinstance(region,dict): draw_bbox(region)

    diff=ImageChops.difference(src,out)
    # Threshold antialias noise. Count only clearly visible channel deviations.
    # A 16:9 validation render contains several million pixels; iterating over every
    # channel in Python made multi-slide export QA take minutes. Use NumPy when
    # available and keep the scalar implementation as a portability fallback.
    threshold=16
    try:
        import numpy as np

        diff_array=np.asarray(diff,dtype=np.uint8)
        mask_array=np.asarray(mask,dtype=bool)
        max_delta_map=diff_array.max(axis=2)
        checked=~mask_array
        total=int(np.count_nonzero(checked))
        changed_map=checked & (max_delta_map > threshold)
        changed=int(np.count_nonzero(changed_map))
        if total:
            checked_values=max_delta_map[checked]
            sum_delta=int(checked_values.sum(dtype=np.uint64))
            max_delta=int(checked_values.max(initial=0))
        else:
            sum_delta=0; max_delta=0
        if debug_path:
            visual=np.full((src.height,src.width,3),255,dtype=np.uint8)
            visual[mask_array]=(235,235,235)
            visual[changed_map]=(0,0,0)
            debug_path.parent.mkdir(parents=True,exist_ok=True)
            Image.fromarray(visual,'RGB').save(debug_path)
    except ImportError:
        dpx=diff.load(); mpx=mask.load(); total=changed=0; max_delta=0; sum_delta=0
        vis=Image.new('RGB',src.size,(255,255,255)) if debug_path else None
        vpx=vis.load() if vis else None
        for y in range(src.height):
            for x in range(src.width):
                if mpx[x,y]:
                    if vpx: vpx[x,y]=(235,235,235)
                    continue
                total += 1
                delta=max(dpx[x,y]); sum_delta += delta; max_delta=max(max_delta,delta)
                if delta > threshold:
                    changed += 1
                    if vpx: vpx[x,y]=(0,0,0)
        if debug_path and vis:
            debug_path.parent.mkdir(parents=True,exist_ok=True); vis.save(debug_path)
    ratio=(changed/total) if total else 0.0
    return {
        'available':True,'changed_pixels':changed,'checked_pixels':total,'changed_ratio':ratio,
        'mean_max_channel_delta':(sum_delta/total) if total else 0.0,'max_channel_delta':max_delta,
        'threshold':threshold,'mask_padding_px':pad,'debug':str(debug_path) if debug_path else None,
    }

def _intersection_area(a: tuple[float,float,float,float], b: tuple[float,float,float,float]) -> float:
    w=max(0.0,min(a[2],b[2])-max(a[0],b[0])); h=max(0.0,min(a[3],b[3])-max(a[1],b[1]))
    return w*h


def _shape_bbox_in(shape) -> dict[str,float]:
    return {
        'x': float(shape.left) / EMU,
        'y': float(shape.top) / EMU,
        'w': float(shape.width) / EMU,
        'h': float(shape.height) / EMU,
    }


def _bbox_delta(actual: dict[str,float], expected: dict[str,Any]) -> float:
    try:
        return max(abs(float(actual[key]) - float(expected[key])) for key in ('x','y','w','h'))
    except Exception:
        return float('inf')


def _photo_layout_checks(add, *, sid: str, slide, out_obj: dict[int,Any], layout_plan: dict[str,Any], chrome_assets: list[dict[str,Any]], sw: int, sh: int) -> None:
    family=str(layout_plan.get('family') or '')
    if family not in {'photo_text','full_bleed'}:
        return
    role='background_visual' if family == 'full_bleed' else 'primary_visual'
    generated=next((row for row in layout_plan.get('generated_objects',[]) if row.get('role') == role),None)
    slots=layout_plan.get('image_slots') or []
    slot=slots[0] if slots else None
    if generated is None or slot is None:
        add('PHOTO_LAYOUT_PRIMARY_VISUAL','fail','blocking',slide_id=sid,family=family,reason='generated_visual_or_slot_missing')
        return
    shape=out_obj.get(generated.get('output_shape_id'))
    is_picture=bool(shape is not None and getattr(shape,'shape_type',None) == MSO_SHAPE_TYPE.PICTURE)
    add('PHOTO_LAYOUT_PRIMARY_VISUAL','pass' if is_picture else 'fail','info' if is_picture else 'blocking',
        slide_id=sid,family=family,shape_id=generated.get('output_shape_id'))
    if not is_picture:
        return
    actual=_shape_bbox_in(shape); expected=slot.get('bbox') or {}
    delta=_bbox_delta(actual,expected)
    add('PHOTO_LAYOUT_SLOT_GEOMETRY','pass' if delta <= 0.035 else 'fail','info' if delta <= 0.035 else 'blocking',
        slide_id=sid,family=family,max_delta_in=round(delta,4),actual=actual,expected=expected)

    source_object_id=slot.get('source_object_id')
    source_shape_id=_shape_id(source_object_id) if source_object_id else None
    reuse_requested=bool(slot.get('reuse_source_image_frame'))
    reused=bool(reuse_requested and source_shape_id is not None and shape.shape_id == source_shape_id and generated.get('source_picture_reused'))
    if reuse_requested:
        add('PHOTO_LAYOUT_SOURCE_FRAME_REUSED','pass' if reused else 'fail','info' if reused else 'blocking',
            slide_id=sid,source_object_id=source_object_id,output_shape_id=shape.shape_id)

    geom=layout_plan.get('geometry') or {}
    image_box=geom.get('image_box') or expected
    text_box=geom.get('text_box') or {}
    side=(layout_plan.get('layout_intent') or {}).get('resolved_text_side') or geom.get('resolved_text_side')
    directional_ok=True
    if family == 'photo_text' and image_box and text_box:
        eps=0.035
        if side == 'right': directional_ok=float(text_box.get('x',0)) >= float(image_box.get('x',0))+float(image_box.get('w',0))-eps
        elif side == 'left': directional_ok=float(text_box.get('x',0))+float(text_box.get('w',0)) <= float(image_box.get('x',0))+eps
        elif side == 'bottom': directional_ok=float(text_box.get('y',0)) >= float(image_box.get('y',0))+float(image_box.get('h',0))-eps
        elif side == 'top': directional_ok=float(text_box.get('y',0))+float(text_box.get('h',0)) <= float(image_box.get('y',0))+eps
        add('PHOTO_TEXT_NEGATIVE_SPACE_DIRECTION','pass' if directional_ok else 'fail','info' if directional_ok else 'blocking',
            slide_id=sid,resolved_text_side=side,image_box=image_box,text_box=text_box)

    if family == 'full_bleed':
        tolerance=0.06
        covers=actual['x'] <= tolerance and actual['y'] <= tolerance and actual['x']+actual['w'] >= sw/EMU-tolerance and actual['y']+actual['h'] >= sh/EMU-tolerance
        add('FULL_BLEED_COVERAGE','pass' if covers else 'fail','info' if covers else 'blocking',
            slide_id=sid,actual=actual,slide={'w':sw/EMU,'h':sh/EMU})
        order=[row.shape_id for row in slide.shapes]
        picture_index=order.index(shape.shape_id) if shape.shape_id in order else len(order)
        chrome_ids=[_shape_id(row.get('source_object_id')) for row in chrome_assets]
        relevant=[cid for cid in chrome_ids if cid is not None and cid in order]
        above=all(picture_index < order.index(cid) for cid in relevant) if relevant else True
        add('FULL_BLEED_BEHIND_TEMPLATE_CHROME','pass' if above else 'fail','info' if above else 'blocking',
            slide_id=sid,picture_shape_id=shape.shape_id,chrome_shape_ids=relevant)


def _editable_layout_checks(add, *, sid: str, out_obj: dict[int,Any], layout_plan: dict[str,Any]) -> None:
    family=str(layout_plan.get('family') or '')
    generated=layout_plan.get('generated_objects') or []
    if family == 'comparison_matrix':
        matrix=next((row for row in generated if row.get('role') == 'comparison_matrix'),None)
        shape=out_obj.get((matrix or {}).get('output_shape_id')) if matrix else None
        editable=bool(shape is not None and getattr(shape,'has_table',False))
        add('EDITABLE_LAYOUT_COMPARISON_MATRIX','pass' if editable else 'fail','info' if editable else 'blocking',
            slide_id=sid,shape_id=(matrix or {}).get('output_shape_id'))
    if family == 'relationship':
        node_rows=[row for row in generated if row.get('role') == 'relationship_node']
        edge_rows=[row for row in generated if row.get('role') == 'relationship_edge']
        node_shapes=[out_obj.get(row.get('output_shape_id')) for row in node_rows]
        edge_shapes=[out_obj.get(row.get('output_shape_id')) for row in edge_rows]
        nodes_native=bool(node_rows) and all(shape is not None and getattr(shape,'shape_type',None) == MSO_SHAPE_TYPE.AUTO_SHAPE for shape in node_shapes)
        edges_native=bool(edge_rows) and all(shape is not None and getattr(shape,'shape_type',None) == MSO_SHAPE_TYPE.LINE for shape in edge_shapes)
        add('EDITABLE_RELATIONSHIP_NODES','pass' if nodes_native else 'fail','info' if nodes_native else 'blocking',
            slide_id=sid,node_count=len(node_rows))
        add('EDITABLE_RELATIONSHIP_CONNECTORS','pass' if edges_native else 'fail','info' if edges_native else 'blocking',
            slide_id=sid,edge_count=len(edge_rows))

def validate(project: str | Path, pptx_rel: str, *, kind='draft') -> dict[str, Any]:
    root=Path(project).resolve(); pptx=(root/pptx_rel).resolve()
    if not str(pptx).startswith(str(root)+'/'): raise ValueError('pptx outside project')
    if not pptx.exists(): raise FileNotFoundError(pptx)
    prs=Presentation(pptx)
    deck=read_json(root/'planning/deck_plan.json',{})
    matches=read_json(root/'planning/slide_matching.json',{})
    assets=read_json(root/'analysis/source_assets.json',{})
    source=read_json(root/'analysis/source_slides.json',{})
    template=read_json(root/'analysis/template_profile.json',{})
    deck_by_order=sorted(deck.get('slides',[]),key=lambda x:x.get('order',9999))
    match_by={x['slide_id']:x for x in matches.get('slides',[])}
    src_by={x['id']:x for x in source.get('slides',[])}
    chrome_by_src={}
    for a in assets.get('assets',[]):
        if a.get('classification')=='template_chrome': chrome_by_src.setdefault(a['source_slide_id'],[]).append(a)

    # For subset sample decks, recover slide ids from compose report when available.
    report_path=root/'output'/(pptx.stem+'_compose_report.json')
    compose_report=read_json(report_path,{})
    output_ids=[x.get('slide_id') for x in compose_report.get('slides',[]) if x.get('slide_id')]
    if not output_ids:
        output_ids=[x['slide_id'] for x in deck_by_order[:len(prs.slides)]]

    checks=[]
    def add(code,result,severity='blocking',**kw):
        checks.append({'check_id':code,'result':result,'severity':severity,**kw})

    if len(output_ids)!=len(prs.slides): add('SLIDE_COUNT_MAP','fail','blocking',expected=len(output_ids),actual=len(prs.slides))
    else: add('SLIDE_COUNT_MAP','pass','info',actual=len(prs.slides))

    sw,sh=prs.slide_width,prs.slide_height
    for idx, slide in enumerate(prs.slides):
        sid=output_ids[idx] if idx<len(output_ids) else f'PAGE-{idx+1}'
        match=match_by.get(sid,{})
        sel=match.get('selected') or (match.get('candidates') or [None])[0]
        srcid=sel.get('source_slide_id') if sel else None
        srcslide=src_by.get(srcid,{})
        src_obj={o.get('shape_id'):o for o in srcslide.get('objects',[])}
        out_obj={s.shape_id:s for s in slide.shapes}
        plan=read_json(root/f'planning/slide_plans/{sid}.json',{})
        mutations=plan.get('source_object_mutations',[])
        mutated_ids={_shape_id(m.get('source_object_id')) for m in mutations}

        # Template chrome must survive unchanged unless user has explicitly changed its classification upstream.
        for a in chrome_by_src.get(srcid,[]):
            shp_id=_shape_id(a.get('source_object_id'))
            if shp_id not in out_obj:
                add('TEMPLATE_CHROME_PRESENT','fail','blocking',slide_id=sid,source_object_id=a.get('source_object_id'))
                continue
            src=src_obj.get(shp_id); out=out_obj[shp_id]
            if src:
                bbox=src['bbox']; expected=[round(bbox[k]*EMU) for k in ['x','y','w','h']]
                actual=[out.left,out.top,out.width,out.height]
                # Analyzer bboxes are rounded to 4 decimals, so allow the corresponding
                # sub-point rounding error while still requiring effectively lossless cloning.
                if any(abs(a0-e0)>150 for a0,e0 in zip(actual,expected)):
                    add('TEMPLATE_CHROME_GEOMETRY','fail','blocking',slide_id=sid,source_object_id=a.get('source_object_id'))
                else:
                    add('TEMPLATE_CHROME_GEOMETRY','pass','info',slide_id=sid,source_object_id=a.get('source_object_id'))

        # Mutated source text boxes should still be in slide bounds. Existing source decorations may intentionally bleed.
        for mid in mutated_ids:
            shp=out_obj.get(mid)
            if shp is None: continue
            if shp.left < 0 or shp.top < 0 or shp.left+shp.width > sw or shp.top+shp.height > sh:
                add('MUTATED_OBJECT_OUT_OF_BOUNDS','fail','blocking',slide_id=sid,shape_id=mid)
            else: add('MUTATED_OBJECT_OUT_OF_BOUNDS','pass','info',slide_id=sid,shape_id=mid)

        layout_plan=read_json(root/f'render/layout_plans/{sid}.json',{})
        if layout_plan:
            q=layout_plan.get('quality') or {}
            policy=layout_plan.get('fill_policy') or q.get('fill_policy') or {}
            add('LAYOUT_FIT_TO_INTENT','pass' if q.get('status')=='pass' else 'fail','blocking' if q.get('status')!='pass' else 'info',
                slide_id=sid,family=layout_plan.get('family'),score=q.get('score'),assessment=q.get('assessment'),
                fill_policy=policy.get('mode'),fill_policy_source=policy.get('source'),
                intentional_whitespace=q.get('intentional_whitespace'),accidental_dead_space=q.get('accidental_dead_space'),
                reason_codes=q.get('reason_codes'),content_extent_width_ratio=q.get('content_extent_width_ratio'),
                content_extent_height_ratio=q.get('content_extent_height_ratio'),occupied_area_ratio=q.get('occupied_area_ratio'),
                center_offset_x_ratio=q.get('center_offset_x_ratio'),center_offset_y_ratio=q.get('center_offset_y_ratio'),thresholds=q.get('thresholds'))
            for issue in layout_plan.get('semantic_layout_issues') or []:
                severity = 'blocking' if issue.get('severity') == 'blocking' else 'review'
                add(issue.get('check_id') or 'SEMANTIC_LAYOUT_MISMATCH', 'fail', severity,
                    slide_id=sid, relation=issue.get('relation'), family=issue.get('family'),
                    allowed_families=issue.get('allowed_families'))
            for obj in layout_plan.get('generated_objects',[]):
                oid=obj.get('output_shape_id'); shp=out_obj.get(oid)
                if shp is None:
                    add('LAYOUT_GENERATED_OBJECT_PRESENT','fail','blocking',slide_id=sid,shape_id=oid,role=obj.get('role')); continue
                # Layout coordinates are stored in inches and converted back to
                # integer EMUs. A full-bleed width can differ from the exact slide
                # width by a few dozen EMUs due to decimal rounding; allow the same
                # sub-point tolerance used for cloned template chrome.
                tolerance_emu = 150
                bad=(
                    shp.left < -tolerance_emu or shp.top < -tolerance_emu
                    or shp.left + shp.width > sw + tolerance_emu
                    or shp.top + shp.height > sh + tolerance_emu
                )
                add('LAYOUT_GENERATED_OBJECT_BOUNDS','fail' if bad else 'pass','blocking' if bad else 'info',slide_id=sid,shape_id=oid,role=obj.get('role'))
                role = str(obj.get('role') or '')
                minimum = ROLE_MIN_FONT_PT.get(role)
                sizes = _shape_font_sizes(shp)
                if minimum is not None and sizes:
                    actual_min = min(sizes)
                    add('MINIMUM_PRESENTATION_FONT', 'pass' if actual_min + 0.05 >= minimum else 'fail',
                        'info' if actual_min + 0.05 >= minimum else 'review',
                        slide_id=sid, shape_id=oid, role=role, actual_pt=round(actual_min,2), minimum_pt=minimum)
            _photo_layout_checks(
                add,sid=sid,slide=slide,out_obj=out_obj,layout_plan=layout_plan,
                chrome_assets=chrome_by_src.get(srcid,[]),sw=sw,sh=sh,
            )
            _editable_layout_checks(add,sid=sid,out_obj=out_obj,layout_plan=layout_plan)

    render_dir=root/'render'/'output_validation'/pptx.stem
    shutil.rmtree(render_dir,ignore_errors=True)
    rendered=[]
    try:
        # Reuse the renderer; names are source-like but directory identifies output deck.
        rendered=render_pptx(pptx,render_dir)
        add('LIBREOFFICE_RENDER','pass' if len(rendered)==len(prs.slides) else 'fail','blocking',rendered_pages=len(rendered))
    except Exception as e:
        add('LIBREOFFICE_RENDER','fail','blocking',error=str(e))

    # Pixel-level guard: outside the source-object mutation regions, output should
    # render identically (within tiny antialias noise). This catches accidental
    # movement/deletion of unrelated template/content objects during partial edits.
    pres=template.get('presentation',{})
    slide_width_in=float(pres.get('width_in') or (sw/EMU))
    slide_height_in=float(pres.get('height_in') or (sh/EMU))
    # Analysis normally creates these renders. If an earlier conversion was
    # interrupted or the project was moved without render caches, rebuild them
    # once here instead of degrading the pixel guard to a warning.
    required_source_ids=[]
    for sid in output_ids:
        preplan=read_json(root/f'planning/slide_plans/{sid}.json',{})
        presrc=preplan.get('source_slide_id')
        if not presrc:
            prematch=match_by.get(sid,{})
            presel=prematch.get('selected') or (prematch.get('candidates') or [None])[0]
            presrc=presel.get('source_slide_id') if presel else None
        if presrc:
            required_source_ids.append(presrc)
    source_render_dir=root/'render'/'source_renders'
    if any(not (source_render_dir/f'{srcid}.png').exists() for srcid in required_source_ids):
        source_pptx=root/'inputs'/'source.pptx'
        if source_pptx.exists():
            try:
                render_pptx(source_pptx,source_render_dir)
            except Exception:
                # The normal warning path below preserves observability if the
                # recovery render is unavailable in the current environment.
                pass

    for idx,sid in enumerate(output_ids):
        plan=read_json(root/f'planning/slide_plans/{sid}.json',{})
        srcid=plan.get('source_slide_id')
        if not srcid:
            match=match_by.get(sid,{})
            sel=match.get('selected') or (match.get('candidates') or [None])[0]
            srcid=sel.get('source_slide_id') if sel else None
        srcslide=src_by.get(srcid,{})
        source_png=root/'render'/'source_renders'/f'{srcid}.png' if srcid else Path('__missing__')
        output_png=Path(rendered[idx]) if idx < len(rendered) else Path('__missing__')
        debug=render_dir/'diff'/f'{sid}.png'
        layout_plan=read_json(root/f'render/layout_plans/{sid}.json',{})
        info=_pixel_diff_outside_mutations(source_png,output_png,srcslide,plan,slide_width_in,slide_height_in,debug,layout_plan=layout_plan)
        if not info.get('available'):
            add('UNEXPECTED_PIXEL_DIFF','fail','warning',slide_id=sid,source_slide_id=srcid,reason=info.get('reason','render_missing'))
            continue
        # A small count tolerance protects against renderer-specific isolated pixels;
        # a meaningful mutation outside the mask will exceed both gates rapidly.
        # LibreOffice may rasterize unchanged text with slightly different glyph
        # antialiasing when a single cloned slide is exported instead of the full
        # source deck. Require both a meaningful affected area and a coherent
        # channel delta before treating it as an unauthorized mutation. Real
        # movement/deletion still exceeds these gates by a wide margin.
        meaningful_area = info['changed_pixels'] > 50 and info['changed_ratio'] > 0.00005
        coherent_delta = info['mean_max_channel_delta'] > 0.20 or info['max_channel_delta'] > 48
        unexpected = meaningful_area and coherent_delta
        add('UNEXPECTED_PIXEL_DIFF','fail' if unexpected else 'pass','blocking' if unexpected else 'info',
            slide_id=sid,source_slide_id=srcid,changed_pixels=info['changed_pixels'],
            checked_pixels=info['checked_pixels'],changed_ratio=round(info['changed_ratio'],8),
            mean_max_channel_delta=round(info['mean_max_channel_delta'],4),max_channel_delta=info['max_channel_delta'],
            threshold=info['threshold'],mask_padding_px=info['mask_padding_px'],
            meaningful_area=meaningful_area,coherent_delta=coherent_delta,
            coherent_delta_thresholds={'mean_max_channel_delta':0.20,'max_channel_delta':48},debug=info['debug'])

    pages=[]
    try:
        pages,pdf_path=_extract_pdf_text(pptx,render_dir)
        bbox_pages=_extract_pdf_words(Path(pdf_path)) if pdf_path else []
        for idx,sid in enumerate(output_ids):
            plan=read_json(root/f'planning/slide_plans/{sid}.json',{})
            page_text=_norm(pages[idx] if idx<len(pages) else '')
            slide=prs.slides[idx] if idx < len(prs.slides) else None
            by_shape={s.shape_id:s for s in slide.shapes} if slide else {}
            rendered_mutation_boxes=[]
            for mut in plan.get('source_object_mutations',[]):
                if mut.get('action')=='replace_table':
                    for r,row in enumerate(mut.get('table_data') or []):
                        for c,cell in enumerate(row):
                            expected_cell=_norm(cell)
                            if not expected_cell: continue
                            ok_cell=expected_cell in page_text
                            add('RENDER_TABLE_TEXT_PRESENT','pass' if ok_cell else 'fail','blocking',slide_id=sid,source_object_id=mut.get('source_object_id'),row=r,column=c,expected=cell)
                    continue
                if mut.get('action')=='replace_chart':
                    expected_series=mut.get('series') or []
                    charts=[shp.chart for shp in slide.shapes if getattr(shp,'has_chart',False)] if slide else []
                    def chart_matches(chart):
                        if len(chart.series) != len(expected_series): return False
                        for got,exp in zip(chart.series,expected_series):
                            try: vals=[float(x) for x in got.values]
                            except Exception: return False
                            target=[float(x) for x in exp.get('values',[])]
                            if len(vals)!=len(target) or any(abs(a-b)>1e-8 for a,b in zip(vals,target)): return False
                        return True
                    ok_chart=any(chart_matches(ch) for ch in charts)
                    add('EDITABLE_CHART_DATA','pass' if ok_chart else 'fail','blocking',slide_id=sid,source_object_id=mut.get('source_object_id'),series_count=len(expected_series))
                    continue
                if mut.get('action')!='replace_text': continue
                expected=_norm(mut.get('render_text',''))
                if not expected: continue

                shp=by_shape.get(_shape_id(mut.get('source_object_id')))
                bbox_page=bbox_pages[idx] if idx < len(bbox_pages) else None
                rb=None; shape_box=None
                matched_words=None
                if shp and bbox_page:
                    # PowerPoint EMU -> PDF point mapping is slide-relative. Using the
                    # rendered page size avoids assuming exactly 72dpi.
                    sx=bbox_page['width']/sw; sy=bbox_page['height']/sh
                    shape_box=(shp.left*sx, shp.top*sy, (shp.left+shp.width)*sx, (shp.top+shp.height)*sy)
                    matched_words=_match_rendered_words(_words_near_box(bbox_page['words'],shape_box),mut.get('render_text',''),shape_box)
                    if matched_words:
                        rb=(min(w['xMin'] for w in matched_words),min(w['yMin'] for w in matched_words),max(w['xMax'] for w in matched_words),max(w['yMax'] for w in matched_words))

                # Global PDF text order interleaves rows across multi-column layouts, so
                # prefer shape-local word matching whenever bbox data is available.
                ok = rb is not None if (shp and bbox_page) else (expected in page_text)
                add('RENDER_TEXT_PRESENT','pass' if ok else 'fail','blocking',slide_id=sid,source_object_id=mut.get('source_object_id'),expected=mut.get('render_text'))
                if rb is None or shape_box is None:
                    if shp and bbox_page: add('RENDER_TEXT_BBOX_MATCH','fail','warning',slide_id=sid,source_object_id=mut.get('source_object_id'))
                    continue
                rendered_mutation_boxes.append((mut.get('source_object_id'),rb))
                tolerance=12.0  # glyph ascender/descender allowance in PDF points
                overflow = rb[0] < shape_box[0]-tolerance or rb[1] < shape_box[1]-tolerance or rb[2] > shape_box[2]+tolerance or rb[3] > shape_box[3]+tolerance
                add('RENDER_TEXT_OVERFLOW','fail' if overflow else 'pass','blocking' if overflow else 'info',slide_id=sid,source_object_id=mut.get('source_object_id'),rendered_bbox=[round(x,2) for x in rb],shape_bbox=[round(x,2) for x in shape_box])
                cb_map={x.get('id'):x for x in plan.get('content_blocks',[])}
                role=(cb_map.get(mut.get('content_block_id')) or {}).get('role')
                for code,severity,detail in _wrap_issues(_group_rendered_lines(matched_words or []),mut.get('render_text',''),role):
                    add(code,'fail',severity,slide_id=sid,source_object_id=mut.get('source_object_id'),**detail)

            layout_plan=read_json(root/f'render/layout_plans/{sid}.json',{})
            for obj in layout_plan.get('generated_objects',[]):
                expected_raw=str(obj.get('render_text') or '')
                if not expected_raw: continue
                shp=by_shape.get(obj.get('output_shape_id'))
                bbox_page=bbox_pages[idx] if idx < len(bbox_pages) else None
                rb=None; shape_box=None; matched_words=None
                if shp and bbox_page:
                    sx=bbox_page['width']/sw; sy=bbox_page['height']/sh
                    shape_box=(shp.left*sx,shp.top*sy,(shp.left+shp.width)*sx,(shp.top+shp.height)*sy)
                    matched_words=_match_rendered_words(_words_near_box(bbox_page['words'],shape_box),expected_raw,shape_box)
                    if matched_words:
                        rb=(min(w['xMin'] for w in matched_words),min(w['yMin'] for w in matched_words),max(w['xMax'] for w in matched_words),max(w['yMax'] for w in matched_words))
                ok=rb is not None if (shp and bbox_page) else (_norm(expected_raw) in page_text)
                add('RENDER_LAYOUT_TEXT_PRESENT','pass' if ok else 'fail','blocking',slide_id=sid,shape_id=obj.get('output_shape_id'),role=obj.get('role'),expected=expected_raw)
                if rb is None or shape_box is None: continue
                rendered_mutation_boxes.append((f"LAYOUT-{obj.get('output_shape_id')}",rb))
                tolerance=12.0
                overflow=rb[0] < shape_box[0]-tolerance or rb[1] < shape_box[1]-tolerance or rb[2] > shape_box[2]+tolerance or rb[3] > shape_box[3]+tolerance
                add('RENDER_LAYOUT_TEXT_OVERFLOW','fail' if overflow else 'pass','blocking' if overflow else 'info',slide_id=sid,shape_id=obj.get('output_shape_id'),role=obj.get('role'),rendered_bbox=[round(x,2) for x in rb],shape_bbox=[round(x,2) for x in shape_box])
                role=obj.get('role')
                wrap_role='title' if role=='title' else None
                for code,severity,detail in _wrap_issues(_group_rendered_lines(matched_words or []),expected_raw,wrap_role):
                    add(code,'fail',severity,slide_id=sid,shape_id=obj.get('output_shape_id'),role=role,**detail)

            for i in range(len(rendered_mutation_boxes)):
                for j in range(i+1,len(rendered_mutation_boxes)):
                    ida,a=rendered_mutation_boxes[i]; idb,b=rendered_mutation_boxes[j]
                    ix=max(0.0,min(a[2],b[2])-max(a[0],b[0])); iy=max(0.0,min(a[3],b[3])-max(a[1],b[1])); area=ix*iy
                    # Glyph bounding boxes can overlap a few points at adjacent baselines
                    # without visible collision. Require meaningful overlap in both axes.
                    if ix > 4.0 and iy > 8.0 and area > 36.0:
                        add('RENDER_TEXT_OVERLAP','fail','blocking',slide_id=sid,source_object_ids=[ida,idb],intersection_area=round(area,2),intersection_width=round(ix,2),intersection_height=round(iy,2))
    except Exception as e:
        add('PDF_TEXT_EXTRACTION','fail','warning',error=str(e))

    image_report = validate_image_harmony(root, output_ids=output_ids, rendered_files=rendered)
    checks.extend(image_report.get('checks', []))

    for check in checks:
        check.setdefault('category', _qa_category(check.get('check_id')))
    blocking=[c for c in checks if c['result']=='fail' and c['severity']=='blocking']
    warning=[c for c in checks if c['result']=='fail' and c['severity']!='blocking']
    category_counts = {category: {
        'fail': sum(1 for c in checks if c.get('category') == category and c.get('result') == 'fail'),
        'blocking': sum(1 for c in checks if c.get('category') == category and c.get('result') == 'fail' and c.get('severity') == 'blocking'),
    } for category in ('technical','content','design')}
    result={
        'schema_version':'0.10','generated_at':now_iso(),'kind':kind,'pptx':str(pptx),
        'status':'pass' if not blocking else 'fail','blocking_count':len(blocking),'warning_count':len(warning),
        'checks':checks,'category_counts':category_counts,'rendered_pages':rendered,'image_harmony':image_report,
        'notes':['v0.10 QA is grouped into Technical / Content / Design. Design review never triggers unlimited automatic repair. v0.6 QA covers package/render, template chrome geometry, mutation/layout-scoped pixel diff, source and layout-generated object bounds, fit-to-intent layout quality, photo-aware slot geometry and negative-space direction, full-bleed coverage/z-order, native comparison/relationship editability, rendered text presence/overflow/overlap, table/chart editability, image harmony, and first-pass Japanese/Latin wrap heuristics. Semantic line-break quality still benefits from Sol review.']
    }
    out=root/'qa'/f'{pptx.stem}_qa.json'; write_json(out,result); result['report']=str(out)
    return result
